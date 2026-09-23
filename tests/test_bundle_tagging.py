"""
Tests for tag_pipeline.py's v2 bundle mode (spec
docs/superpowers/specs/2026-09-13-tag-pipeline-v2-bundle-design.md).
Mocked: tag_pipeline._call_llm_with_json_retry and _load_source_documents are
stubbed; no live API, no PDFs. Prompt behaviour is measured live by the §9
acceptance test, not here.
"""
import hashlib
import os
import pathlib
import uuid

import openpyxl
import pandas as pd
import pypdf.errors
import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import golden_set_pipeline as gsp
import tag_pipeline as tp
from golden_set_pipeline import LLMClient


def test_imports_from_golden_set_pipeline_resolve():
    """Every name tag_pipeline.py imports from golden_set_pipeline.py — a rename
    there fails here, loudly (tag_pipeline is a second consumer)."""
    from golden_set_pipeline import (  # noqa: F401
        LLMClient,
        _call_llm_with_json_retry,
        export_for_review,
        _is_deterministic_claim_id,
        build_chunk_index,
        parse_claims_file,
    read_filing_entity,
        _load_source_documents,
        _derive_claim_id,
        _claims_with_occurrence,
        _chunk_index_of,
        _valid_memo_id,
        _HUMAN_ADDED_PREFIX,
        _QUOTE_SEPARATOR,
    )


def test_bundle_constants():
    assert tp._BUNDLE_MODEL == "google/gemini-3.1-pro-preview"
    assert tp._BUNDLE_MAX_TOKENS == 16000
    assert tp._CONTEXT_CHARS == 500
    assert tp._MAX_CHUNKS_PER_CLAIM == 40
    assert tp._VERDICTS == ("stated_directly", "needs_combining", "not_supported")
    assert tp._DRAFT_FAILED == "draft_failed"
    # the eval spec keys a dense re-review cross-check on this exact prefix
    assert tp._AUTO_NOT_FOUND_REASON == "auto: found=False, nothing was found by search"


def test_bundle_prompt_carries_appendix_a():
    """Distinctive single-line fragments from every part of spec Appendix A --
    a transcription slip fails here, not silently in the live acceptance run."""
    p = tp._BUNDLE_PROMPT
    for token in ("__ENTITY__", "__CLAIM__", "__PASSAGES__"):
        assert token in p, token
    for fragment in (
        "Choose exactly one verdict for the claim, looking at all passages together:",
        '- "stated_directly": at least one passage, on its own, states the claim\'s',
        "rounding still count (EUR 4,203m supports \"EUR 4.2bn\"; 29.04% supports",
        '- "needs_combining": the claim is true, but only by putting passages together,',
        "by a calculation (for example a part divided by a total), or by describing",
        '- "not_supported": the passages, even together, do not establish the claim.',
        "claim's number. Do not pick the closest passage from a set that does not",
        "Only list a passage if its own PASSAGE text carries what the claim needs. Use",
        "alone.",
        "about the numbers, not something the passage says.",
        "300 / 1,200 = 25%.",
        "Same number, different measure.",
        "Answer with JSON only, no other text:",
    ):
        assert fragment in p, fragment
    for n in range(1, 5):
        assert f"\n{n}. Claim " in p, n


def test_bundle_prompt_sha256_matches_template():
    assert tp.bundle_prompt_sha256() == hashlib.sha256(tp._BUNDLE_PROMPT.encode("utf-8")).hexdigest()


# The digest of the prompt the §9 MEMO-004 acceptance test ran, as recorded in
# docs/prompt-verification-log.md → "§9 acceptance (2026-09-13)".
_ACCEPTED_PROMPT_SHA256 = "13a8ca519cc7e0d1098dc414ac97d902d348b2a6be980fc128606a5f0e108edc"


def test_bundle_prompt_is_the_one_the_acceptance_test_ran():
    """A literal, not a recomputation: any edit to _BUNDLE_PROMPT fails here,
    because it voids the recorded acceptance result (CLAUDE.md decision 18).
    If the change is intended, re-run §9 and record the new digest in the log
    before updating this constant."""
    assert tp.bundle_prompt_sha256() == _ACCEPTED_PROMPT_SHA256


# --- Task 2: load_memo_chunks / build_bundle ---------------------------------

CLAIMS_MD = """---
memo_id: MEMO-T
source_folder: sources/acme
---

## Business Profile

1. Acme sells widgets.
"""


def _index(text: str, doc_id: str = "d.pdf"):
    chunks_by_id = {c["chunk_id"]: c for c in gsp.build_chunk_index([(doc_id, text)])}
    return chunks_by_id, {doc_id: text}


def test_read_claims_file_returns_source_folder_and_ids_in_file_order(tmp_path):
    (tmp_path / "MEMO-T.md").write_text(CLAIMS_MD + "2. Acme has 12 plants.\n2. Acme sells widgets.\n")
    source_folder, claims = tp.read_claims_file("MEMO-T", claims_dir=str(tmp_path))
    assert source_folder == "sources/acme"
    texts = ["Acme sells widgets.", "Acme has 12 plants.", "Acme sells widgets."]
    assert [(s, t) for s, t, _ in claims] == [("Business Profile", t) for t in texts]
    assert [c for _, _, c in claims] == [gsp._derive_claim_id("MEMO-T", "Business Profile", t, n)
                                         for t, n in zip(texts, (0, 0, 1))]


def test_read_claims_file_missing_names_the_path_without_the_memo_prefix(tmp_path):
    with pytest.raises(ValueError, match=r"^\S*MEMO-T\.md not found") as excinfo:
        tp.read_claims_file("MEMO-T", claims_dir=str(tmp_path))
    assert not str(excinfo.value).startswith("MEMO-T:")


def test_read_claims_file_refuses_a_folder_at_the_claims_path(tmp_path):
    """A ValueError (so draft and finalize turn it into their refusal), never
    parse_claims_file's IsADirectoryError traceback."""
    (tmp_path / "MEMO-T.md").mkdir()
    with pytest.raises(ValueError, match=r"^\S*MEMO-T\.md is not a file"):
        tp.read_claims_file("MEMO-T", claims_dir=str(tmp_path))


@pytest.mark.skipif(hasattr(os, "geteuid") and os.geteuid() == 0, reason="root can read a mode-000 file")
def test_read_claims_file_unreadable_file_is_a_value_error(tmp_path):
    """An OSError while reading (here: permission denied) becomes a ValueError,
    so draft and finalize refuse instead of showing a traceback. A non-UTF-8
    file needs nothing: UnicodeDecodeError is already a ValueError."""
    path = tmp_path / "MEMO-T.md"
    path.write_text(CLAIMS_MD)
    path.chmod(0)
    try:
        with pytest.raises(ValueError, match=r"^\S*MEMO-T\.md could not be read: .*Permission denied"):
            tp.read_claims_file("MEMO-T", claims_dir=str(tmp_path))
    finally:
        path.chmod(0o644)


def test_read_claims_file_refuses_memo_id_mismatch(tmp_path):
    (tmp_path / "MEMO-T.md").write_text(CLAIMS_MD.replace("memo_id: MEMO-T", "memo_id: MEMO-OTHER"))
    with pytest.raises(ValueError, match="MEMO-OTHER"):
        tp.read_claims_file("MEMO-T", claims_dir=str(tmp_path))


def test_load_memo_chunks_reads_the_given_source_folder(monkeypatch):
    seen = []

    def fake_load(source_folder, label):
        seen.append((source_folder, label))
        return [("a.pdf", "x" * 1500)]

    monkeypatch.setattr(tp, "_load_source_documents", fake_load)
    chunks_by_id, documents = tp.load_memo_chunks("MEMO-T", "sources/acme")
    assert seen == [("sources/acme", "MEMO-T")]
    assert set(chunks_by_id) == {"a.pdf_0", "a.pdf_1"}      # 1000/200 chunking
    assert chunks_by_id["a.pdf_1"]["start_offset"] == 800
    assert documents == {"a.pdf": "x" * 1500}


@pytest.mark.parametrize("error", [
    PermissionError(13, "Permission denied", "sources/acme/a.pdf"),
    pypdf.errors.PdfStreamError("Stream has ended unexpectedly"),
])
def test_load_memo_chunks_unreadable_or_corrupt_pdf_is_a_value_error(monkeypatch, error):
    """An OSError or a pypdf error while loading the PDFs becomes a ValueError,
    so draft and finalize refuse instead of showing a traceback."""
    def fail(folder, label):
        raise error

    monkeypatch.setattr(tp, "_load_source_documents", fail)
    with pytest.raises(ValueError, match=r"^source_folder 'sources/acme': a PDF could not be read \("):
        tp.load_memo_chunks("MEMO-T", "sources/acme")


def test_label_sequence():
    assert [tp._label(i) for i in (0, 1, 25, 26, 39)] == ["A", "B", "Z", "AA", "AN"]


def test_build_bundle_labels_order_dedupe_and_context():
    text = "0123456789" * 300                                   # 3000 chars, no whitespace
    chunks_by_id, documents = _index(text)                     # offsets 0, 800, 1600, 2400
    bundle = tp.build_bundle("claim-1", ["d.pdf_2", "d.pdf_0", "d.pdf_2"], chunks_by_id, documents)
    assert [b["label"] for b in bundle] == ["A", "B"]          # duplicate chunk_id dropped
    assert [b["chunk_id"] for b in bundle] == ["d.pdf_2", "d.pdf_0"]   # input order kept
    assert bundle[0]["doc_id"] == "d.pdf"
    assert bundle[0]["text"] == text[1600:2600]
    assert bundle[0]["before"] == text[1100:1600]              # exactly _CONTEXT_CHARS
    assert bundle[0]["after"] == text[2600:3000]               # only 400 chars remain
    assert bundle[1]["before"] == ""                           # chunk at offset 0


def test_build_bundle_collapses_whitespace():
    text = "word \n\n  " * 400                                 # 3600 chars
    chunks_by_id, documents = _index(text)
    b = tp.build_bundle("claim-1", ["d.pdf_1"], chunks_by_id, documents)[0]
    for key in ("text", "before", "after"):
        assert "\n" not in b[key] and "  " not in b[key], key
        assert b[key].strip() == b[key], key


def test_build_bundle_refuses_more_than_max_chunks(monkeypatch):
    chunks_by_id, documents = _index("0123456789" * 300)
    monkeypatch.setattr(tp, "_MAX_CHUNKS_PER_CLAIM", 2)
    with pytest.raises(ValueError, match=r"claim-9.*3 found chunks.*_MAX_CHUNKS_PER_CLAIM"):
        tp.build_bundle("claim-9", ["d.pdf_0", "d.pdf_1", "d.pdf_2"], chunks_by_id, documents)


def test_build_bundle_refuses_chunk_id_missing_from_index():
    chunks_by_id, documents = _index("0123456789" * 300)
    with pytest.raises(ValueError, match=r"claim-9.*'d\.pdf_77'.*chunk index"):
        tp.build_bundle("claim-9", ["d.pdf_0", "d.pdf_77"], chunks_by_id, documents)


# --- Task 3: _build_bundle_prompt ---------------------------------------------

BUNDLE = [
    {"label": "A", "chunk_id": "x.pdf_1", "doc_id": "x.pdf",
     "text": "PIECE ONE", "before": "BEFORE ONE", "after": "AFTER ONE"},
    {"label": "B", "chunk_id": "y.pdf_7", "doc_id": "y.pdf",
     "text": "PIECE TWO", "before": "", "after": "AFTER TWO"},
]


def test_build_bundle_prompt_fills_every_slot():
    p = tp._build_bundle_prompt("Acme sells widgets.", "Acme Ltd", BUNDLE)
    for token in ("__ENTITY__", "__CLAIM__", "__PASSAGES__"):
        assert token not in p, token
    assert "memo about Acme Ltd. In the passages" in p
    assert "Company\" mean Acme Ltd." in p
    assert "CLAIM: Acme sells widgets.\n" in p
    assert "[A] document: x.pdf\n  before: BEFORE ONE\n  PASSAGE: PIECE ONE\n  after: AFTER ONE" in p
    assert "[B] document: y.pdf\n  before: \n  PASSAGE: PIECE TWO\n  after: AFTER TWO" in p
    assert p.index("[A] document") < p.index("[B] document")
    assert "x.pdf_1" not in p and "y.pdf_7" not in p          # labels only, never chunk_ids
    assert p.rstrip().endswith('the numbers used"}')


def test_build_bundle_prompt_never_rescans_an_inserted_value():
    """The slots are filled in one pass: a claim or entity that happens to
    contain a slot token is inserted as-is, never filled a second time."""
    p = tp._build_bundle_prompt("Claim naming __PASSAGES__ and __ENTITY__.", "Acme __CLAIM__ Ltd", BUNDLE)
    assert "CLAIM: Claim naming __PASSAGES__ and __ENTITY__.\n" in p
    assert "memo about Acme __CLAIM__ Ltd. In the passages" in p
    assert p.count("[A] document: x.pdf") == 1


# --- Task 4: _validate_bundle_answer / draft_claim ----------------------------

LABELS = ["A", "B", "C"]
CLIENT = LLMClient(base_url="http://unused.invalid", api_key="unused", model="test-model")


def test_validate_accepts_and_normalises():
    verdict, needed, reason = tp._validate_bundle_answer(
        {"verdict": " Needs_Combining ", "needed": ["b", "A", "B"], "reason": " B over A. "}, LABELS)
    assert (verdict, needed, reason) == ("needs_combining", ["B", "A"], "B over A.")


@pytest.mark.parametrize("parsed, message", [
    (["a list"], "JSON object"),
    ({"verdict": "maybe", "needed": [], "reason": "x"}, "verdict"),
    ({"verdict": "stated_directly", "needed": "A", "reason": "x"}, "must be a list"),
    ({"verdict": "stated_directly", "needed": ["Z"], "reason": "x"}, "unknown passage label"),
    ({"verdict": "not_supported", "needed": ["A"], "reason": "x"}, "not_supported must list no passages"),
    ({"verdict": "stated_directly", "needed": [], "reason": "x"}, "must list at least one passage"),
    ({"verdict": "needs_combining", "needed": ["A"], "reason": "   "}, "reason"),
    ({"verdict": "needs_combining", "needed": ["A"]}, "reason"),
])
def test_validate_rejects(parsed, message):
    with pytest.raises(ValueError, match=message):
        tp._validate_bundle_answer(parsed, LABELS)


def _fake_llm(responses, calls):
    def fake(prompt, llm_client, context, max_attempts=2, max_tokens=4096):
        calls.append({"prompt": prompt, "model": llm_client.model, "context": context,
                      "max_attempts": max_attempts, "max_tokens": max_tokens})
        response = responses[len(calls) - 1]
        if isinstance(response, Exception):
            raise response
        return response
    return fake


def test_draft_claim_maps_labels_to_chunk_ids(monkeypatch):
    calls = []
    monkeypatch.setattr(tp, "_call_llm_with_json_retry", _fake_llm(
        [{"verdict": "needs_combining", "needed": ["B", "A"], "reason": "B divided by A."}], calls))
    out = tp.draft_claim("Acme sells widgets.", "Acme Ltd", BUNDLE, CLIENT)
    assert out == {"verdict": "needs_combining", "needed_chunk_ids": ["y.pdf_7", "x.pdf_1"],
                   "reason": "B divided by A.", "attempts": 1}
    assert len(calls) == 1
    assert calls[0]["max_attempts"] == 1                   # our loop owns the one retry
    assert calls[0]["max_tokens"] == tp._BUNDLE_MAX_TOKENS
    assert calls[0]["model"] == "test-model"
    assert "CLAIM: Acme sells widgets." in calls[0]["prompt"]


def test_draft_claim_retries_once_after_invalid_answer(monkeypatch):
    calls = []
    monkeypatch.setattr(tp, "_call_llm_with_json_retry", _fake_llm([
        {"verdict": "not_supported", "needed": ["A"], "reason": "inconsistent"},
        {"verdict": "stated_directly", "needed": ["A"], "reason": "A says it."},
    ], calls))
    out = tp.draft_claim("Acme sells widgets.", "Acme Ltd", BUNDLE, CLIENT)
    assert out == {"verdict": "stated_directly", "needed_chunk_ids": ["x.pdf_1"],
                   "reason": "A says it.", "attempts": 2}
    assert len(calls) == 2


def test_draft_claim_marks_draft_failed_after_two_bad_attempts(monkeypatch):
    calls = []
    monkeypatch.setattr(tp, "_call_llm_with_json_retry", _fake_llm([
        ValueError("Expecting value: line 1 column 1"),
        {"verdict": "maybe", "needed": [], "reason": "x"},
    ], calls))
    out = tp.draft_claim("Acme sells widgets.", "Acme Ltd", BUNDLE, CLIENT)
    assert out["verdict"] == tp._DRAFT_FAILED
    assert out["needed_chunk_ids"] == []
    assert out["attempts"] == 2
    assert out["reason"].startswith("DRAFT FAILED: ")
    assert "Expecting value" in out["reason"] and "verdict" in out["reason"]
    assert len(calls) == 2                                  # never a third call


def test_draft_claim_empty_bundle_makes_no_call(monkeypatch):
    monkeypatch.setattr(tp, "_call_llm_with_json_retry",
                        lambda *a, **k: pytest.fail("no LLM call for a claim with no chunks"))
    out = tp.draft_claim("Acme sells widgets.", "Acme Ltd", [], CLIENT)
    assert out == {"verdict": "not_supported", "needed_chunk_ids": [],
                   "reason": "auto: found=False, nothing was found by search", "attempts": 0}


def test_draft_claim_never_makes_a_third_call_to_call_llm(monkeypatch, caplog):
    """Mocks the one HTTP contact point itself, so the retry helper's own
    attempt loop is inside the test: two garbage answers -> exactly two calls."""
    calls = []

    def fake_call_llm(prompt, llm_client, temperature=0.0, max_tokens=4096):
        calls.append(max_tokens)
        return "not json at all"

    monkeypatch.setattr(gsp, "call_llm", fake_call_llm)
    with caplog.at_level("WARNING"):
        out = tp.draft_claim("Acme sells widgets in forty countries.", "Acme Ltd", BUNDLE, CLIENT)
    assert out["verdict"] == tp._DRAFT_FAILED
    assert calls == [tp._BUNDLE_MAX_TOKENS, tp._BUNDLE_MAX_TOKENS]
    assert "Acme sells widgets in forty countries." in caplog.text     # the log names the claim


# --- Plan A: prepare_draft ----------------------------------------------------

MEMO = "MEMO-T"
SECTION = "Business Profile"
C1, C2 = "Acme sells widgets.", "Acme has 12 plants."
CLAIMS_MD_TWO = CLAIMS_MD + "2. Acme has 12 plants.\n"
DOC_TEXT = "0123456789" * 300                 # d.pdf -> chunks d.pdf_0 .. d.pdf_3
# What `build` stored for each chunk; prepare_draft refuses a checkpoint whose
# chunk_text differs from the rebuilt index.
REAL_CHUNK_TEXT = {c["chunk_id"]: c["chunk_text"] for c in gsp.build_chunk_index([("d.pdf", DOC_TEXT)])}


def _cid(claim_text, memo=MEMO):
    return gsp._derive_claim_id(memo, SECTION, claim_text, 0)


def _ckpt_row(claim_text, chunk_id=None, memo=MEMO, **overrides):
    """One golden-set checkpoint row: a found row when chunk_id is given, else
    the claim's single found=False row."""
    found = chunk_id is not None
    row = {"claim_id": _cid(claim_text, memo), "memo_id": memo, "section": SECTION,
           "claim_text": claim_text,
           "doc_id": "d.pdf" if found else None, "chunk_id": chunk_id,
           "chunk_text": REAL_CHUNK_TEXT.get(chunk_id, f"text of {chunk_id}") if found else None,
           "bm25_score": 1.5 if found else None,
           "evidence_span": f"span of {chunk_id}" if found else None,
           "found": found, "confidence": "high" if found else None,
           "ambiguous_match": False, "verbatim_match": True if found else None,
           "human_reviewed": False, "tag": None}
    row.update(overrides)
    return row


ROWS = [
    _ckpt_row(C2),                                             # file order, not checkpoint order, wins
    _ckpt_row(C1, "d.pdf_2"),
    _ckpt_row(C1, "d.pdf_0"),
]


def _setup(tmp_path, monkeypatch, rows, claims=None, entity="Acme Ltd"):
    """Checkpoint and claims files at their default names in tmp_path (made the
    cwd); every claims file gets `filing_entity: <entity>` in its frontmatter
    (none when entity is None); the PDF loader stubbed."""
    monkeypatch.chdir(tmp_path)
    pd.DataFrame(rows).to_parquet("golden_set_checkpoint.parquet")
    (tmp_path / "claims").mkdir()
    for memo_id, md in (claims if claims is not None else {MEMO: CLAIMS_MD_TWO}).items():
        if entity is not None:
            md = md.replace("\nsource_folder:", f"\nfiling_entity: {entity}\nsource_folder:", 1)
        (tmp_path / "claims" / f"{memo_id}.md").write_text(md)
    monkeypatch.setattr(tp, "_load_source_documents", lambda folder, label: [("d.pdf", DOC_TEXT)])


def _no_pdfs(monkeypatch):
    monkeypatch.setattr(tp, "_load_source_documents",
                        lambda *a: pytest.fail("must refuse before reading any PDF"))


def _existing_sheet(tmp_path, text="the user's checked work"):
    sheet = tmp_path / "review" / f"{MEMO}.xlsx"
    sheet.parent.mkdir(exist_ok=True)
    sheet.write_text(text)
    return sheet


def test_prepare_draft_builds_every_bundle_in_claims_file_order(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch, ROWS)
    plan = tp.prepare_draft([])
    assert len(plan) == 1
    memo = plan[0]
    assert (memo["memo_id"], memo["action"], memo["entity"]) == (MEMO, "draft", "Acme Ltd")
    assert memo["review_path"] == os.path.join("review", f"{MEMO}.xlsx")
    c1, c2 = memo["claims"]
    assert (c1["number"], c1["claim_id"], c1["section"], c1["claim_text"]) == (1, _cid(C1), SECTION, C1)
    assert [b["chunk_id"] for b in c1["bundle"]] == ["d.pdf_2", "d.pdf_0"]
    assert [r["chunk_id"] for r in c1["rows"]] == ["d.pdf_2", "d.pdf_0"]     # aligned with bundle
    assert c1["rows"][0]["evidence_span"] == "span of d.pdf_2"
    assert set(c1["rows"][0]) == set(tp._CHECKPOINT_FIELDS)
    assert (c2["number"], c2["claim_text"], c2["bundle"], c2["rows"]) == (2, C2, [], [])


def test_prepare_draft_refuses_missing_checkpoint(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    with pytest.raises(tp.DraftRefused, match="golden_set_checkpoint.parquet not found"):
        tp.prepare_draft([])


def test_prepare_draft_refuses_unknown_memo_id(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch, ROWS)
    _no_pdfs(monkeypatch)
    with pytest.raises(tp.DraftRefused, match=r"MEMO-NOPE: not in golden_set_checkpoint\.parquet"):
        tp.prepare_draft([MEMO, "MEMO-NOPE"])


def test_prepare_draft_refuses_memo_without_filing_entity(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch, ROWS, entity=None)
    _no_pdfs(monkeypatch)
    with pytest.raises(tp.DraftRefused, match=r"MEMO-T: claims.MEMO-T\.md has no filing_entity"):
        tp.prepare_draft([])


def test_prepare_draft_checks_every_filing_entity_before_reading_any_pdf(tmp_path, monkeypatch):
    """The SECOND memo lacking filing_entity refuses the run before the FIRST
    memo's PDFs are read (they take ~26 s on a real filing)."""
    other = "MEMO-U"
    rows = ROWS + [_ckpt_row(C1, "d.pdf_2", memo=other), _ckpt_row(C2, memo=other)]
    _setup(tmp_path, monkeypatch, rows, claims={MEMO: CLAIMS_MD_TWO})
    (tmp_path / "claims" / f"{other}.md").write_text(CLAIMS_MD_TWO.replace("memo_id: MEMO-T", f"memo_id: {other}"))
    _no_pdfs(monkeypatch)
    with pytest.raises(tp.DraftRefused, match=r"MEMO-U: .*has no filing_entity"):
        tp.prepare_draft([])


def test_prepare_draft_refuses_blank_filing_entity(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch, ROWS, entity="''")
    _no_pdfs(monkeypatch)
    with pytest.raises(tp.DraftRefused, match="MEMO-T: .*'filing_entity' must be a non-blank one-line string"):
        tp.prepare_draft([])


def test_prepare_draft_refuses_non_uuid5_claim_id(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch, ROWS + [_ckpt_row("Acme is old.", claim_id=str(uuid.uuid4()))])
    _no_pdfs(monkeypatch)
    with pytest.raises(tp.DraftRefused, match=r"MEMO-T: 1 claim_id\(s\) are not deterministic uuid5"):
        tp.prepare_draft([])


def test_prepare_draft_refuses_too_many_chunks_even_when_sheet_exists(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch, ROWS)
    _no_pdfs(monkeypatch)
    _existing_sheet(tmp_path)
    monkeypatch.setattr(tp, "_MAX_CHUNKS_PER_CLAIM", 1)
    with pytest.raises(tp.DraftRefused, match=rf"MEMO-T: claim {_cid(C1)} has 2 found chunks"):
        tp.prepare_draft([])


def test_prepare_draft_refuses_a_checkpoint_that_cannot_be_read(tmp_path, monkeypatch):
    """pyarrow raises ArrowInvalid (a ValueError _main does not catch) for a
    damaged parquet file: a refusal naming the fix, not a traceback."""
    _setup(tmp_path, monkeypatch, ROWS)
    _no_pdfs(monkeypatch)
    pathlib.Path("golden_set_checkpoint.parquet").write_bytes(b"not a parquet file")
    with pytest.raises(tp.DraftRefused, match=r"^golden_set_checkpoint\.parquet could not be read \(ArrowInvalid: "
                                               r".*\) — re-run `python golden_set_pipeline\.py build`"):
        tp.prepare_draft([MEMO])


def test_prepare_draft_refuses_a_checkpoint_missing_a_column(tmp_path, monkeypatch):
    """An older or hand-edited checkpoint without a column draft reads is
    refused up front, not a KeyError partway through the checks."""
    _setup(tmp_path, monkeypatch, [{k: v for k, v in row.items() if k != "confidence"} for row in ROWS])
    _no_pdfs(monkeypatch)
    with pytest.raises(tp.DraftRefused, match=r"^golden_set_checkpoint\.parquet is missing column\(s\) "
                                               r"\['confidence'\] — re-run `python golden_set_pipeline\.py build`"):
        tp.prepare_draft([MEMO])


def test_prepare_draft_refuses_checkpoint_claim_text_that_differs_from_claims_file(tmp_path, monkeypatch):
    """The id sets agree, but one checkpoint row under C1's id carries other
    text: the bundle would be judged against the claims file's claim while
    its evidence rows came from something else. Refused before any PDF."""
    rows = [ROWS[0], dict(ROWS[1], claim_text="Acme sells gadgets."), ROWS[2]]
    _setup(tmp_path, monkeypatch, rows)
    _no_pdfs(monkeypatch)
    with pytest.raises(tp.DraftRefused, match=rf"^{MEMO}: claim {_cid(C1)}: the checkpoint's claim_text "
                                               rf"'Acme sells gadgets\.' is not claims.{MEMO}\.md's "
                                               r"'Acme sells widgets\.' .* re-run `python golden_set_pipeline\.py build`"):
        tp.prepare_draft([MEMO])


def test_prepare_draft_refuses_a_review_folder_that_is_a_file(tmp_path, monkeypatch):
    """run_draft could only create review/ after every LLM call was spent;
    a file in its place is refused before any PDF, credential or call."""
    _setup(tmp_path, monkeypatch, ROWS)
    _no_pdfs(monkeypatch)
    (tmp_path / "review").write_text("not a folder")
    with pytest.raises(tp.DraftRefused, match=r"^review exists but is not a folder"):
        tp.prepare_draft([MEMO])


def test_prepare_draft_refuses_duplicate_found_rows_for_one_chunk(tmp_path, monkeypatch):
    """A checkpoint built before _rows_for_claim collapsed same-chunk_id matches
    can hold two rows for one (claim, chunk). Keeping either would guess which
    confidence/span is right, so build must be re-run. Refused before any PDF."""
    _setup(tmp_path, monkeypatch, [*ROWS, dict(ROWS[1], confidence="low")])
    _no_pdfs(monkeypatch)
    with pytest.raises(tp.DraftRefused, match=rf"^{MEMO}: claim {_cid(C1)}: chunk 'd\.pdf_2' appears in 2 found "
                                               r"rows .* re-run `python golden_set_pipeline\.py build`"):
        tp.prepare_draft([MEMO])


@pytest.mark.parametrize("where, message", [
    ("review", r"^review exists but is not a folder"),
    (os.path.join("review", f"{MEMO}.xlsx"), rf"^{MEMO}: review.{MEMO}\.xlsx exists but is not a file"),
])
def test_prepare_draft_refuses_a_dangling_symlink_destination(tmp_path, monkeypatch, where, message):
    """A symlink to nothing (os.path.exists is False for it, os.path.lexists
    True) is refused by prepare_draft, before any PDF, credential or LLM
    call — rather than failing at makedirs / publish after the calls."""
    _setup(tmp_path, monkeypatch, ROWS)
    _no_pdfs(monkeypatch)
    os.makedirs(os.path.dirname(where) or ".", exist_ok=True)
    os.symlink("missing-target", where)
    with pytest.raises(tp.DraftRefused, match=message):
        tp.prepare_draft([MEMO])


def test_prepare_draft_refuses_a_review_path_that_is_not_a_file(tmp_path, monkeypatch):
    """A folder named review/<memo_id>.xlsx is not a finished sheet: skipping it
    would skip the memo on every run. Refused before any PDF is read."""
    _setup(tmp_path, monkeypatch, ROWS)
    _no_pdfs(monkeypatch)
    (tmp_path / "review" / f"{MEMO}.xlsx").mkdir(parents=True)
    with pytest.raises(tp.DraftRefused, match=rf"{MEMO}: review.{MEMO}\.xlsx exists but is not a file"):
        tp.prepare_draft([MEMO])


def test_prepare_draft_skips_memo_whose_sheet_exists_without_reading_it(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch, ROWS, claims={})              # not even a claims file
    _no_pdfs(monkeypatch)
    _existing_sheet(tmp_path)
    assert tp.prepare_draft([MEMO]) == [{"memo_id": MEMO, "action": "skip",
                                         "review_path": os.path.join("review", f"{MEMO}.xlsx"),
                                         "entity": None, "claims": []}]


def test_prepare_draft_refuses_found_values_that_are_not_booleans(tmp_path, monkeypatch):
    """A hand-edited or malformed checkpoint with found as text: 'FALSE' would
    fail .eq(True) and the claim would silently become auto not-found,
    dropping its evidence. Refused before any PDF."""
    _setup(tmp_path, monkeypatch, [dict(row, found="TRUE" if row["found"] else "FALSE") for row in ROWS])
    _no_pdfs(monkeypatch)
    with pytest.raises(tp.DraftRefused, match=rf"^{MEMO}: 3 checkpoint row\(s\) have a found value that is not "
                                               r"TRUE/FALSE \(first: 'FALSE'\) — re-run `python golden_set_pipeline\.py build`"):
        tp.prepare_draft([MEMO])


def test_prepare_draft_refuses_build_error_rows(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch, ROWS + [_ckpt_row("Acme is old.", found=None, confidence="error")])
    _no_pdfs(monkeypatch)
    with pytest.raises(tp.DraftRefused, match=r"MEMO-T: 1 build-error row\(s\).*re-run"):
        tp.prepare_draft([])


@pytest.mark.parametrize("claims_md, extra_rows, message", [
    (CLAIMS_MD_TWO + "3. Acme is old.\n", [], r"1 claim\(s\) only in the claims file"),
    (CLAIMS_MD_TWO, [_ckpt_row("Acme is old.")], r"1 only in the checkpoint"),
])
def test_prepare_draft_refuses_claims_file_checkpoint_mismatch(tmp_path, monkeypatch,
                                                               claims_md, extra_rows, message):
    _setup(tmp_path, monkeypatch, ROWS + extra_rows, claims={MEMO: claims_md})
    _no_pdfs(monkeypatch)
    with pytest.raises(tp.DraftRefused, match=rf"MEMO-T: claims file and checkpoint disagree.*{message}"):
        tp.prepare_draft([])


def test_prepare_draft_refuses_missing_claims_file(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch, ROWS, claims={})
    _no_pdfs(monkeypatch)
    with pytest.raises(tp.DraftRefused, match=r"MEMO-T: claims[/\\]MEMO-T\.md not found") as excinfo:
        tp.prepare_draft([])
    assert str(excinfo.value).count("MEMO-T:") == 1          # memo prefix added once, not doubled


def test_prepare_draft_refuses_malformed_claims_file(tmp_path, monkeypatch):
    """A claims file parse_claims_file itself rejects (here: stray prose glued
    directly under the last claim, no blank line before it) must surface as
    DraftRefused, not an uncaught ValueError traceback."""
    malformed = CLAIMS_MD_TWO + "some stray prose line\n"
    _setup(tmp_path, monkeypatch, ROWS, claims={MEMO: malformed})
    _no_pdfs(monkeypatch)
    with pytest.raises(tp.DraftRefused, match=r"MEMO-T: "):
        tp.prepare_draft([])


def test_prepare_draft_checks_every_memo_before_returning(tmp_path, monkeypatch):
    """A missing chunk in the SECOND memo refuses the whole run, named by memo —
    the first memo's bundles being fine does not let a run start."""
    other = "MEMO-U"
    rows = ROWS + [_ckpt_row(C1, "d.pdf_77", memo=other), _ckpt_row(C2, memo=other)]
    _setup(tmp_path, monkeypatch, rows,
           claims={MEMO: CLAIMS_MD_TWO, other: CLAIMS_MD_TWO.replace("memo_id: MEMO-T", f"memo_id: {other}")})
    with pytest.raises(tp.DraftRefused, match=r"MEMO-U: claim .*'d\.pdf_77'.*chunk index"):
        tp.prepare_draft([])


@pytest.mark.parametrize("override, what", [
    ({"chunk_text": "text from a different chunking"}, "text"),
    ({"doc_id": "other.pdf"}, "document"),
])
def test_prepare_draft_refuses_chunk_that_differs_from_checkpoint(tmp_path, monkeypatch, override, what):
    """The same chunk_id can name a different passage (non-default chunk_size /
    overlap, or an edited PDF). Refused, so the draft never judges text the
    sheet does not carry."""
    rows = list(ROWS)
    rows[1] = {**rows[1], **override}                        # C1's d.pdf_2 row
    _setup(tmp_path, monkeypatch, rows)
    with pytest.raises(tp.DraftRefused, match=rf"MEMO-T: claim {_cid(C1)}: chunk 'd\.pdf_2' .*different {what}"):
        tp.prepare_draft([])


# --- Plan A: write_review_sheet -------------------------------------------------

def _drafted_claims(tmp_path, monkeypatch):
    """prepare_draft's claims for ROWS plus a third one-passage claim, each given
    a draft: needs_combining on B; no pieces; draft failed."""
    rows = ROWS + [_ckpt_row("Acme is old.", "d.pdf_3", evidence_span="=SUM(A1) looks like a formula")]
    _setup(tmp_path, monkeypatch, rows, claims={MEMO: CLAIMS_MD_TWO + "3. Acme is old.\n"})
    claims = tp.prepare_draft([])[0]["claims"]
    drafts = [
        {"verdict": "needs_combining", "needed_chunk_ids": ["d.pdf_0"], "reason": "B over the total.", "attempts": 1},
        {"verdict": "not_supported", "needed_chunk_ids": [], "reason": tp._AUTO_NOT_FOUND_REASON, "attempts": 0},
        {"verdict": tp._DRAFT_FAILED, "needed_chunk_ids": [], "reason": "DRAFT FAILED: a | b", "attempts": 2},
    ]
    return [{**claim, "draft": draft} for claim, draft in zip(claims, drafts, strict=True)]


def _sheet_records(path):
    """The bundles tab as one {column key: value} dict per row below the header."""
    ws = load_workbook(path)["bundles"]
    key_by_header = {header: key for key, header in tp._SHEET_HEADERS.items()}
    keys = [key_by_header[cell.value] for cell in ws[1]]
    return [dict(zip(keys, (cell.value for cell in row), strict=True)) for row in ws.iter_rows(min_row=2)]


def test_sheet_columns_contract():
    cols = tp._SHEET_COLUMNS
    assert len(cols) == len(set(cols))
    assert cols[:10] == ("block", "text", "document", "before", "after", "quoted_span",
                         "verdict_needed", "ai_reason", "checked", "note")
    assert cols[10:] == tp._SHEET_HIDDEN
    for key in ("row_kind", "claim_id", "memo_id", "section", "claim_text",
                *tp._CHECKPOINT_FIELDS, "ai_verdict", "ai_needed", "chunk_count", "source_docs"):
        assert key in tp._SHEET_HIDDEN, key
    assert set(tp._SHEET_HEADERS) == set(cols)
    assert len(set(tp._SHEET_HEADERS.values())) == len(cols)        # unique -> readable back by header


def test_review_sheet_rows_and_values(tmp_path, monkeypatch):
    claims = _drafted_claims(tmp_path, monkeypatch)
    path = str(tmp_path / "out.xlsx")
    tp.write_review_sheet(path, MEMO, claims)

    wb = load_workbook(path)
    assert wb.sheetnames == ["how to", "bundles"]
    how_to = [cell.value for cell in wb["how to"]["A"]]
    assert how_to[0] == f"TAG REVIEW — {MEMO}"
    assert "Save with the SAME file name." in how_to
    assert "5. DRAFT FAILED? The AI could not answer that claim. Pick the verdict" in how_to
    assert how_to.index("5. DRAFT FAILED? The AI could not answer that claim. Pick the verdict") < \
        how_to.index("Save with the SAME file name.")

    recs = _sheet_records(path)
    assert [r["row_kind"] for r in recs] == ["claim", "chunk", "chunk", "add", "add",
                                            "claim", "add", "add",
                                            "claim", "chunk", "add", "add"]
    assert all(r["memo_id"] == MEMO and r["claim_id"] for r in recs)
    # each claim row carries how many chunk rows draft wrote under it; no other row does
    assert [r["chunk_count"] for r in recs] == [2, None, None, None, None, 0, None, None, 1, None, None, None]
    # and the memo's PDFs draft read (sorted doc_ids), so finalize can tell if they changed
    sources = tp._source_docs({"d.pdf": DOC_TEXT})
    assert sources.startswith("d.pdf (") and len(sources) == len("d.pdf (") + 12 + 1   # name + text fingerprint
    assert [r["source_docs"] for r in recs] == [sources, None, None, None, None, sources, None, None,
                                                sources, None, None, None]

    claim1, a, b = recs[0], recs[1], recs[2]
    assert (claim1["block"], claim1["text"], claim1["verdict_needed"], claim1["ai_verdict"],
            claim1["ai_reason"], claim1["checked"], claim1["claim_id"]) == \
        (1, C1, "needs combining", "needs_combining", "B over the total.", None, _cid(C1))
    assert (a["block"], a["chunk_id"], a["verdict_needed"], a["ai_needed"]) == ("A", "d.pdf_2", None, None)
    assert (b["block"], b["chunk_id"], b["verdict_needed"], b["ai_needed"]) == ("B", "d.pdf_0", "yes", "yes")
    passage = claims[0]["bundle"][1]
    assert (b["text"], b["before"], b["after"], b["document"]) == \
        (passage["text"], None, passage["after"], "d.pdf")              # chunk at offset 0: empty before
    assert (b["chunk_text"], b["evidence_span"], b["quoted_span"], b["bm25_score"], b["found"],
            b["confidence"], b["ambiguous_match"], b["verbatim_match"], b["section"], b["claim_text"]) == \
        (REAL_CHUNK_TEXT["d.pdf_0"], "span of d.pdf_0", "span of d.pdf_0", 1.5, True, "high", False, True, SECTION, C1)
    assert (recs[3]["block"], recs[3]["verdict_needed"], recs[3]["claim_id"]) == ("+", None, _cid(C1))

    claim2 = recs[5]
    assert (claim2["block"], claim2["verdict_needed"], claim2["ai_verdict"]) == (2, "not supported", "not_supported")
    assert claim2["ai_reason"].startswith("auto: found=False")

    claim3, piece = recs[8], recs[9]
    assert (claim3["block"], claim3["verdict_needed"], claim3["ai_verdict"]) == (3, "draft failed", "draft_failed")
    assert piece["evidence_span"] == "=SUM(A1) looks like a formula"    # stored as text, not a formula

    # Values alone don't prove it: openpyxl reloads a real formula cell with
    # the same string value, so also check the cell's stored type — find the
    # row by content rather than hard-coding a row number.
    ws = wb["bundles"]
    col = {key: i for i, key in enumerate(tp._SHEET_COLUMNS, start=1)}
    formula_row = next(r for r in range(2, ws.max_row + 1)
                       if ws.cell(row=r, column=col["evidence_span"]).value == "=SUM(A1) looks like a formula")
    for key in ("evidence_span", "quoted_span"):
        assert ws.cell(row=formula_row, column=col[key]).data_type == "s"


def test_review_sheet_strips_characters_excel_cannot_store(tmp_path, monkeypatch):
    """pypdf can emit NUL and other control bytes (real MEMO-001/002 PDFs have
    them); openpyxl refuses such text, so it is cleaned, never a crash."""
    claims = _drafted_claims(tmp_path, monkeypatch)
    passage, ckpt = claims[0]["bundle"][1], claims[0]["rows"][1]
    passage["text"], passage["after"] = "piece\x00 text", "after\x0b text"
    ckpt["chunk_text"] = "raw\x1f text"
    path = str(tmp_path / "out.xlsx")
    tp.write_review_sheet(path, MEMO, claims)
    b = _sheet_records(path)[2]
    assert (b["text"], b["after"], b["chunk_text"]) == ("piece text", "after text", "raw text")


def test_review_sheet_hidden_columns_dropdowns_and_fill(tmp_path, monkeypatch):
    path = str(tmp_path / "out.xlsx")
    tp.write_review_sheet(path, MEMO, _drafted_claims(tmp_path, monkeypatch))
    ws = load_workbook(path)["bundles"]
    col = {key: i for i, key in enumerate(tp._SHEET_COLUMNS, start=1)}

    hidden = set()
    for dim in ws.column_dimensions.values():        # openpyxl stores alike neighbouring columns as one entry
        if dim.hidden:
            hidden.update(range(dim.min, dim.max + 1))
    assert hidden == {col[key] for key in tp._SHEET_HIDDEN}

    verdict, checked = get_column_letter(col["verdict_needed"]), get_column_letter(col["checked"])
    lists = {dv.formula1: dv for dv in ws.data_validations.dataValidation}
    verdict_dv, yes_dv = lists['"stated directly,needs combining,not supported"'], lists['"yes"']
    assert f"{verdict}2" in verdict_dv and f"{checked}2" in yes_dv      # claim row
    assert f"{verdict}3" in yes_dv and f"{verdict}5" in yes_dv          # chunk row, add row
    assert f"{verdict}3" not in verdict_dv

    assert ws["A2"].fill.fgColor.rgb.endswith(tp._CLAIM_FILL)           # claim row is blue
    assert not ws["A3"].fill.fgColor.rgb.endswith(tp._CLAIM_FILL)
    assert ws.freeze_panes == "C2"


def test_review_sheet_never_overwrites(tmp_path, monkeypatch):
    claims = _drafted_claims(tmp_path, monkeypatch)
    path = tmp_path / "out.xlsx"
    path.write_text("the user's checked work")
    with pytest.raises(FileExistsError, match="never overwritten"):
        tp.write_review_sheet(str(path), MEMO, claims)
    assert path.read_text() == "the user's checked work"


def test_review_sheet_never_replaces_a_sheet_that_appears_during_the_save(tmp_path, monkeypatch):
    """Another run publishes the same sheet after the exists-check but before
    this publish: the publish fails instead of replacing that sheet."""
    claims = _drafted_claims(tmp_path, monkeypatch)
    path = tmp_path / "out.xlsx"
    real_save = openpyxl.Workbook.save

    def save_then_another_run_publishes(self, filename):
        real_save(self, filename)
        path.write_text("the other run's sheet")

    monkeypatch.setattr(openpyxl.Workbook, "save", save_then_another_run_publishes)
    with pytest.raises(FileExistsError):
        tp.write_review_sheet(str(path), MEMO, claims)
    assert path.read_text() == "the other run's sheet"
    assert not list(tmp_path.glob("*.tmp"))


def test_review_sheet_failed_save_leaves_no_file(tmp_path, monkeypatch):
    claims = _drafted_claims(tmp_path, monkeypatch)

    def half_written_then_boom(self, filename):
        pathlib.Path(filename).write_text("half a workbook")
        raise OSError("disk full")

    monkeypatch.setattr(openpyxl.Workbook, "save", half_written_then_boom)
    path = tmp_path / "out.xlsx"
    with pytest.raises(OSError, match="disk full"):
        tp.write_review_sheet(str(path), MEMO, claims)
    assert not path.exists()
    assert not list(tmp_path.glob("*.tmp"))


# --- Plan A: run_draft ----------------------------------------------------------

def _stated_first(bundle):
    return {"verdict": "stated_directly", "needed_chunk_ids": [bundle[0]["chunk_id"]],
            "reason": "A says it.", "attempts": 1}


def _failed(bundle):
    return {"verdict": tp._DRAFT_FAILED, "needed_chunk_ids": [], "reason": "DRAFT FAILED: x | y", "attempts": 2}


def _fake_drafts(monkeypatch, answer, calls):
    """Stub draft_claim: `answer(bundle)` for a claim with pieces; the real
    no-call path for a claim without."""
    real = tp.draft_claim

    def fake(claim_text, entity, bundle, llm_client):
        calls.append((claim_text, entity, [b["chunk_id"] for b in bundle], llm_client))
        return answer(bundle) if bundle else real(claim_text, entity, bundle, llm_client)

    monkeypatch.setattr(tp, "draft_claim", fake)


def test_run_draft_writes_sheet_and_summary(tmp_path, monkeypatch, caplog):
    _setup(tmp_path, monkeypatch, ROWS)
    calls = []
    _fake_drafts(monkeypatch, _stated_first, calls)
    with caplog.at_level("INFO"):
        counts = tp.run_draft(tp.prepare_draft([]), CLIENT)
    assert counts == {"written": 1, "skipped": 0, "failed": 0}
    assert calls == [(C1, "Acme Ltd", ["d.pdf_2", "d.pdf_0"], CLIENT), (C2, "Acme Ltd", [], CLIENT)]
    sheet = tmp_path / "review" / f"{MEMO}.xlsx"
    assert sheet.exists() and not list((tmp_path / "review").glob("*.tmp"))
    assert f"{MEMO}: written (1 drafted, 1 with no pieces, 0 draft failed)" in caplog.text
    recs = _sheet_records(str(sheet))
    assert [r["verdict_needed"] for r in recs if r["row_kind"] == "claim"] == ["stated directly", "not supported"]


def test_run_draft_writes_memo_with_some_failed_claims(tmp_path, monkeypatch, caplog):
    _setup(tmp_path, monkeypatch, ROWS[1:] + [_ckpt_row(C2, "d.pdf_1")])    # both claims have pieces
    _fake_drafts(monkeypatch,
                 lambda bundle: _failed(bundle) if bundle[0]["chunk_id"] == "d.pdf_1" else _stated_first(bundle),
                 [])
    with caplog.at_level("INFO"):
        counts = tp.run_draft(tp.prepare_draft([]), CLIENT)
    assert counts == {"written": 1, "skipped": 0, "failed": 0}
    assert f"{MEMO}: written (1 drafted, 0 with no pieces, 1 draft failed)" in caplog.text


def test_run_draft_memo_fails_when_every_call_fails_despite_a_no_pieces_claim(tmp_path, monkeypatch, caplog):
    """An outage: the one claim that needed no call must not turn the memo into
    "written" — a written sheet full of `draft failed` would be skipped by every
    later run instead of retried."""
    _setup(tmp_path, monkeypatch, ROWS)                     # C1 has pieces, C2 has none
    _fake_drafts(monkeypatch, _failed, [])
    with caplog.at_level("INFO"):
        counts = tp.run_draft(tp.prepare_draft([]), CLIENT)
    assert counts == {"written": 0, "skipped": 0, "failed": 1}
    assert not (tmp_path / "review" / f"{MEMO}.xlsx").exists()
    assert f"{MEMO}: failed (all 1 LLM calls came back draft failed)" in caplog.text


def test_run_draft_memo_fails_when_every_claim_fails(tmp_path, monkeypatch, caplog):
    _setup(tmp_path, monkeypatch, ROWS[1:], claims={MEMO: CLAIMS_MD})     # one claim, with pieces
    _fake_drafts(monkeypatch, _failed, [])
    with caplog.at_level("INFO"):
        counts = tp.run_draft(tp.prepare_draft([]), CLIENT)
    assert counts == {"written": 0, "skipped": 0, "failed": 1}
    assert not (tmp_path / "review" / f"{MEMO}.xlsx").exists()
    assert f"{MEMO}: failed" in caplog.text


def test_run_draft_skipped_memo_makes_no_calls(tmp_path, monkeypatch, caplog):
    _setup(tmp_path, monkeypatch, ROWS)
    sheet = _existing_sheet(tmp_path)
    monkeypatch.setattr(tp, "draft_claim", lambda *a: pytest.fail("no calls for a skipped memo"))
    with caplog.at_level("INFO"):
        counts = tp.run_draft(tp.prepare_draft([]), None)
    assert counts == {"written": 0, "skipped": 1, "failed": 0}
    assert sheet.read_text() == "the user's checked work"
    assert f"{MEMO}: skipped" in caplog.text


@pytest.mark.parametrize("error, outcome, detail", [
    (OSError("disk full"), "failed", "could not save"),
    (FileExistsError("published by another run"), "skipped", "appeared while drafting"),
])
def test_run_draft_save_problem_is_counted_and_later_memos_still_run(tmp_path, monkeypatch, caplog,
                                                                     error, outcome, detail):
    _setup(tmp_path, monkeypatch, ROWS)
    _fake_drafts(monkeypatch, _stated_first, [])

    def save_fails(path, memo_id, claims):
        raise error

    monkeypatch.setattr(tp, "write_review_sheet", save_fails)
    later = {"memo_id": "MEMO-LATER", "action": "skip", "review_path": os.path.join("review", "MEMO-LATER.xlsx"),
             "entity": None, "claims": []}
    with caplog.at_level("INFO"):
        counts = tp.run_draft(tp.prepare_draft([]) + [later], CLIENT)
    expected = {"written": 0, "skipped": 1, "failed": 0}
    expected[outcome] += 1
    assert counts == expected
    assert f"{MEMO}: {outcome} — " in caplog.text and detail in caplog.text
    assert "MEMO-LATER: skipped" in caplog.text                     # the run went on


def test_run_draft_review_folder_blocked_by_a_file_is_failed_not_skipped(tmp_path, monkeypatch, caplog):
    """`review` existing as a regular file makes os.makedirs raise
    FileExistsError. No sheet exists, so it is a failure (exit 1, a re-run
    retries), never the "sheet appeared" skip."""
    _setup(tmp_path, monkeypatch, ROWS)
    _fake_drafts(monkeypatch, _stated_first, [])
    plan = tp.prepare_draft([])
    (tmp_path / "review").write_text("not a folder")
    with caplog.at_level("INFO"):
        counts = tp.run_draft(plan, CLIENT)
    assert counts == {"written": 0, "skipped": 0, "failed": 1}
    assert f"{MEMO}: failed — could not create the folder" in caplog.text


def test_run_draft_end_to_end_through_call_llm(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch, ROWS)
    prompts = []

    def fake_call_llm(prompt, llm_client, temperature=0.0, max_tokens=4096):
        prompts.append(prompt)
        return '{"verdict": "needs_combining", "needed": ["B"], "reason": "B over A."}'

    monkeypatch.setattr(gsp, "call_llm", fake_call_llm)
    counts = tp.run_draft(tp.prepare_draft([]), CLIENT)
    assert counts["written"] == 1
    assert len(prompts) == 1                                        # the no-pieces claim made no call
    assert "CLAIM: Acme sells widgets." in prompts[0]
    recs = _sheet_records(str(tmp_path / "review" / f"{MEMO}.xlsx"))
    assert [(r["block"], r["verdict_needed"]) for r in recs if r["claim_id"] == _cid(C1)] == \
        [(1, "needs combining"), ("A", None), ("B", "yes"), ("+", None), ("+", None)]


# --- Plan A: CLI ------------------------------------------------------------------

def _count_clients(monkeypatch, client=CLIENT):
    """Replace tp._bundle_client; the returned list grows by one per call."""
    built = []

    def bundle_client():
        built.append(1)
        return client

    monkeypatch.setattr(tp, "_bundle_client", bundle_client)
    return built


@pytest.mark.parametrize("argv, message", [
    (["tag_pipeline.py"], "usage: python tag_pipeline.py draft"),
    (["tag_pipeline.py", "build"], "unknown command 'build'"),
    (["tag_pipeline.py", "finalize"], "finalize takes exactly one memo_id"),
])
def test_main_rejects_commands_before_client(tmp_path, monkeypatch, argv, message):
    monkeypatch.chdir(tmp_path)
    built = _count_clients(monkeypatch)
    with pytest.raises(SystemExit, match=message):
        tp._main(argv)
    assert built == []


@pytest.mark.parametrize("case", ["bad memo id", "no filing entity", "non-uuid5 claim_id", "missing chunk_id"])
def test_main_draft_refusals_never_build_a_client(tmp_path, monkeypatch, case):
    """Spec §5.2 check 1 and §12: every refusal happens before LLMClient.from_env()."""
    rows, entity, argv = list(ROWS), "Acme Ltd", ["tag_pipeline.py", "draft"]
    if case == "bad memo id":
        argv.append("MEMO-NOPE")
    elif case == "no filing entity":
        entity = None
    elif case == "non-uuid5 claim_id":
        rows[1] = {**rows[1], "claim_id": str(uuid.uuid4())}
    else:
        rows[1] = {**rows[1], "chunk_id": "d.pdf_77"}
    _setup(tmp_path, monkeypatch, rows, entity=entity)
    built = _count_clients(monkeypatch)
    with pytest.raises(SystemExit, match="draft refused"):
        tp._main(argv)
    assert built == []


def test_main_draft_builds_bundle_model_client_after_checks(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch, ROWS)
    built = _count_clients(monkeypatch)
    seen = []
    monkeypatch.setattr(tp, "run_draft",
                        lambda plan, client: seen.append((plan, client)) or {"written": 1, "skipped": 0, "failed": 0})
    tp._main(["tag_pipeline.py", "draft", MEMO])
    assert built == [1]
    [(plan, client)] = seen
    assert [memo["memo_id"] for memo in plan] == [MEMO]
    assert client is CLIENT


def test_bundle_client_pins_the_model_and_needs_no_llm_model(monkeypatch):
    """.env.example says draft ignores LLM_MODEL: an environment without it works."""
    monkeypatch.setenv("LLM_BASE_URL", "http://endpoint.invalid")
    monkeypatch.setenv("LLM_API_KEY", "key")
    monkeypatch.delenv("LLM_MODEL", raising=False)
    client = tp._bundle_client()
    assert (client.base_url, client.api_key, client.model) == ("http://endpoint.invalid", "key", tp._BUNDLE_MODEL)
    monkeypatch.setenv("LLM_MODEL", "some/other-model")
    assert tp._bundle_client().model == tp._BUNDLE_MODEL


def test_bundle_client_names_only_the_missing_endpoint_variables(monkeypatch):
    monkeypatch.setenv("LLM_BASE_URL", "http://endpoint.invalid")
    monkeypatch.delenv("LLM_API_KEY", raising=False)
    monkeypatch.delenv("LLM_MODEL", raising=False)
    with pytest.raises(EnvironmentError) as excinfo:
        tp._bundle_client()
    assert "Missing required environment variable(s): LLM_API_KEY —" in str(excinfo.value)
    assert "LLM_BASE_URL," not in str(excinfo.value)


def test_main_draft_exits_1_when_a_memo_failed(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch, ROWS)
    _count_clients(monkeypatch)
    monkeypatch.setattr(tp, "run_draft", lambda plan, client: {"written": 0, "skipped": 0, "failed": 1})
    with pytest.raises(SystemExit) as excinfo:
        tp._main(["tag_pipeline.py", "draft"])
    assert excinfo.value.code == 1


def test_main_draft_all_skipped_needs_no_credentials(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch, ROWS)
    _existing_sheet(tmp_path)
    built = _count_clients(monkeypatch)
    seen = []
    monkeypatch.setattr(tp, "run_draft",
                        lambda plan, client: seen.append(client) or {"written": 0, "skipped": 1, "failed": 0})
    tp._main(["tag_pipeline.py", "draft"])
    assert built == [] and seen == [None]


def test_main_draft_memo_with_no_found_chunks_needs_no_credentials(tmp_path, monkeypatch):
    """Every claim found=False: draft_claim answers each without an LLM, so no
    client is built and the no-evidence sheet is still written."""
    _setup(tmp_path, monkeypatch, [_ckpt_row(C2), _ckpt_row(C1)])
    built = _count_clients(monkeypatch)
    monkeypatch.setattr(tp, "_call_llm_with_json_retry", lambda *a, **k: pytest.fail("no LLM call expected"))
    tp._main(["tag_pipeline.py", "draft", MEMO])
    assert built == []
    recs = _sheet_records(str(tmp_path / "review" / f"{MEMO}.xlsx"))
    assert [r["ai_verdict"] for r in recs if r["row_kind"] == "claim"] == ["not_supported", "not_supported"]
