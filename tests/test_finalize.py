"""
Tests for tag_pipeline.py's finalize command (spec
docs/superpowers/specs/2026-09-13-tag-pipeline-v2-bundle-design.md §7 as
refined by §13). No LLM and no real PDF, except one test that reads a real
filing when sources/sample.pdf is on disk (skipped otherwise). A review sheet
is made the way `draft` makes it — prepare_draft + write_review_sheet on a
stubbed memo — then edited cell by cell, as a reviewer would.
"""
import hashlib
import io
import os
import pathlib
import zipfile

import pandas as pd
import pytest
from openpyxl import load_workbook

import eval_pipeline as ep
import golden_set_pipeline as gsp
import tag_pipeline as tp
from golden_set_pipeline import LLMClient

MEMO = "MEMO-T"
SECTION = "Business Profile"
C1, C2, C3 = "Acme sells widgets.", "Acme has 12 plants.", "Acme is old."
CLAIMS_MD = f"""---
memo_id: {MEMO}
filing_entity: Acme Ltd
source_folder: sources/acme
---

## {SECTION}

1. {C1}
2. {C2}
3. {C3}
"""
# Distinct sentences, so a 40-character slice has exactly one home.
DOC_TEXT = " ".join(f"Sentence {i} reports that plant {i} made {i * 7} widgets." for i in range(110))
CHUNKS = {c["chunk_id"]: c for c in gsp.build_chunk_index([("d.pdf", DOC_TEXT)])}
SHEET = os.path.join("review", f"{MEMO}.xlsx")
REVIEWED = os.path.join("reviewed", f"{MEMO}.xlsx")
DATE = "2026-09-14"


def _cid(claim_text):
    return gsp._derive_claim_id(MEMO, SECTION, claim_text, 0)


def _ckpt_row(claim_text, chunk_id=None, **overrides):
    found = chunk_id is not None
    row = {"claim_id": _cid(claim_text), "memo_id": MEMO, "section": SECTION, "claim_text": claim_text,
           "doc_id": "d.pdf" if found else None, "chunk_id": chunk_id,
           "chunk_text": CHUNKS[chunk_id]["chunk_text"] if found else None,
           "bm25_score": 1.5 if found else None, "evidence_span": f"span of {chunk_id}" if found else None,
           "found": found, "confidence": "high" if found else None,
           "ambiguous_match": False, "verbatim_match": True if found else None,
           "human_reviewed": False, "tag": None}
    row.update(overrides)
    return row


ROWS = [_ckpt_row(C1, "d.pdf_2"), _ckpt_row(C1, "d.pdf_0"), _ckpt_row(C2),
        _ckpt_row(C3, "d.pdf_3", evidence_span="=SUM(A1) looks like a formula")]
DRAFTS = {
    C1: {"verdict": "needs_combining", "needed_chunk_ids": ["d.pdf_0"], "reason": "B over the total.", "attempts": 1},
    C2: {"verdict": "not_supported", "needed_chunk_ids": [], "reason": tp._AUTO_NOT_FOUND_REASON, "attempts": 0},
    C3: {"verdict": tp._DRAFT_FAILED, "needed_chunk_ids": [], "reason": "DRAFT FAILED: a | b", "attempts": 2},
}


@pytest.fixture
def memo(tmp_path, monkeypatch):
    """A drafted, unreviewed memo in tmp_path (the cwd). The sheet: claim 1
    (A = d.pdf_2, B = d.pdf_0 marked), claim 2 (no pieces), claim 3 (A =
    d.pdf_3, draft failed); each claim followed by two '+' rows."""
    monkeypatch.chdir(tmp_path)
    pd.DataFrame(ROWS).to_parquet("golden_set_checkpoint.parquet")
    (tmp_path / "claims").mkdir()
    (tmp_path / "claims" / f"{MEMO}.md").write_text(CLAIMS_MD)
    monkeypatch.setattr(tp, "_load_source_documents", lambda folder, label: [("d.pdf", DOC_TEXT)])
    claims = tp.prepare_draft([])[0]["claims"]
    (tmp_path / "review").mkdir()
    tp.write_review_sheet(SHEET, MEMO, [{**c, "draft": DRAFTS[c["claim_text"]]} for c in claims])
    return tmp_path


def _edit(changes, path=SHEET):
    """Set cells as a reviewer would. `changes` maps (claim number, place) to
    {column key: value}; place is "claim", a chunk letter ("A"), or "+1"/"+2"."""
    wb = load_workbook(path)
    ws = wb["bundles"]
    key_by_header = {header: key for key, header in tp._SHEET_HEADERS.items()}
    col = {key_by_header[cell.value]: i for i, cell in enumerate(ws[1]) if cell.value in key_by_header}
    number, plus = None, 0
    for row in ws.iter_rows(min_row=2):
        kind = row[col["row_kind"]].value
        if kind == "claim":
            number, plus, place = row[col["block"]].value, 0, "claim"
        elif kind == "chunk":
            place = row[col["block"]].value
        else:
            plus += 1
            place = f"+{plus}"
        for key, value in changes.get((number, place), {}).items():
            row[col[key]].value = value
    wb.save(path)


APPROVE = {(1, "claim"): {"checked": "yes"}, (2, "claim"): {"checked": "yes"},
           (3, "claim"): {"checked": "yes", "verdict_needed": "stated directly"}, (3, "A"): {"verdict_needed": "yes"}}


def _approve(extra=None):
    """Every claim CHECKED; claim 3's draft failed replaced by 'stated
    directly' on A. `extra` changes are merged cell by cell on top."""
    changes = {where: dict(values) for where, values in APPROVE.items()}
    for where, values in (extra or {}).items():
        changes.setdefault(where, {}).update(values)
    _edit(changes)


# --- read_review_sheet -------------------------------------------------------------

def test_read_review_sheet_returns_records_with_sheet_rows(memo):
    records = tp.read_review_sheet(SHEET)
    assert [(r["sheet_row"], r["row_kind"], r["block"]) for r in records] == [
        (2, "claim", 1), (3, "chunk", "A"), (4, "chunk", "B"), (5, "add", "+"), (6, "add", "+"),
        (7, "claim", 2), (8, "add", "+"), (9, "add", "+"),
        (10, "claim", 3), (11, "chunk", "A"), (12, "add", "+"), (13, "add", "+")]
    assert set(records[0]) == set(tp._SHEET_COLUMNS) | {"sheet_row"}
    assert (records[2]["chunk_id"], records[2]["verdict_needed"], records[2]["found"]) == ("d.pdf_0", "yes", True)


def test_read_review_sheet_survives_a_numbers_round_trip(memo):
    """What Apple Numbers did to the real MEMO-004 sheet (spec §13.1): an
    'Export Summary' sheet inserted first, hidden columns made visible, an
    extra column with an empty header — plus blank trailing rows."""
    before = tp.read_review_sheet(SHEET)
    wb = load_workbook(SHEET)
    wb.create_sheet("Export Summary", 0)["A1"] = "This document was exported from Numbers."
    ws = wb["bundles"]
    for dim in ws.column_dimensions.values():
        dim.hidden = False
    extra = ws.max_column + 1
    ws.cell(row=2, column=extra, value="stray")
    ws.cell(row=40, column=1, value="   ")
    wb.save(SHEET)
    assert load_workbook(SHEET).active.title == "Export Summary"      # a naive reader would read this
    assert tp.read_review_sheet(SHEET) == before


@pytest.mark.parametrize("damage, message", [
    ("rename tab", "no 'bundles' tab"),
    ("rename header", r"missing column header\(s\) \['CHECKED'\]"),
    ("repeat header", "'CHECKED' appears twice"),
    ("not a workbook", "not a readable .xlsx workbook"),
    ("corrupt sheet xml", "not a readable .xlsx workbook"),
])
def test_read_review_sheet_refusals(memo, damage, message):
    if damage == "not a workbook":
        pathlib.Path(SHEET).write_text("not a spreadsheet")
    elif damage == "corrupt sheet xml":
        # A valid zip whose worksheet XML is cut in half: openpyxl's XML parser
        # raises ParseError, which must be a refusal, not a traceback.
        original = pathlib.Path(SHEET).read_bytes()
        pathlib.Path(SHEET).unlink()
        with zipfile.ZipFile(io.BytesIO(original)) as zin, zipfile.ZipFile(SHEET, "w") as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename.startswith("xl/worksheets/"):
                    data = data[: len(data) // 2]
                zout.writestr(item, data)
    else:
        wb = load_workbook(SHEET)
        ws = wb["bundles"]
        checked = next(cell for cell in ws[1] if cell.value == "CHECKED")
        if damage == "rename tab":
            ws.title = "Bundles 1"
        elif damage == "rename header":
            checked.value = "checked?"
        else:
            ws.cell(row=1, column=ws.max_column + 1, value="CHECKED")
        wb.save(SHEET)
    with pytest.raises(tp.FinalizeRefused, match=message):
        tp.read_review_sheet(SHEET)


# --- _resolve_quote ------------------------------------------------------------------

SMALL = {
    "a.pdf_0": {"doc_id": "a.pdf", "chunk_text": "Revenue for FY24 was EUR4,203m, up\n3% on a like-for-like basis."},
    "a.pdf_1": {"doc_id": "a.pdf", "chunk_text": "up 3% on a like-for-like basis. Net debt was stable at EUR 19bn."},
    "a.pdf_5": {"doc_id": "a.pdf", "chunk_text": "Figures in EUR million unless stated otherwise. Offices: 12."},
    "b.pdf_0": {"doc_id": "b.pdf", "chunk_text": "Figures in EUR million unless stated otherwise. Retail: 40."},
    "b.pdf_1": {"doc_id": "b.pdf", "chunk_text": "The Group opened two centres in Spain during the year."},
}


def test_resolve_quote_too_short():
    assert tp._resolve_quote("  EUR 19bn   stable  ", SMALL) == "quote is 15 characters — paste at least 25"


def test_resolve_quote_refuses_the_quote_separator():
    # finalize joins one chunk's quotes with " | " and the eval splits there again.
    assert "separate quotes" in tp._resolve_quote("Revenue for FY24 | was EUR 4,203m", SMALL)


@pytest.mark.parametrize("quote", ["Revenue for FY24\n|\nwas EUR 4,203m",     # stored with the line breaks collapsed
                                   "Revenue for FY24 was EUR 4,203m |",        # joined: "… | | next"
                                   "| Revenue for FY24 was EUR 4,203m"])
def test_resolve_quote_refuses_the_separator_as_it_will_be_stored(quote):
    assert "separate quotes" in tp._resolve_quote(quote, SMALL)


def test_resolve_quote_ignores_all_whitespace_differences():
    """A viewer copy has a space pypdf dropped ('EUR 4,203m') and no line break."""
    assert tp._resolve_quote("Revenue for FY24 was EUR 4,203m", SMALL) == ["a.pdf_0"]
    assert tp._resolve_quote("Net debt was\n\n stable at EUR 19bn", SMALL) == ["a.pdf_1"]


def test_resolve_quote_not_found_names_both_likely_causes():
    problem = tp._resolve_quote("The Group opened three centres in Spain", SMALL)
    assert problem.startswith("quote not found") and "ligatures, hyphenation" in problem


def test_resolve_quote_in_the_shared_overlap_of_neighbours_returns_both():
    assert tp._resolve_quote("up 3% on a like-for-like basis", SMALL) == ["a.pdf_0", "a.pdf_1"]
    reordered = {"a.pdf_1": SMALL["a.pdf_1"], "a.pdf_0": SMALL["a.pdf_0"]}
    assert tp._resolve_quote("up 3% on a like-for-like basis", reordered) == ["a.pdf_0", "a.pdf_1"]


@pytest.mark.parametrize("chunks", [
    {k: SMALL[k] for k in ("a.pdf_5", "b.pdf_0")},                                       # two documents
    {"a.pdf_0": {**SMALL["a.pdf_5"], "doc_id": "a.pdf"},
     "a.pdf_2": {**SMALL["b.pdf_0"], "doc_id": "a.pdf"}},                                # same doc, not neighbours
])
def test_resolve_quote_a_generic_phrase_in_two_places_is_refused(chunks):
    problem = tp._resolve_quote("Figures in EUR million unless stated otherwise", chunks)
    assert problem.startswith("quote found in 2 places") and problem.endswith("paste a longer quote")


# Any small real filing: sources/ is gitignored, so point this path (a symlink
# is fine) at one locally to run the test.
SAMPLE_PDF = pathlib.Path(__file__).resolve().parent.parent / "sources" / "sample.pdf"


@pytest.mark.skipif(not SAMPLE_PDF.exists(), reason="real filing not on disk (sources/ is gitignored)")
def test_resolve_quote_on_real_pdf_text_across_line_breaks():
    """Real pypdf text breaks lines mid-sentence. A quote typed or copied with
    a space where the line broke must still resolve to its chunk. Uses the
    sample filing (~1 s for a short one). No PDF text is stored in the repo."""
    chunks = {c["chunk_id"]: c for c in gsp.build_chunk_index([(SAMPLE_PDF.name, gsp.load_pdf_text(str(SAMPLE_PDF)))])}
    resolved = 0
    for chunk_id, chunk in chunks.items():
        body = chunk["chunk_text"][250:750]                   # clear of both 200-character overlaps
        for at in [i for i, ch in enumerate(body) if ch == "\n"][:5]:
            window = body[max(0, at - 25):at + 25]
            if len(" ".join(window.split())) < tp._MIN_QUOTE_CHARS:
                continue
            result = tp._resolve_quote(window.replace("\n", " "), chunks)
            if isinstance(result, list):                     # a repeated table line may be refused; fine
                assert chunk_id in result
                resolved += 1
    assert resolved >= 5


# --- check_review --------------------------------------------------------------------

def _check():
    _, file_claims = tp.read_claims_file(MEMO)
    return tp.check_review(MEMO, tp.read_review_sheet(SHEET), file_claims, CHUNKS, {"d.pdf": DOC_TEXT})


def _only_in(chunk_id):
    """A 40-character slice of DOC_TEXT that sits in `chunk_id` and no other chunk."""
    start = CHUNKS[chunk_id]["start_offset"]
    quote = DOC_TEXT[start + 300:start + 340]
    assert [c for c, chunk in CHUNKS.items() if quote in chunk["chunk_text"]] == [chunk_id]
    return quote


def test_check_review_untouched_draft_lists_every_problem_at_once(memo):
    problems, _ = _check()
    assert problems == [
        "row 2 (claim 1): CHECKED is not 'yes'",
        "row 7 (claim 2): CHECKED is not 'yes'",
        "row 10 (claim 3): CHECKED is not 'yes'",
        "row 10 (claim 3): VERDICT is still 'draft failed' — pick a verdict yourself and mark the pieces "
        "that support it",
    ]


def test_check_review_approved_sheet_has_no_problems(memo):
    _approve()
    problems, claims = _check()
    assert problems == []
    assert [(c["number"], c["verdict"], [ch["marked"] for ch in c["chunks"]]) for c in claims] == [
        (1, "needs_combining", [False, True]), (2, "not_supported", []), (3, "stated_directly", [True])]


def test_check_review_compares_yes_and_verdicts_stripped_and_case_folded(memo):
    _approve({(1, "claim"): {"checked": " Yes", "verdict_needed": "Needs Combining "},
              (1, "B"): {"verdict_needed": "YES"}})
    assert _check()[0] == []


@pytest.mark.parametrize("changes, expected", [
    ({(1, "claim"): {"verdict_needed": "maybe"}}, "row 2 (claim 1): VERDICT 'maybe' is not one of"),
    ({(1, "claim"): {"verdict_needed": "not supported"}}, "row 2 (claim 1): 'not supported' with pieces marked"),
    ({(1, "B"): {"verdict_needed": None}}, "row 2 (claim 1): 'needs combining' with no piece marked"),
    ({(1, "A"): {"verdict_needed": "no"}}, "row 3 (claim 1): NEEDED? 'no' is not 'yes' or blank"),
    ({(1, "+1"): {"text": "Sentence 3 reports that plant 3 made 21"}}, "row 5 (claim 1): a quote without NEEDED? 'yes'"),
    ({(1, "+2"): {"verdict_needed": "yes"}}, "row 6 (claim 1): NEEDED? 'yes' on a '+' row without a quote"),
    ({(1, "A"): {"row_kind": "piece"}}, "row 3: unknown row_kind 'piece'"),
    ({(1, "B"): {"claim_id": _cid(C3)}}, "row 4 (claim 1): claim_id " + repr(_cid(C3)) + " differs from its claim row's"),
    ({(3, "A"): {"memo_id": "MEMO-OTHER"}}, "row 11 (claim 3): memo_id 'MEMO-OTHER' is not MEMO-T"),
    ({(1, "A"): {"chunk_id": "d.pdf_9"}}, "row 3 (claim 1): chunk_id 'd.pdf_9' is not in the chunk index"),
    ({(1, "A"): {"chunk_id": "d.pdf_1"}}, "row 3 (claim 1): chunk 'd.pdf_1': its text differs from the PDFs'"),
    ({(1, "A"): {"chunk_id": "d.pdf_0"}}, "row 4 (claim 1): chunk 'd.pdf_0' appears twice under this claim"),
    ({(1, "A"): {"doc_id": "e.pdf"}}, "row 3 (claim 1): chunk 'd.pdf_2': document 'e.pdf' differs"),
    # all three hidden cells replaced with another real chunk; the passage shown stays d.pdf_2's
    ({(1, "A"): {"chunk_id": "d.pdf_3", "chunk_text": CHUNKS["d.pdf_3"]["chunk_text"]}},
     "row 3 (claim 1): chunk 'd.pdf_3': the passage shown in 'text' is not this chunk's text (hidden cells "
     "replaced with another chunk's?)"),
    ({(1, "A"): {"document": "e.pdf"}}, "row 3 (claim 1): chunk 'd.pdf_2': the 'document' shown is 'e.pdf', "
                                        "not 'd.pdf' (an edited cell?)"),
    ({(1, "A"): {"found": "TRUE"}}, "row 3 (claim 1): found 'TRUE' is not TRUE/FALSE"),
    ({(1, "A"): {"found": False}}, "row 3 (claim 1): found is FALSE, but draft writes chunk rows only for "
                                   "found chunks (an edited cell?)"),
    ({(1, "A"): {"bm25_score": "high"}}, "row 3 (claim 1): bm25_score 'high' is not a number"),
])
def test_check_review_problem(memo, changes, expected):
    _approve(changes)
    problems, _ = _check()
    assert any(p.startswith(expected) for p in problems), problems


def _delete_row(sheet_row):
    """Delete one worksheet row of the bundles tab, as a reviewer might."""
    wb = load_workbook(SHEET)
    wb["bundles"].delete_rows(sheet_row)
    wb.save(SHEET)


def test_check_review_deleted_chunk_row_leaves_a_gap_in_the_letters(memo):
    """Claim 1's unmarked row A deleted: B is still marked, so every other
    check passes — without this refusal, d.pdf_2 would silently vanish from
    the golden set (CLAUDE.md decision 3: a deleted row is never silent)."""
    _approve()
    _delete_row(3)
    problems, _ = _check()
    assert problems == [
        "row 3 (claim 1): chunk rows are lettered ['B'], not ['A'] as draft wrote them — a chunk row was "
        "deleted or moved: restore it from the drafted sheet, or delete review/MEMO-T.xlsx and re-run "
        "`python tag_pipeline.py draft MEMO-T`"]


def test_check_review_deleted_only_chunk_row_of_a_claim_search_found_evidence_for(memo):
    """Claim 3's only chunk row deleted and the claim approved 'not supported':
    without this refusal it would be written found=False with the 'auto:
    found=False, nothing was found by search' rationale — false, search found
    d.pdf_3. Claim 2 (really nothing found) has no chunk rows and passes."""
    _approve({(3, "claim"): {"verdict_needed": "not supported"}, (3, "A"): {"verdict_needed": None}})
    _delete_row(11)
    problems, _ = _check()
    assert problems == [
        f"row 10 (claim 3): no chunk rows, but its AI reason is not '{tp._AUTO_NOT_FOUND_REASON}', so "
        "search found chunks for it — a chunk row was deleted: restore it from the drafted sheet, or "
        "delete review/MEMO-T.xlsx and re-run `python tag_pipeline.py draft MEMO-T`"]


def test_check_review_deleted_last_chunk_row_is_caught_by_the_drafted_count(memo):
    """Claim 1's LAST row (B) deleted, with A marked instead: the letters read
    ['A'] and every other check passes, so only the claim row's hidden
    chunk_count (2, written by draft) shows that d.pdf_0 was dropped."""
    _approve({(1, "A"): {"verdict_needed": "yes"}})
    _delete_row(4)
    problems, _ = _check()
    assert problems == [
        "row 2 (claim 1): 1 chunk row(s) under this claim, but draft wrote 2 — a chunk row was deleted: "
        "restore it from the drafted sheet, or delete review/MEMO-T.xlsx and re-run "
        "`python tag_pipeline.py draft MEMO-T`"]


@pytest.mark.parametrize("value", ["two", None, 2.5, True])
def test_check_review_edited_chunk_count_is_refused(memo, value):
    _approve({(1, "claim"): {"chunk_count": value}})
    problems, _ = _check()
    assert problems == [f"row 2 (claim 1): chunk_count {value!r} is not a whole number (an edited cell?)"]


def test_check_review_accepts_a_whole_number_chunk_count_saved_as_a_float(memo):
    """A spreadsheet app may save 2 back as 2.0."""
    _approve({(1, "claim"): {"chunk_count": 2.0}})
    assert _check()[0] == []


@pytest.mark.parametrize("now", [
    {"d.pdf": DOC_TEXT, "e.pdf": " ".join(f"Line {i} of another filing about something else." for i in range(60))},
    {"d.pdf": DOC_TEXT, "e.pdf": ""},                                    # added PDF with no extractable text
    {"d.pdf": DOC_TEXT.replace("Sentence 105 ", "Sentence one-oh-five ")},  # same name, text edited in place
], ids=["pdf added", "empty-text pdf added", "same name, text edited"])
def test_check_review_pdfs_changed_since_drafting_is_one_problem(memo, now):
    """The memo's PDFs changed after drafting — one added (even one with no
    text, which yields no chunks), or one edited in place under its old name
    outside every drafted chunk. Every chunk row still verifies, but a claim
    drafted with nothing found was never searched in the new text, so its
    'auto: found=False' would be false. Refused once, not per claim row."""
    _approve()
    _, file_claims = tp.read_claims_file(MEMO)
    chunks = {c["chunk_id"]: c for c in gsp.build_chunk_index(list(now.items()))}
    problems, _ = tp.check_review(MEMO, tp.read_review_sheet(SHEET), file_claims, chunks, now)
    drafted, current = tp._source_docs({"d.pdf": DOC_TEXT}), tp._source_docs(now)
    assert drafted != current
    assert problems == [
        f"row 2: this sheet was drafted from the PDFs {drafted!r}, but claims/MEMO-T.md's source_folder now "
        f"gives {current!r} (names and extracted text are compared) — run `python golden_set_pipeline.py build`, "
        f"then `python tag_pipeline.py draft MEMO-T`, and review again (answers are not carried over)"]


def test_run_finalize_sheet_drafted_before_source_docs_warns_and_still_finalizes(memo, caplog):
    """MEMO-004's sheet predates the source_docs column: still finalized, and
    the check that cannot run is named in a warning."""
    _approve()
    wb = load_workbook(SHEET)
    ws = wb["bundles"]
    ws.delete_cols(next(cell.column for cell in ws[1] if cell.value == "source_docs"))
    wb.save(SHEET)
    caplog.set_level("WARNING", logger="tag")
    tp.run_finalize(MEMO, run_date=DATE)
    assert os.path.exists(REVIEWED)
    assert any("has no source_docs column" in r.getMessage() for r in caplog.records), caplog.text


def test_run_finalize_sheet_drafted_before_chunk_count_warns_and_still_finalizes(memo, caplog):
    """A sheet drafted before the chunk_count column existed (MEMO-004's) has no
    such column: it is still read and finalized, and the one check that
    cannot run is named in a warning."""
    _approve()
    wb = load_workbook(SHEET)
    ws = wb["bundles"]
    ws.delete_cols(next(cell.column for cell in ws[1] if cell.value == "chunk_count"))
    wb.save(SHEET)
    assert "chunk_count" not in tp.read_review_sheet(SHEET)[0]
    caplog.set_level("WARNING", logger="tag")
    tp.run_finalize(MEMO, run_date=DATE)
    assert os.path.exists(REVIEWED)
    assert any("has no chunk_count column" in r.getMessage() for r in caplog.records), caplog.text


def test_check_review_edited_claim_id_names_both_causes_and_the_recovery(memo):
    _approve({(2, "claim"): {"claim_id": _cid("Acme has 13 plants.")}})
    problems, _ = _check()
    assert problems[0].startswith(f"row 7 (claim 2): claim_id {_cid('Acme has 13 plants.')} is not in "
                                  f"claims/MEMO-T.md (an edited cell, or a claim reworded after drafting)")
    assert any(p.startswith(f"claims/MEMO-T.md claim 'Acme has 12 plants.' ({_cid(C2)}) is not in the sheet")
               for p in problems)
    assert problems[-1].startswith("to recover from a changed claims file: delete review/MEMO-T.xlsx")


_SWAPPED = "for this claim_id (an edited cell, or claim ids swapped between claims)"


def test_check_review_edited_claim_text_cell_is_refused(memo):
    """The claim row's text is compared with the claims file's text for its
    claim_id: an edit to it would otherwise be silently ignored."""
    _approve({(1, "claim"): {"text": "Acme sells gadgets."}})
    problems, _ = _check()
    assert problems == [f"row 2 (claim 1): text 'Acme sells gadgets.' is not claims/MEMO-T.md's "
                        f"'Acme sells widgets.' {_SWAPPED}"]


def test_check_review_claim_ids_swapped_between_two_blocks_are_refused(memo):
    """Every row of claim 1 given claim 2's id and vice versa: both ids are
    still in the claims file and the census holds, but each block's verdict
    and marks would land on the other claim."""
    swap = {(1, place): {"claim_id": _cid(C2)} for place in ("claim", "A", "B", "+1", "+2")}
    swap.update({(2, place): {"claim_id": _cid(C1)} for place in ("claim", "+1", "+2")})
    _approve(swap)
    problems, _ = _check()
    assert problems == [
        f"row 2 (claim 1): text 'Acme sells widgets.' is not claims/MEMO-T.md's 'Acme has 12 plants.' {_SWAPPED}",
        f"row 7 (claim 2): text 'Acme has 12 plants.' is not claims/MEMO-T.md's 'Acme sells widgets.' {_SWAPPED}"]


def test_check_review_claim_added_to_claims_file_after_drafting(memo):
    """Spec §13.3 check 5, second direction (§7.2 listed only the first)."""
    _approve()
    (memo / "claims" / f"{MEMO}.md").write_text(CLAIMS_MD + "4. Acme is listed.\n")
    problems, _ = _check()
    assert problems == [
        f"claims/MEMO-T.md claim 'Acme is listed.' ({_cid('Acme is listed.')}) is not in the sheet "
        f"(added to the claims file after drafting?)",
        "to recover from a changed claims file: delete review/MEMO-T.xlsx, re-run `python "
        "golden_set_pipeline.py build`, then `python tag_pipeline.py draft MEMO-T` and review again "
        "(answers are not carried over)",
    ]


def test_check_review_quote_marks_a_bundle_chunk_or_adds_a_new_one(memo):
    _approve({(1, "+1"): {"text": _only_in("d.pdf_2"), "verdict_needed": "yes"},
              (1, "+2"): {"text": _only_in("d.pdf_4"), "verdict_needed": "yes", "note": "from the table"}})
    problems, claims = _check()
    assert problems == []
    a, b = claims[0]["chunks"]
    assert (a["marked"], a["by_quote"], b["by_quote"]) == (False, "1 chunk", None)
    assert claims[0]["added"] == [{"chunk_id": "d.pdf_4", "quotes": [_only_in("d.pdf_4")],
                                   "notes": ["from the table"], "how": "1 chunk"}]


def test_check_review_two_quotes_for_one_new_chunk_become_one_entry(memo):
    start = CHUNKS["d.pdf_4"]["start_offset"]
    q1, q2 = DOC_TEXT[start + 300:start + 340], DOC_TEXT[start + 500:start + 540]
    _approve({(1, "+1"): {"text": q1, "verdict_needed": "yes", "note": "n1"},
              (1, "+2"): {"text": q2, "verdict_needed": "yes", "note": "n2"}})
    problems, claims = _check()
    assert problems == []
    assert claims[0]["added"] == [{"chunk_id": "d.pdf_4", "quotes": [q1, q2], "notes": ["n1", "n2"], "how": "1 chunk"}]


def test_check_review_quote_in_neighbour_overlap_is_labelled(memo):
    start = CHUNKS["d.pdf_5"]["start_offset"]
    overlap = DOC_TEXT[start + 50:start + 90]                  # inside d.pdf_4's last 200 characters
    _approve({(1, "+1"): {"text": overlap, "verdict_needed": "yes"}})
    problems, claims = _check()
    assert problems == []
    assert [(a["chunk_id"], a["how"]) for a in claims[0]["added"]] == [
        ("d.pdf_4", "2 adjacent chunks — check for repeated text"),
        ("d.pdf_5", "2 adjacent chunks — check for repeated text")]


def test_check_review_quote_problems_are_listed_with_the_others(memo):
    _edit({(1, "+1"): {"text": "too short", "verdict_needed": "yes"}})
    problems, _ = _check()
    assert "row 5 (claim 1): quote is 9 characters — paste at least 25" in problems
    assert "row 2 (claim 1): CHECKED is not 'yes'" in problems


def test_check_review_not_supported_with_a_quote_is_refused_even_if_it_resolves(memo):
    _approve({(2, "+1"): {"text": _only_in("d.pdf_5"), "verdict_needed": "yes"}})
    assert "row 7 (claim 2): 'not supported' with pieces marked — clear the marks or change the verdict" in _check()[0]


# --- build_reviewed_rows + run_finalize ----------------------------------------------

def _reviewed():
    return pd.read_excel(REVIEWED)


def test_run_finalize_writes_rows_tags_and_rationale(memo):
    _approve({(1, "claim"): {"note": "claim note"}, (1, "B"): {"note": "row note"}})
    result = tp.run_finalize(MEMO, run_date=DATE)
    assert result["problems"] == [] and result["written"] == REVIEWED
    df = _reviewed()
    assert list(df.columns) == list(tp._REVIEWED_COLUMNS)
    got = [(r.claim_text, r.chunk_id, r.found, r.tag, r.tag_draft, r.tag_rationale)
           for r in df.astype(object).where(df.notna(), None).itertuples()]
    assert got == [
        (C1, "d.pdf_2", True, "unverifiable", "unverifiable",
         f"bundle review {DATE}: verdict 'needs combining'; chunk not marked; AI: B over the total.; note: claim note"),
        (C1, "d.pdf_0", True, "synthesized", "synthesized",
         f"bundle review {DATE}: verdict 'needs combining'; chunk marked; AI: B over the total.; "
         f"note: claim note; note: row note"),
        (C2, None, False, "unverifiable", "unverifiable", "auto: found=False, nothing was found by search"),
        (C3, "d.pdf_3", True, "extractive", None,
         f"bundle review {DATE}: verdict 'stated directly'; chunk marked; AI: DRAFT FAILED: a | b"),
    ]
    assert df["human_reviewed"].eq(True).all()
    assert df.loc[3, "evidence_span"] == "=SUM(A1) looks like a formula"       # text, not a formula
    assert df.loc[0, "chunk_text"] == CHUNKS["d.pdf_2"]["chunk_text"]
    assert df.loc[0, "claim_id"] == _cid(C1) and df.loc[0, "bm25_score"] == 1.5
    assert (result["verdicts_changed"], result["marks_changed"]) == (1, 1)


def test_run_finalize_human_added_rows_and_quote_marks(memo, caplog):
    _approve({(1, "+1"): {"text": _only_in("d.pdf_2"), "verdict_needed": "yes"},
              (2, "claim"): {"verdict_needed": "stated directly", "note": "found it"},
              (2, "+1"): {"text": _only_in("d.pdf_5"), "verdict_needed": "yes", "note": "page 4"}})
    with caplog.at_level("INFO", logger="tag"):
        result = tp.run_finalize(MEMO, run_date=DATE)
    assert (result["marked_by_quote"], result["human_added"]) == (1, 1)
    df = _reviewed().astype(object)
    a = df[df["chunk_id"].eq("d.pdf_2")].iloc[0]
    assert (a["tag"], a["tag_draft"]) == ("synthesized", "unverifiable")
    assert "chunk marked by quote" in a["tag_rationale"]
    c2 = df[df["claim_text"].eq(C2)]
    assert len(c2) == 1                                                   # placeholder dropped
    row = c2.iloc[0]
    assert (row["chunk_id"], row["found"], row["tag"], row["evidence_span"]) == \
        ("d.pdf_5", True, "extractive", " ".join(_only_in("d.pdf_5").split()))
    assert pd.isna(row["tag_draft"]) and pd.isna(row["bm25_score"]) and pd.isna(row["confidence"])
    assert row["tag_rationale"] == (f"human-added {DATE}: quote '{' '.join(_only_in('d.pdf_5').split())}'; "
                                    f"verdict 'stated directly'; note: found it; note: page 4")
    assert "claim 1: chunk d.pdf_2 marked by quote (1 chunk)" in caplog.text
    assert f"claim 2: human-added chunk d.pdf_5 (1 chunk): {' '.join(CHUNKS['d.pdf_5']['chunk_text'].split())[:80]!r}" \
        in caplog.text


def test_run_finalize_keeps_every_note_typed_on_a_plus_row(memo):
    """Spec §7.4: nothing typed in a note cell is dropped — also the note of a
    '+' row whose quote marks a chunk already in the bundle, and a '+' row
    holding only a note (a comment on the claim, kept on every row of it)."""
    _approve({(1, "+1"): {"text": _only_in("d.pdf_2"), "verdict_needed": "yes", "note": "quote note"},
              (1, "+2"): {"note": "about claim 1"},
              (2, "+1"): {"note": "about claim 2"}})
    assert tp.run_finalize(MEMO, run_date=DATE)["problems"] == []
    df = _reviewed()
    rationale = dict(zip(df["chunk_id"].fillna(""), df["tag_rationale"]))
    assert rationale["d.pdf_2"] == (f"bundle review {DATE}: verdict 'needs combining'; chunk marked by quote; "
                                    f"AI: B over the total.; note: about claim 1; note: quote note")
    assert rationale["d.pdf_0"] == (f"bundle review {DATE}: verdict 'needs combining'; chunk marked; "
                                    f"AI: B over the total.; note: about claim 1")
    assert rationale[""] == "auto: found=False, nothing was found by search; note: about claim 2"


def test_run_finalize_with_problems_writes_nothing_and_logs_each(memo, caplog):
    with caplog.at_level("INFO", logger="tag"):
        result = tp.run_finalize(MEMO, run_date=DATE)
    assert len(result["problems"]) == 4 and result["written"] is None
    assert not os.path.exists("reviewed")
    assert "MEMO-T: row 10 (claim 3): VERDICT is still 'draft failed'" in caplog.text
    assert "MEMO-T: refused — 4 problem(s); nothing written" in caplog.text


def test_run_finalize_refuses_a_reviewed_path_that_is_not_a_folder(memo, monkeypatch):
    """os.makedirs would raise FileExistsError at write time, a traceback after
    the PDFs were read; refused first instead, before the sheet is read."""
    _approve()
    pathlib.Path("reviewed").write_text("not a folder")
    monkeypatch.setattr(tp, "read_review_sheet", lambda path: pytest.fail("must refuse before reading"))
    with pytest.raises(tp.FinalizeRefused, match=r"^reviewed exists but is not a folder"):
        tp.run_finalize(MEMO, run_date=DATE)


def test_run_finalize_refuses_a_reviewed_file_path_that_is_a_folder(memo, monkeypatch):
    """os.replace onto a folder raises IsADirectoryError, and only after the
    PDFs were read; refused first instead, before the sheet is read."""
    _approve()
    os.makedirs(REVIEWED)
    monkeypatch.setattr(tp, "read_review_sheet", lambda path: pytest.fail("must refuse before reading"))
    with pytest.raises(tp.FinalizeRefused, match=r"^reviewed.MEMO-T\.xlsx exists but is not a file"):
        tp.run_finalize(MEMO, run_date=DATE)


@pytest.mark.parametrize("where, message", [
    ("reviewed", r"^reviewed exists but is not a folder"),
    (REVIEWED, r"^reviewed.MEMO-T\.xlsx exists but is not a file"),
])
def test_run_finalize_refuses_a_dangling_symlink_destination(memo, monkeypatch, where, message):
    """A symlink to nothing (os.path.exists is False for it, os.path.lexists
    True) is refused by run_finalize before the sheet or the PDFs are read —
    rather than failing with a raw OSError at write time."""
    _approve()
    os.makedirs(os.path.dirname(where) or ".", exist_ok=True)
    os.symlink("missing-target", where)
    monkeypatch.setattr(tp, "read_review_sheet", lambda path: pytest.fail("must refuse before reading"))
    with pytest.raises(tp.FinalizeRefused, match=message):
        tp.run_finalize(MEMO, run_date=DATE)


def test_run_finalize_overwrites_reviewed_and_never_touches_review(memo):
    _approve()
    tp.run_finalize(MEMO, run_date=DATE)
    _edit({(1, "claim"): {"note": "second pass"}})
    sheet_bytes = hashlib.sha256(pathlib.Path(SHEET).read_bytes()).hexdigest()
    tp.run_finalize(MEMO, run_date=DATE)
    assert hashlib.sha256(pathlib.Path(SHEET).read_bytes()).hexdigest() == sheet_bytes
    assert _reviewed().loc[0, "tag_rationale"].endswith("; note: second pass")
    assert sorted(os.listdir("reviewed")) == [f"{MEMO}.xlsx"]              # no temp file left


@pytest.mark.parametrize("setup, message", [
    ("numbered copy", r"\['MEMO-T 2.xlsx'\] next to review.MEMO-T.xlsx: Excel or Numbers saved a copy"),
    # Only the copy is left: the answers may be in it, so never "run draft" (≈ paid calls).
    ("numbered copy, no sheet", r"\['MEMO-T 2.xlsx'\] next to review.MEMO-T.xlsx: Excel or Numbers saved a copy"),
    ("no sheet", r"review.MEMO-T.xlsx not found — run `python tag_pipeline.py draft MEMO-T` first"),
    ("no review folder", r"review.MEMO-T.xlsx not found — run `python tag_pipeline.py draft MEMO-T` first"),
    ("no claims file", r"MEMO-T: claims.MEMO-T.md not found"),
])
def test_run_finalize_refusals(memo, setup, message):
    if setup.startswith("numbered copy"):
        pathlib.Path("review", f"{MEMO} 2.xlsx").write_bytes(pathlib.Path(SHEET).read_bytes())
        if setup.endswith("no sheet"):
            os.remove(SHEET)
    elif setup == "no review folder":
        os.remove(SHEET)
        os.rmdir("review")
    elif setup == "no sheet":
        os.remove(SHEET)
    else:
        os.remove(os.path.join("claims", f"{MEMO}.md"))
    with pytest.raises(tp.FinalizeRefused, match=message):
        tp.run_finalize(MEMO, run_date=DATE)


def test_write_reviewed_strips_characters_excel_cannot_store(tmp_path):
    path = str(tmp_path / "reviewed" / "m.xlsx")
    df = pd.DataFrame([{col: None for col in tp._REVIEWED_COLUMNS}])
    df.loc[0, ["chunk_text", "tag", "tag_rationale"]] = ["pypdf\x00 text", "extractive", "=1+1; note: x"]
    tp._write_reviewed(df, path)
    back = pd.read_excel(path)
    assert (back.loc[0, "chunk_text"], back.loc[0, "tag_rationale"]) == ("pypdf text", "=1+1; note: x")


def test_write_reviewed_failed_write_leaves_no_file(tmp_path, monkeypatch):
    path = str(tmp_path / "m.xlsx")

    def half_write(df, target):
        pathlib.Path(target).write_text("partial")
        raise OSError("disk full")

    monkeypatch.setattr(tp, "export_for_review", half_write)
    with pytest.raises(OSError, match="disk full"):
        tp._write_reviewed(pd.DataFrame(columns=list(tp._REVIEWED_COLUMNS)), path)
    assert os.listdir(tmp_path) == []


def test_write_reviewed_uses_a_unique_temp_file_per_call(tmp_path, monkeypatch):
    """Two finalize runs of one memo at once must not share a temp file: with a
    shared name one run could publish, or delete, the other's half-written file."""
    targets = []
    real_export = tp.export_for_review
    monkeypatch.setattr(tp, "export_for_review", lambda df, target: targets.append(target) or real_export(df, target))
    path = str(tmp_path / "m.xlsx")
    df = pd.DataFrame(columns=list(tp._REVIEWED_COLUMNS))
    tp._write_reviewed(df, path)
    tp._write_reviewed(df, path)
    assert len(set(targets)) == 2
    assert all(os.path.dirname(t) == str(tmp_path) and os.path.basename(t).startswith("m.")
               and t.endswith(".tmp.xlsx") for t in targets)
    assert os.listdir(tmp_path) == ["m.xlsx"]


def test_reviewed_file_meets_the_eval_ground_truth_contract(memo):
    """
    The eval spec's ground-truth rules (2026-09-10-retrieval-eval-harness-
    design.md, "Ground truth" and "Tags"), asserted on the file finalize
    writes, read back the way the eval reads it. The tag and census rules are
    now asserted by eval_pipeline's own validator.

    Blank means pd.isna(), never `is None` or == "": an empty xlsx cell reads
    back as NaN, so a None/"" check would pass here on a DataFrame and miss a
    blank tag in the real file.
    """
    _approve({(1, "+1"): {"text": _only_in("d.pdf_4"), "verdict_needed": "yes"}})
    tp.run_finalize(MEMO, run_date=DATE)
    df = pd.read_excel(REVIEWED)
    assert set(df.columns) == {"claim_id", "memo_id", "section", "claim_text", "doc_id", "chunk_id", "chunk_text",
                               "bm25_score", "evidence_span", "found", "confidence", "ambiguous_match",
                               "verbatim_match", "human_reviewed", "tag", "tag_draft", "tag_rationale"}
    gt = ep.load_ground_truth(MEMO)                             # raises on any contract breach
    assert set(df["claim_id"]) == set(gt.claims["claim_id"])    # every claim has a row, no row outside
    assert not df.duplicated(subset=["claim_id", "chunk_id"]).any()
    not_found = df[df["found"].eq(False)]
    assert len(not_found) == 1
    assert not_found["tag_rationale"].str.startswith("auto: found=False").all()


# --- CLI -------------------------------------------------------------------------------

def _count_clients(monkeypatch):
    """Replace tp._bundle_client; the returned list grows by one per call."""
    built = []

    def bundle_client():
        built.append(1)
        return LLMClient(base_url="http://unused.invalid", api_key="unused", model="unused")

    monkeypatch.setattr(tp, "_bundle_client", bundle_client)
    return built


@pytest.mark.parametrize("argv", [["tag_pipeline.py", "finalize"], ["tag_pipeline.py", "finalize", MEMO, "MEMO-U"]])
def test_main_finalize_takes_exactly_one_memo_id(tmp_path, monkeypatch, argv):
    monkeypatch.chdir(tmp_path)
    built = _count_clients(monkeypatch)
    with pytest.raises(SystemExit, match="finalize takes exactly one memo_id"):
        tp._main(argv)
    assert built == []


def test_main_finalize_never_builds_a_client(memo, monkeypatch):
    built = _count_clients(monkeypatch)
    with pytest.raises(SystemExit) as excinfo:
        tp._main(["tag_pipeline.py", "finalize", MEMO])               # untouched sheet: problems
    assert excinfo.value.code == 1
    _approve()
    tp._main(["tag_pipeline.py", "finalize", MEMO])                   # returns normally: exit 0
    assert os.path.exists(REVIEWED)
    assert built == []


def test_main_finalize_refusal_becomes_the_exit_message(memo, monkeypatch):
    os.remove(SHEET)
    with pytest.raises(SystemExit, match="finalize refused: review.MEMO-T.xlsx not found"):
        tp._main(["tag_pipeline.py", "finalize", MEMO])


@pytest.mark.parametrize("memo_id", [f"../{MEMO}", f"/tmp/{MEMO}", "MEMO T", ""])
def test_run_finalize_refuses_a_memo_id_that_is_not_a_plain_filename(memo, memo_id, monkeypatch):
    """The id names review/, claims/ and reviewed/ files: a path in it could read
    or overwrite files outside those folders. Refused before any file is read."""
    monkeypatch.setattr(tp, "read_review_sheet", lambda path: pytest.fail("must refuse before reading"))
    with pytest.raises(tp.FinalizeRefused, match=r"is not a valid memo id \(letters, digits"):
        tp.run_finalize(memo_id)
    with pytest.raises(SystemExit, match="finalize refused: .* is not a valid memo id"):
        tp._main(["tag_pipeline.py", "finalize", memo_id])
