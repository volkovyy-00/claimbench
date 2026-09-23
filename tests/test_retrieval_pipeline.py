"""
Tests for retrieval_pipeline.py — the standalone dense-retrieval prototype.
Mocked: rp.call_embeddings and gsp.load_pdf_text are stubbed; no live API,
no real PDF parsing (CLAUDE.md, "Testing convention").
"""
import hashlib
import os
import textwrap

import numpy as np
import pandas as pd
import pytest
import yaml

import golden_set_pipeline as gsp
import retrieval_pipeline as rp


def test_imports_from_golden_set_pipeline_resolve():
    # A rename in golden_set_pipeline must fail loudly here, not at first run.
    from golden_set_pipeline import (  # noqa: F401
        chunk_document,
        build_chunk_index,
        _load_source_documents,
        _valid_memo_id,
        _section_name_issue,
        _scan_claims_file_sections,
        _tokenize,
        parse_claims_file,
        _derive_claim_id,
        _claims_with_occurrence,
    )


def test_validate_depth_config_rejects_shallow_artifact():
    rp._validate_depth_config(retrieve_depth=20, top_k=5)  # ok, no raise
    with pytest.raises(ValueError, match="_RETRIEVE_DEPTH"):
        rp._validate_depth_config(retrieve_depth=3, top_k=5)


@pytest.mark.parametrize(
    "raw,expected",
    [
        (None, False), ("", False), ("false", False), ("False", False),
        ("0", False), ("no", False),
        ("true", True), ("True", True), ("1", True), ("yes", True), ("on", True),
    ],
)
def test_parse_bool_env(raw, expected):
    assert rp._parse_bool_env(raw) is expected


def test_embedding_client_from_env(monkeypatch):
    monkeypatch.setenv("EMBED_BASE_URL", "https://x/v1")
    monkeypatch.setenv("EMBED_API_KEY", "k")
    monkeypatch.setenv("EMBED_MODEL", "m")
    monkeypatch.delenv("EMBED_ASYMMETRIC", raising=False)
    c = rp.EmbeddingClient.from_env()
    assert (c.base_url, c.api_key, c.model, c.asymmetric) == ("https://x/v1", "k", "m", False)
    monkeypatch.setenv("EMBED_ASYMMETRIC", "true")
    assert rp.EmbeddingClient.from_env().asymmetric is True


def test_embedding_client_from_env_missing_var(monkeypatch):
    monkeypatch.delenv("EMBED_BASE_URL", raising=False)
    monkeypatch.setenv("EMBED_API_KEY", "k")
    monkeypatch.setenv("EMBED_MODEL", "m")
    with pytest.raises(EnvironmentError, match="EMBED_BASE_URL"):
        rp.EmbeddingClient.from_env()


import requests  # noqa: E402


class _FakeResp:
    def __init__(self, payload, status=200):
        self._payload = payload
        self.status_code = status

    def raise_for_status(self):
        if self.status_code >= 400:
            err = requests.HTTPError(f"{self.status_code} error for url")
            err.response = self
            raise err

    def json(self):
        return self._payload


def _ok_payload(texts):
    # Return embeddings out of order to prove call_embeddings sorts by index.
    data = [{"embedding": [float(len(t)), 1.0], "index": i} for i, t in enumerate(texts)]
    return {"data": list(reversed(data))}


def test_call_embeddings_no_input_type_by_default(monkeypatch):
    seen = {}

    def fake_post(url, headers=None, json=None, timeout=None):
        seen["url"] = url
        seen["body"] = json
        return _FakeResp(_ok_payload(json["input"]))

    monkeypatch.setattr(rp.requests, "post", fake_post)
    client = rp.EmbeddingClient("https://x/v1", "k", "m", asymmetric=False)
    out = rp.call_embeddings(["ab", "cde"], client)
    assert seen["url"] == "https://x/v1/embeddings"
    assert "input_type" not in seen["body"]
    assert out == [[2.0, 1.0], [3.0, 1.0]]  # sorted back into input order


def test_call_embeddings_sends_input_type_when_given(monkeypatch):
    seen = {}
    monkeypatch.setattr(
        rp.requests, "post",
        lambda url, headers=None, json=None, timeout=None: (
            seen.update(body=json) or _FakeResp(_ok_payload(json["input"]))
        ),
    )
    client = rp.EmbeddingClient("https://x/v1", "k", "m", asymmetric=True)
    rp.call_embeddings(["a"], client, input_type="passage")
    assert seen["body"]["input_type"] == "passage"


def test_call_embeddings_retries_transient_then_raises(monkeypatch):
    calls = {"n": 0}

    def flaky_post(url, headers=None, json=None, timeout=None):
        calls["n"] += 1
        return _FakeResp({"error": {"code": 503, "message": "upstream down"}}, status=503)

    monkeypatch.setattr(rp.requests, "post", flaky_post)
    monkeypatch.setattr(rp.time, "sleep", lambda s: None)
    client = rp.EmbeddingClient("https://x/v1", "k", "m", asymmetric=False)
    with pytest.raises(RuntimeError, match="upstream down"):
        rp.call_embeddings(["a"], client)
    assert calls["n"] == 2  # one retry on a 5xx


def test_call_embeddings_does_not_retry_client_error(monkeypatch):
    calls = {"n": 0}

    def bad_request(url, headers=None, json=None, timeout=None):
        calls["n"] += 1
        return _FakeResp(
            {"error": {"code": 400, "message": "invalid model", "metadata": {"error_type": "x"}}},
            status=400,
        )

    monkeypatch.setattr(rp.requests, "post", bad_request)
    client = rp.EmbeddingClient("https://x/v1", "k", "m", asymmetric=False)
    with pytest.raises(RuntimeError, match="invalid model"):
        rp.call_embeddings(["a"], client)
    assert calls["n"] == 1  # 400 is not retried


def test_call_embeddings_surfaces_error_on_200_body(monkeypatch):
    # OpenRouter has been seen returning 200 + {"error": {...}} and no data.
    monkeypatch.setattr(
        rp.requests, "post",
        lambda url, headers=None, json=None, timeout=None: _FakeResp(
            {"error": {"code": "x", "message": "model unavailable", "metadata": {}}}, status=200
        ),
    )
    monkeypatch.setattr(rp.time, "sleep", lambda s: None)
    client = rp.EmbeddingClient("https://x/v1", "k", "m", asymmetric=False)
    with pytest.raises(RuntimeError, match="model unavailable"):
        rp.call_embeddings(["a"], client)


def test_call_embeddings_retries_on_missing_data(monkeypatch):
    calls = {"n": 0}
    monkeypatch.setattr(
        rp.requests, "post",
        lambda url, headers=None, json=None, timeout=None: (
            calls.__setitem__("n", calls["n"] + 1) or _FakeResp({"no_data_key": True})
        ),
    )
    monkeypatch.setattr(rp.time, "sleep", lambda s: None)
    client = rp.EmbeddingClient("https://x/v1", "k", "m", asymmetric=False)
    with pytest.raises(RuntimeError):
        rp.call_embeddings(["a"], client)
    assert calls["n"] == 2


def test_call_embeddings_retries_on_non_dict_row(monkeypatch):
    # A malformed 200 body whose "data" rows aren't dicts (e.g. bare strings)
    # must hit the same "unparseable 200 -> retry once" path as a missing
    # key, not escape as a bare AttributeError from `r.get(...)`.
    calls = {"n": 0}
    monkeypatch.setattr(
        rp.requests, "post",
        lambda url, headers=None, json=None, timeout=None: (
            calls.__setitem__("n", calls["n"] + 1) or _FakeResp({"data": ["not-a-dict"]})
        ),
    )
    monkeypatch.setattr(rp.time, "sleep", lambda s: None)
    client = rp.EmbeddingClient("https://x/v1", "k", "m", asymmetric=False)
    with pytest.raises(RuntimeError):
        rp.call_embeddings(["a"], client)
    assert calls["n"] == 2  # retried once, then raised — not a bare AttributeError


def test_embed_texts_batches(monkeypatch):
    seen_batches = []
    monkeypatch.setattr(
        rp, "call_embeddings",
        lambda texts, client, *, input_type=None: (
            seen_batches.append(list(texts)) or [[float(i), 0.0] for i, _ in enumerate(texts)]
        ),
    )
    client = rp.EmbeddingClient("https://x/v1", "k", "m", asymmetric=False)
    out = rp._embed_texts([f"t{i}" for i in range(5)], client, input_type=None, batch_size=2)
    assert [len(b) for b in seen_batches] == [2, 2, 1]
    assert len(out) == 5


def test_call_embeddings_retries_on_transport_error(monkeypatch):
    calls = {"n": 0}

    def flaky_transport(url, headers=None, json=None, timeout=None):
        calls["n"] += 1
        raise requests.ConnectionError("connection refused")

    monkeypatch.setattr(rp.requests, "post", flaky_transport)
    monkeypatch.setattr(rp.time, "sleep", lambda s: None)
    client = rp.EmbeddingClient("https://x/v1", "k", "m", asymmetric=False)
    with pytest.raises(RuntimeError, match="transport error"):
        rp.call_embeddings(["a"], client)
    assert calls["n"] == 2  # one retry on transport error


def test_call_embeddings_raises_on_vector_count_mismatch(monkeypatch):
    # Response has fewer vectors than inputs
    monkeypatch.setattr(
        rp.requests, "post",
        lambda url, headers=None, json=None, timeout=None: _FakeResp(
            {"data": [{"embedding": [1.0, 2.0], "index": 0}]}
        ),
    )
    client = rp.EmbeddingClient("https://x/v1", "k", "m", asymmetric=False)
    with pytest.raises(RuntimeError, match="returned 1 vectors for 3 inputs"):
        rp.call_embeddings(["a", "b", "c"], client)


def test_call_embeddings_raises_on_empty_vector(monkeypatch):
    # Response contains an empty embedding
    monkeypatch.setattr(
        rp.requests, "post",
        lambda url, headers=None, json=None, timeout=None: _FakeResp(
            {"data": [
                {"embedding": [1.0, 2.0], "index": 0},
                {"embedding": [], "index": 1},
                {"embedding": [3.0, 4.0], "index": 2},
            ]}
        ),
    )
    client = rp.EmbeddingClient("https://x/v1", "k", "m", asymmetric=False)
    with pytest.raises(RuntimeError, match="empty vector"):
        rp.call_embeddings(["a", "b", "c"], client)


def _phrase_file(tmp_path, name, body):
    p = tmp_path / name
    p.write_text(textwrap.dedent(body), encoding="utf-8")
    return str(p)


_GOOD_PHRASE_YAML = """\
memo_id: MEMO-001
source_folder: sources/acme
sections:
  Business Profile:
    - "core business and asset portfolio"
    - "market cap and ranking"
  Ownership:
    - "major shareholders"
"""


def test_read_phrase_config_ok(tmp_path):
    path = _phrase_file(tmp_path, "MEMO-001.yaml", _GOOD_PHRASE_YAML)
    cfg = rp._read_phrase_config(path)
    assert cfg["memo_id"] == "MEMO-001"
    assert cfg["source_folder"] == "sources/acme"
    assert list(cfg["sections"]) == ["Business Profile", "Ownership"]
    assert cfg["sections"]["Ownership"] == ["major shareholders"]


@pytest.mark.parametrize(
    "name,body,match",
    [
        ("MEMO-001.yaml", "source_folder: s\nsections:\n  A:\n    - x\n", "memo_id"),
        ("MEMO-001.yaml", "memo_id: MEMO-001\nsource_folder: s\n", "sections"),
        ("OTHER.yaml", _GOOD_PHRASE_YAML, "must equal the filename stem"),
        ("MEMO-001.yaml", "memo_id: 'bad id'\nsource_folder: s\nsections:\n  A:\n    - x\n", r"\[A-Za-z0-9"),
        ("MEMO-001.yaml",
         "memo_id: MEMO-001\nsource_folder: s\nsections:\n  A: not-a-list\n", "must be a list"),
        ("MEMO-001.yaml",
         "memo_id: MEMO-001\nsource_folder: s\nsections:\n  A:\n    - '   '\n", "blank"),
        ("MEMO-001.yaml",
         "memo_id: MEMO-001\nsource_folder: s\nsections:\n  A: []\n", "1 to 20"),
        ("MEMO-001.yaml",
         "memo_id: MEMO-001\nsource_folder: s\nsections:\n  A:\n"
         + "".join(f"    - p{i}\n" for i in range(21)), "1 to 20"),
        ("MEMO-001.yaml",
         "memo_id: MEMO-001\nsource_folder: s\nnonsense: 1\nsections:\n  A:\n    - x\n",
         "unrecognized"),
        ("MEMO-001.yaml",
         "memo_id: MEMO-001\nsource_folder: s\nsections:\n  ' A ':\n    - x\n", "section name"),
        ("MEMO-001.yaml", "MEMO-001", "top level must be a mapping"),
        ("MEMO-001.yaml", "memo_id: MEMO-001\nsource_folder: 's\n", "not valid YAML"),
        ("MEMO-001.yaml",
         "memo_id: MEMO-001\nsections:\n  A:\n    - x\n", "source_folder"),
        ("MEMO-001.yaml",
         "memo_id: MEMO-001\nsource_folder: s\nsections: [1, 2]\n", "must be a mapping"),
    ],
)
def test_read_phrase_config_rejects(tmp_path, name, body, match):
    path = _phrase_file(tmp_path, name, body)
    with pytest.raises(ValueError, match=match):
        rp._read_phrase_config(path)


# ---------------------------------------------------------------------------
# Section 5: embed — chunk, cache, embed, atomic write
# ---------------------------------------------------------------------------


def _fake_vec(text: str, dim: int = 8) -> list[float]:
    h = hashlib.sha256(text.encode()).digest()
    return [b / 255.0 for b in h[:dim]]


@pytest.fixture
def stub_embed(monkeypatch):
    monkeypatch.setattr(
        rp, "call_embeddings",
        lambda texts, client, *, input_type=None: [_fake_vec(t) for t in texts],
    )


@pytest.fixture
def stub_pdf(monkeypatch):
    # one long doc -> deterministic chunk count via the real chunker
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "para. " * 400)


def _src_folder(tmp_path, text_by_name, folder_name="src"):
    d = tmp_path / folder_name
    d.mkdir(exist_ok=True)
    for n in text_by_name:
        (d / n).write_bytes(b"%PDF-1.4 stub")
    return str(d)


def _write_retrieval_yaml(tmp_path, memo_id, source_folder, sections):
    rdir = tmp_path / "retrieval"
    rdir.mkdir(exist_ok=True)
    body = {"memo_id": memo_id, "source_folder": source_folder, "sections": sections}
    (rdir / f"{memo_id}.yaml").write_text(yaml.safe_dump(body, sort_keys=False), encoding="utf-8")
    return str(rdir)


def _client():
    return rp.EmbeddingClient("https://x/v1", "k", "m", asymmetric=False)


def test_embed_fresh(tmp_path, stub_embed, monkeypatch):
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "para. " * 400)
    src = _src_folder(tmp_path, {"a.pdf": None})
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-1", src, {"Sec": ["q one", "q two"]})
    idir = str(tmp_path / "retrieval_index")
    rp.run_embed(rdir, idir, client=_client())

    idx = pd.read_parquet(os.path.join(idir, "MEMO-1.parquet"))
    assert list(idx.columns) == [
        "chunk_id", "doc_id", "chunk_text", "start_offset", "embedding", "model",
    ]
    assert len(idx) > 1
    assert idx["chunk_id"].is_unique
    assert not os.path.exists(os.path.join(idir, "MEMO-1.parquet.tmp"))
    assert all(len(v) == 8 for v in idx["embedding"])


def test_embed_incremental_reuses_unchanged(tmp_path, monkeypatch):
    calls = {"texts": []}
    monkeypatch.setattr(
        rp, "call_embeddings",
        lambda texts, client, *, input_type=None: (
            calls["texts"].extend(texts) or [_fake_vec(t) for t in texts]
        ),
    )
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "alpha. " * 200)
    src = _src_folder(tmp_path, {"a.pdf": None})
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-1", src, {"Sec": ["q"]})
    idir = str(tmp_path / "retrieval_index")
    rp.run_embed(rdir, idir, client=_client())
    first_count = len(calls["texts"])
    assert first_count > 0

    # change the PDF so exactly the last chunk's text differs
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "alpha. " * 200 + "OMEGA")
    calls["texts"].clear()
    rp.run_embed(rdir, idir, client=_client())
    # far fewer than a full re-embed — only the changed tail chunk(s)
    assert 0 < len(calls["texts"]) < first_count


def test_embed_drops_stale_chunk_ids(tmp_path, stub_embed, monkeypatch):
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "alpha. " * 200)
    src = _src_folder(tmp_path, {"a.pdf": None})
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-1", src, {"Sec": ["q"]})
    idir = str(tmp_path / "retrieval_index")
    rp.run_embed(rdir, idir, client=_client())
    big = len(pd.read_parquet(os.path.join(idir, "MEMO-1.parquet")))

    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "alpha. " * 20)  # much shorter
    rp.run_embed(rdir, idir, client=_client())
    small = len(pd.read_parquet(os.path.join(idir, "MEMO-1.parquet")))
    assert small < big


def test_embed_dimension_mismatch_raises(tmp_path, monkeypatch):
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "alpha. " * 200)
    src = _src_folder(tmp_path, {"a.pdf": None})
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-1", src, {"Sec": ["q"]})
    idir = tmp_path / "retrieval_index"
    idir.mkdir()
    # pre-write an index with 4-dim vectors for the chunk_ids that will be produced
    recs = rp._build_chunk_records(src, "MEMO-1")
    pd.DataFrame(
        [{**r, "embedding": [0.0, 0.0, 0.0, 0.0], "model": "m"} for r in recs]
    )[["chunk_id", "doc_id", "chunk_text", "start_offset", "embedding", "model"]].to_parquet(
        idir / "MEMO-1.parquet"
    )
    monkeypatch.setattr(
        rp, "call_embeddings",
        lambda texts, client, *, input_type=None: [[0.1, 0.2, 0.3] for _ in texts],  # 3-dim
    )
    # force a re-embed by changing one chunk
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "alpha. " * 200 + "OMEGA")
    with pytest.raises(ValueError, match="MEMO-1.parquet"):
        rp.run_embed(str(rdir), str(idir), client=_client())


def test_embed_records_model_provenance(tmp_path, stub_embed, monkeypatch):
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "para. " * 400)
    src = _src_folder(tmp_path, {"a.pdf": None})
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-1", src, {"Sec": ["q"]})
    idir = str(tmp_path / "retrieval_index")
    rp.run_embed(rdir, idir, client=_client())  # _client() model == "m"

    idx = pd.read_parquet(os.path.join(idir, "MEMO-1.parquet"))
    assert "model" in idx.columns
    assert (idx["model"] == "m").all()


def test_embed_model_change_same_dimension_raises_named_error(tmp_path, monkeypatch):
    # The dangerous case findings A calls out: EMBED_MODEL changes to a
    # DIFFERENT model that happens to embed at the SAME dimension. The old
    # dim-only guard would miss this entirely and silently mix vector spaces.
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "alpha. " * 200)
    monkeypatch.setattr(
        rp, "call_embeddings",
        lambda texts, client, *, input_type=None: [_fake_vec(t) for t in texts],
    )
    src = _src_folder(tmp_path, {"a.pdf": None})
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-1", src, {"Sec": ["q"]})
    idir = str(tmp_path / "retrieval_index")
    rp.run_embed(rdir, idir, client=_client())  # writes model "m"

    # Change EMBED_MODEL (same 8-dim vectors) and force a re-embed of one chunk.
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "alpha. " * 200 + "OMEGA")
    other_client = rp.EmbeddingClient("https://x/v1", "k", "m2", asymmetric=False)
    with pytest.raises(ValueError) as ei:
        rp.run_embed(rdir, idir, client=other_client)
    msg = str(ei.value)
    assert "MEMO-1.parquet" in msg
    assert "'m'" in msg and "'m2'" in msg
    assert "re-run embed" in msg


def test_embed_missing_model_column_warns_and_does_not_raise(tmp_path, monkeypatch, caplog):
    # An index written before model provenance was tracked has no 'model'
    # column at all. There's no way to know what wrote it — WARNING naming
    # the memo, never a hard failure.
    import logging

    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "alpha. " * 200)
    src = _src_folder(tmp_path, {"a.pdf": None})
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-1", src, {"Sec": ["q"]})
    idir = tmp_path / "retrieval_index"
    idir.mkdir()
    recs = rp._build_chunk_records(src, "MEMO-1")
    pd.DataFrame(
        [{**r, "embedding": _fake_vec(r["chunk_text"])} for r in recs]
    )[["chunk_id", "doc_id", "chunk_text", "start_offset", "embedding"]].to_parquet(
        idir / "MEMO-1.parquet"
    )
    monkeypatch.setattr(
        rp, "call_embeddings",
        lambda texts, client, *, input_type=None: [_fake_vec(t) for t in texts],
    )

    caplog.set_level(logging.WARNING, logger="retrieval")
    rp.run_embed(str(rdir), str(idir), client=_client())  # must not raise
    assert any(
        "MEMO-1" in rec.getMessage() and "model" in rec.getMessage()
        for rec in caplog.records
    )


def test_run_embed_sends_passage_input_type_when_asymmetric(tmp_path, monkeypatch):
    # Every other run_embed test uses asymmetric=False, so the
    # `"passage" if client.asymmetric else None` conditional in run_embed
    # only ever executes its None branch there — swapping "passage"/"query"
    # would pass the whole suite otherwise.
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "alpha. " * 200)
    seen = {"input_types": []}
    monkeypatch.setattr(
        rp, "call_embeddings",
        lambda texts, client, *, input_type=None: (
            seen["input_types"].append(input_type) or [_fake_vec(t) for t in texts]
        ),
    )
    src = _src_folder(tmp_path, {"a.pdf": None})
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-1", src, {"Sec": ["q"]})
    idir = str(tmp_path / "retrieval_index")
    client = rp.EmbeddingClient("https://x/v1", "k", "m", asymmetric=True)
    rp.run_embed(rdir, idir, client=client)
    assert seen["input_types"]  # at least one batch call happened
    assert all(t == "passage" for t in seen["input_types"])


def test_run_retrieve_sends_query_input_type_when_asymmetric(tmp_path, monkeypatch):
    # Same asymmetric-wiring gap on the retrieve side.
    rdir, idir = _embed_and_get_index(tmp_path, {"Sec": ["q"]}, monkeypatch)
    seen = {"input_types": []}
    monkeypatch.setattr(
        rp, "call_embeddings",
        lambda texts, client, *, input_type=None: (
            seen["input_types"].append(input_type) or [_fake_vec(t) for t in texts]
        ),
    )
    client = rp.EmbeddingClient("https://x/v1", "k", "m", asymmetric=True)
    rp.run_retrieve(rdir, idir, str(tmp_path / "r.parquet"),
                    claims_dir=str(tmp_path / "x"), client=client)
    assert seen["input_types"]
    assert all(t == "query" for t in seen["input_types"])


def test_embed_two_memos_writes_two_distinct_indexes(tmp_path, stub_embed, monkeypatch):
    # The actual production shape of a whole-corpus run: nothing exercised
    # run_embed writing more than one memo's index in the same call.
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "alpha. " * 200)
    src_a = _src_folder(tmp_path, {"a.pdf": None}, folder_name="src_a")
    src_b = _src_folder(tmp_path, {"b.pdf": None}, folder_name="src_b")
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-A", src_a, {"Sec": ["q"]})
    _write_retrieval_yaml(tmp_path, "MEMO-B", src_b, {"Sec": ["q1", "q2"]})
    idir = str(tmp_path / "retrieval_index")

    rp.run_embed(rdir, idir, client=_client())

    idx_a_path = os.path.join(idir, "MEMO-A.parquet")
    idx_b_path = os.path.join(idir, "MEMO-B.parquet")
    assert os.path.exists(idx_a_path)
    assert os.path.exists(idx_b_path)
    idx_a = pd.read_parquet(idx_a_path)
    idx_b = pd.read_parquet(idx_b_path)
    assert len(idx_a) > 0
    assert len(idx_b) > 0
    # distinct source PDFs -> distinct chunk_ids, no cross-contamination
    assert set(idx_a["chunk_id"]).isdisjoint(set(idx_b["chunk_id"]))


def test_retrieve_two_memos_concatenates_and_sorts_across_memos(tmp_path, monkeypatch):
    # The actual production shape of retrieve's whole-corpus concatenation
    # (line ~749) and cross-memo sort (line ~750) — previously exercised
    # only incidentally inside the two pre-pass-failure tests, never on the
    # happy path where both memos actually produce rows.
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "alpha. " * 200)
    monkeypatch.setattr(
        rp, "call_embeddings",
        lambda texts, client, *, input_type=None: [_fake_vec(t) for t in texts],
    )
    src_a = _src_folder(tmp_path, {"a.pdf": None}, folder_name="src_a")
    src_b = _src_folder(tmp_path, {"b.pdf": None}, folder_name="src_b")
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-A", src_a, {"Sec": ["q"]})
    _write_retrieval_yaml(tmp_path, "MEMO-B", src_b, {"Sec": ["q1", "q2"]})
    idir = str(tmp_path / "retrieval_index")
    rp.run_embed(rdir, idir, client=_client())

    df = rp.run_retrieve(rdir, idir, str(tmp_path / "r.parquet"),
                         claims_dir=str(tmp_path / "x"), client=_client())

    assert set(df["memo_id"]) == {"MEMO-A", "MEMO-B"}
    n_a = len(pd.read_parquet(os.path.join(idir, "MEMO-A.parquet")))
    n_b = len(pd.read_parquet(os.path.join(idir, "MEMO-B.parquet")))
    per_phrase_a = min(rp._RETRIEVE_DEPTH, n_a)
    per_phrase_b = min(rp._RETRIEVE_DEPTH, n_b)
    dense = df[df.method == "dense"]
    assert len(dense[dense.memo_id == "MEMO-A"]) == 1 * per_phrase_a  # 1 phrase
    assert len(dense[dense.memo_id == "MEMO-B"]) == 2 * per_phrase_b  # 2 phrases
    # actually sorted across memos by the documented key, not just concatenated
    expected = df.sort_values(
        ["memo_id", "section", "method", "phrase_index", "rank", "chunk_id"], kind="stable"
    ).reset_index(drop=True)
    pd.testing.assert_frame_equal(df, expected)


def test_embed_log_line_reports_batch_count(tmp_path, stub_embed, monkeypatch, caplog):
    import logging

    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "para. " * 400)
    src = _src_folder(tmp_path, {"a.pdf": None})
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-1", src, {"Sec": ["q"]})
    idir = str(tmp_path / "retrieval_index")

    caplog.set_level(logging.INFO, logger="retrieval")
    rp.run_embed(rdir, idir, client=_client())

    n_chunks = len(pd.read_parquet(os.path.join(idir, "MEMO-1.parquet")))
    expected_batches = -(-n_chunks // rp._EMBED_BATCH_SIZE)  # everything is new on a fresh run
    summary = [r.getMessage() for r in caplog.records if r.getMessage().startswith("embed MEMO-1")]
    assert len(summary) == 1
    assert f"{expected_batches}" in summary[0]
    assert "batch" in summary[0]


def test_embed_prepass_rejects_before_any_api_call(tmp_path, monkeypatch):
    # Two configs; the SECOND is malformed. The first must not be embedded.
    calls = {"n": 0}
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "alpha. " * 200)
    monkeypatch.setattr(
        rp, "call_embeddings",
        lambda texts, client, *, input_type=None: (
            calls.__setitem__("n", calls["n"] + 1) or [_fake_vec(t) for t in texts]
        ),
    )
    src = _src_folder(tmp_path, {"a.pdf": None})
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-1", src, {"Sec": ["q"]})
    (tmp_path / "retrieval" / "MEMO-2.yaml").write_text(
        "memo_id: MEMO-2\nsource_folder: s\n", encoding="utf-8"  # missing 'sections'
    )
    idir = str(tmp_path / "retrieval_index")
    with pytest.raises(ValueError, match="MEMO-2"):
        rp.run_embed(rdir, idir, client=_client())
    assert calls["n"] == 0
    assert not os.path.exists(os.path.join(idir, "MEMO-1.parquet"))


def test_embed_empty_pdfs_raise_named_error(tmp_path, stub_embed, monkeypatch):
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "")  # scanned / empty
    monkeypatch.setattr(gsp, "warn_if_text_suspiciously_short", lambda *a, **k: None)
    src = _src_folder(tmp_path, {"a.pdf": None})
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-1", src, {"Sec": ["q"]})
    with pytest.raises(ValueError, match="MEMO-1.*no chunks"):
        rp.run_embed(rdir, str(tmp_path / "retrieval_index"), client=_client())


# --- additional coverage: branches the brief's own test block doesn't reach ---


def test_embedding_dim_helper_branches():
    assert rp._embedding_dim([]) is None
    assert rp._embedding_dim([[1.0, 2.0], [3.0, 4.0]]) == 2
    with pytest.raises(ValueError, match="inconsistent dimensions"):
        rp._embedding_dim([[1.0, 2.0], [3.0, 4.0, 5.0]])


def test_load_index_absent_file_returns_empty(tmp_path):
    assert rp._load_index(str(tmp_path / "nope.parquet")) == {}


def test_load_index_reads_existing(tmp_path):
    p = tmp_path / "idx.parquet"
    pd.DataFrame(
        [{"chunk_id": "a_0", "doc_id": "a", "chunk_text": "hello", "start_offset": 0,
          "embedding": [1.0, 2.0], "model": "m"}]
    )[rp._INDEX_COLUMNS].to_parquet(p)
    loaded = rp._load_index(str(p))
    assert loaded == {"a_0": ("hello", [1.0, 2.0])}


def test_discover_phrase_configs_sorted_and_filtered(tmp_path):
    rdir = tmp_path / "retrieval"
    rdir.mkdir()
    for m in ["MEMO-B", "MEMO-A", "MEMO-C"]:
        (rdir / f"{m}.yaml").write_text(f"memo_id: {m}\n", encoding="utf-8")
    (rdir / "not_yaml.txt").write_text("x", encoding="utf-8")

    all_paths = rp._discover_phrase_configs(str(rdir), None)
    assert [os.path.basename(p) for p in all_paths] == ["MEMO-A.yaml", "MEMO-B.yaml", "MEMO-C.yaml"]

    filtered = rp._discover_phrase_configs(str(rdir), ["MEMO-C", "MEMO-A"])
    assert [os.path.basename(p) for p in filtered] == ["MEMO-A.yaml", "MEMO-C.yaml"]


def test_discover_phrase_configs_not_a_directory(tmp_path):
    with pytest.raises(ValueError, match="not a directory"):
        rp._discover_phrase_configs(str(tmp_path / "missing"), None)


def test_discover_phrase_configs_no_yaml_files(tmp_path):
    rdir = tmp_path / "retrieval"
    rdir.mkdir()
    (rdir / "readme.txt").write_text("x", encoding="utf-8")
    with pytest.raises(ValueError, match="no .yaml phrase configs"):
        rp._discover_phrase_configs(str(rdir), None)


def test_discover_phrase_configs_missing_memo_id(tmp_path):
    rdir = tmp_path / "retrieval"
    rdir.mkdir()
    (rdir / "MEMO-A.yaml").write_text("memo_id: MEMO-A\n", encoding="utf-8")
    with pytest.raises(ValueError, match="MEMO-Z"):
        rp._discover_phrase_configs(str(rdir), ["MEMO-A", "MEMO-Z"])


def test_discover_phrase_configs_warns_on_non_yaml_file(tmp_path, caplog):
    import logging

    rdir = tmp_path / "retrieval"
    rdir.mkdir()
    (rdir / "MEMO-A.yaml").write_text("memo_id: MEMO-A\n", encoding="utf-8")
    (rdir / "MEMO-B.yml").write_text("memo_id: MEMO-B\n", encoding="utf-8")

    caplog.set_level(logging.WARNING, logger="retrieval")
    paths = rp._discover_phrase_configs(str(rdir), None)

    # The .yml file is excluded, silently as far as the result goes...
    assert [os.path.basename(p) for p in paths] == ["MEMO-A.yaml"]
    # ...but a WARNING names it, so a whole-corpus retrieve can't look
    # complete while quietly skipping a memo.
    assert any(
        "MEMO-B.yml" in rec.getMessage() and rec.levelname == "WARNING"
        for rec in caplog.records
    )


def test_embed_prepass_rejects_bad_source_folder(tmp_path, monkeypatch):
    calls = {"n": 0}
    monkeypatch.setattr(
        rp, "call_embeddings",
        lambda texts, client, *, input_type=None: (
            calls.__setitem__("n", calls["n"] + 1) or [_fake_vec(t) for t in texts]
        ),
    )
    missing_src = str(tmp_path / "does_not_exist")
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-1", missing_src, {"Sec": ["q"]})
    idir = str(tmp_path / "retrieval_index")
    with pytest.raises(ValueError, match="not a directory"):
        rp.run_embed(rdir, idir, client=_client())
    assert calls["n"] == 0
    assert not os.path.exists(idir)


def test_embed_cached_index_itself_ragged_raises(tmp_path, monkeypatch):
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "alpha. " * 200)
    src = _src_folder(tmp_path, {"a.pdf": None})
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-1", src, {"Sec": ["q"]})
    idir = tmp_path / "retrieval_index"
    idir.mkdir()
    recs = rp._build_chunk_records(src, "MEMO-1")
    assert len(recs) >= 2
    rows = [{**r, "embedding": [0.0, 0.0], "model": "m"} for r in recs]
    rows[-1]["embedding"] = [0.0, 0.0, 0.0]  # one vector a different length -> ragged cache
    pd.DataFrame(rows)[rp._INDEX_COLUMNS].to_parquet(idir / "MEMO-1.parquet")

    with pytest.raises(ValueError, match=r"MEMO-1\.parquet.*inconsistent dimensions.*[Dd]elete.*re-run embed"):
        rp.run_embed(str(rdir), str(idir), client=_client())


def test_embed_atomic_write_failure_preserves_previous_index(tmp_path, stub_embed, monkeypatch):
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "alpha. " * 200)
    src = _src_folder(tmp_path, {"a.pdf": None})
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-1", src, {"Sec": ["q"]})
    idir = tmp_path / "retrieval_index"
    rp.run_embed(str(rdir), str(idir), client=_client())
    index_path = idir / "MEMO-1.parquet"
    before = index_path.read_bytes()

    def boom(self, path, *a, **k):
        # Write a few real bytes to the .tmp path before failing, so the
        # os.remove(tmp_path) cleanup branch is genuinely exercised against a
        # leftover file — not merely reachable code the raise never proves ran.
        with open(path, "wb") as f:
            f.write(b"partial")
        raise RuntimeError("disk full")

    monkeypatch.setattr(pd.DataFrame, "to_parquet", boom)
    with pytest.raises(RuntimeError, match="disk full"):
        rp.run_embed(str(rdir), str(idir), client=_client())

    assert index_path.read_bytes() == before  # previous index untouched
    assert not os.path.exists(str(index_path) + ".tmp")


# ---------------------------------------------------------------------------
# Section 6/7: retrieve — cosine search, results table, dedupe_by_section
# ---------------------------------------------------------------------------


def _embed_and_get_index(tmp_path, sections, monkeypatch, doc_text="alpha beta. " * 200):
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: doc_text)
    monkeypatch.setattr(
        rp, "call_embeddings",
        lambda texts, client, *, input_type=None: [_fake_vec(t) for t in texts],
    )
    src = _src_folder(tmp_path, {"a.pdf": None})
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-1", src, sections)
    idir = str(tmp_path / "retrieval_index")
    rp.run_embed(rdir, idir, client=_client())
    return rdir, idir


def test_retrieve_row_per_phrase_chunk(tmp_path, monkeypatch):
    rdir, idir = _embed_and_get_index(
        tmp_path, {"Business Profile": ["q one", "q two"], "Ownership": ["q three"]}, monkeypatch
    )
    results = str(tmp_path / "retrieval_results.parquet")
    df = rp.run_retrieve(rdir, idir, results, claims_dir=str(tmp_path / "no_claims"), client=_client())

    assert list(df.columns) == [
        "memo_id", "section", "phrase", "phrase_index", "method",
        "doc_id", "chunk_id", "chunk_text", "rank", "score",
    ]
    # 3 phrases, each with min(_RETRIEVE_DEPTH, n_chunks) dense rows
    n_chunks = len(pd.read_parquet(os.path.join(idir, "MEMO-1.parquet")))
    per_phrase = min(rp._RETRIEVE_DEPTH, n_chunks)
    dense = df[df.method == "dense"]
    assert len(dense) == 3 * per_phrase
    for (_, _, _), grp in dense.groupby(["memo_id", "section", "phrase"]):
        assert list(grp["rank"]) == list(range(1, per_phrase + 1))
    assert set(dense[dense.section == "Business Profile"]["phrase_index"]) == {0, 1}
    # "q one" etc. share no word with "alpha beta." -> keyword finds nothing,
    # and "both" is the dense list re-scored by rank fusion
    assert not (df.method == "keyword").any()
    assert len(df[df.method == "both"]) == 3 * per_phrase


def test_retrieve_scores_match_manual_cosine(tmp_path, monkeypatch):
    rdir, idir = _embed_and_get_index(tmp_path, {"Sec": ["hello world"]}, monkeypatch)
    df = rp.run_retrieve(
        rdir, idir, str(tmp_path / "r.parquet"),
        claims_dir=str(tmp_path / "nope"), client=_client(),
    )
    idx = pd.read_parquet(os.path.join(idir, "MEMO-1.parquet"))
    M = np.array([np.array(v, float) for v in idx["embedding"]])
    M = M / np.linalg.norm(M, axis=1, keepdims=True)
    q = np.array(_fake_vec("hello world"), float)
    q = q / np.linalg.norm(q)
    expected_top = float(np.sort(M @ q)[::-1][0])
    assert df[df.method == "dense"].iloc[0]["score"] == pytest.approx(expected_top, rel=1e-6)


def test_retrieve_deterministic_and_tie_broken_by_chunk_id(tmp_path, monkeypatch):
    # identical vectors for every chunk -> ties everywhere -> order must be chunk_id
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "same. " * 200)
    monkeypatch.setattr(
        rp, "call_embeddings",
        lambda texts, client, *, input_type=None: [[1.0, 0.0] for _ in texts],
    )
    src = _src_folder(tmp_path, {"a.pdf": None})
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-1", src, {"Sec": ["q"]})
    idir = str(tmp_path / "retrieval_index")
    rp.run_embed(rdir, idir, client=_client())
    r1 = str(tmp_path / "r1.parquet")
    r2 = str(tmp_path / "r2.parquet")
    d1 = rp.run_retrieve(rdir, idir, r1, claims_dir=str(tmp_path / "x"), client=_client())
    d2 = rp.run_retrieve(rdir, idir, r2, claims_dir=str(tmp_path / "x"), client=_client())
    pd.testing.assert_frame_equal(d1, d2)
    top = d1[(d1.method == "dense") & (d1["rank"] <= 3)]["chunk_id"].tolist()
    assert top == sorted(top)


def test_retrieve_depth_lt_chunk_count(tmp_path, monkeypatch):
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "alpha. " * 6)  # ~ a handful of chunks
    monkeypatch.setattr(
        rp, "call_embeddings",
        lambda texts, client, *, input_type=None: [_fake_vec(t) for t in texts],
    )
    src = _src_folder(tmp_path, {"a.pdf": None})
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-1", src, {"Sec": ["q"]})
    idir = str(tmp_path / "retrieval_index")
    rp.run_embed(rdir, idir, client=_client())
    n_chunks = len(pd.read_parquet(os.path.join(idir, "MEMO-1.parquet")))
    assert n_chunks < rp._RETRIEVE_DEPTH
    df = rp.run_retrieve(rdir, idir, str(tmp_path / "r.parquet"),
                         claims_dir=str(tmp_path / "x"), client=_client())
    assert list(df[df.method == "dense"]["rank"]) == list(range(1, n_chunks + 1))


def test_retrieve_before_embed_raises(tmp_path, monkeypatch):
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "alpha. " * 50)
    src = _src_folder(tmp_path, {"a.pdf": None})
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-1", src, {"Sec": ["q"]})
    with pytest.raises(ValueError, match="run .*embed"):
        rp.run_retrieve(rdir, str(tmp_path / "empty_index"),
                        str(tmp_path / "r.parquet"), claims_dir=str(tmp_path / "x"), client=_client())


def test_retrieve_stale_index_both_directions(tmp_path, monkeypatch):
    rdir, idir = _embed_and_get_index(tmp_path, {"Sec": ["q"]}, monkeypatch,
                                      doc_text="alpha. " * 200)
    # (a) a current chunk's text changed
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "alpha. " * 200 + "OMEGA")
    with pytest.raises(ValueError, match="run .*embed"):
        rp.run_retrieve(rdir, idir, str(tmp_path / "r.parquet"),
                        claims_dir=str(tmp_path / "x"), client=_client())
    # (b) the PDF got shorter — index holds chunk_ids no longer produced
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "alpha. " * 10)
    with pytest.raises(ValueError, match="run .*embed"):
        rp.run_retrieve(rdir, idir, str(tmp_path / "r.parquet"),
                        claims_dir=str(tmp_path / "x"), client=_client())


def test_dedupe_by_section(tmp_path, monkeypatch):
    rdir, idir = _embed_and_get_index(
        tmp_path, {"Sec": ["q one", "q two", "q three"]}, monkeypatch
    )
    df = rp.run_retrieve(rdir, idir, str(tmp_path / "r.parquet"),
                         claims_dir=str(tmp_path / "x"), client=_client())
    deduped = rp.dedupe_by_section(df)
    assert deduped.groupby(["memo_id", "section", "method"])["chunk_id"].apply(lambda s: s.is_unique).all()
    n_phrases = 3
    assert len(deduped) <= len(rp._RETRIEVAL_METHODS) * n_phrases * rp._TOP_K
    assert (deduped["rank"] <= rp._TOP_K).all()


def test_retrieve_large_corpus_exercises_top_k_sort(tmp_path, monkeypatch):
    # > _RETRIEVE_DEPTH chunks — the real-corpus path. "word " * 6000 = 30000
    # chars -> ~38 chunks at 1000/200.
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "word " * 6000)
    monkeypatch.setattr(
        rp, "call_embeddings",
        lambda texts, client, *, input_type=None: [_fake_vec(t) for t in texts],
    )
    src = _src_folder(tmp_path, {"a.pdf": None})
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-1", src, {"Sec": ["hello", "world"]})
    idir = str(tmp_path / "retrieval_index")
    rp.run_embed(rdir, idir, client=_client())
    n_chunks = len(pd.read_parquet(os.path.join(idir, "MEMO-1.parquet")))
    assert n_chunks > rp._RETRIEVE_DEPTH

    df = rp.run_retrieve(rdir, idir, str(tmp_path / "r.parquet"),
                         claims_dir=str(tmp_path / "x"), client=_client())
    dense = df[df.method == "dense"]
    for _, grp in dense.groupby(["section", "phrase"]):
        assert list(grp["rank"]) == list(range(1, rp._RETRIEVE_DEPTH + 1))
    # rank-1 row is the true cosine argmax
    idx = pd.read_parquet(os.path.join(idir, "MEMO-1.parquet"))
    M = np.array([np.array(v, float) for v in idx["embedding"]])
    M = M / np.linalg.norm(M, axis=1, keepdims=True)
    q = np.array(_fake_vec("hello"), float)
    q = q / np.linalg.norm(q)
    want = idx.iloc[int(np.argmax(M @ q))]["chunk_id"]
    got = dense[(dense.phrase == "hello") & (dense["rank"] == 1)].iloc[0]["chunk_id"]
    assert got == want


def test_retrieve_prepass_checks_all_indexes_before_embedding(tmp_path, monkeypatch):
    calls = {"n": 0}
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "alpha. " * 200)
    monkeypatch.setattr(
        rp, "call_embeddings",
        lambda texts, client, *, input_type=None: (
            calls.__setitem__("n", calls["n"] + 1) or [_fake_vec(t) for t in texts]
        ),
    )
    src = _src_folder(tmp_path, {"a.pdf": None})
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-1", src, {"Sec": ["q"]})
    _write_retrieval_yaml(tmp_path, "MEMO-2", src, {"Sec": ["q"]})
    idir = str(tmp_path / "retrieval_index")
    rp.run_embed(rdir, idir, ["MEMO-1"], client=_client())  # only MEMO-1 has an index
    calls["n"] = 0
    results = str(tmp_path / "retrieval_results.parquet")
    with pytest.raises(ValueError, match="MEMO-2.parquet"):
        rp.run_retrieve(rdir, idir, results, claims_dir=str(tmp_path / "x"), client=_client())
    assert calls["n"] == 0  # no phrases embedded — pre-pass caught MEMO-2 first
    assert not os.path.exists(results)


def test_retrieve_writes_xlsx(tmp_path, monkeypatch):
    from openpyxl import load_workbook

    rdir, idir = _embed_and_get_index(tmp_path, {"Sec": ["q"]}, monkeypatch)
    results = str(tmp_path / "retrieval_results.parquet")
    rp.run_retrieve(rdir, idir, results, claims_dir=str(tmp_path / "x"), client=_client())
    xlsx = results[:-len(".parquet")] + ".xlsx"
    assert os.path.exists(xlsx)
    ws = load_workbook(xlsx).active
    # phrase + chunk_text widened
    assert ws.column_dimensions["C"].width >= 40  # phrase is col C


def test_retrieve_xlsx_strips_control_chars_parquet_keeps_raw(tmp_path, monkeypatch):
    # pypdf emits NUL/control bytes for some PDFs' embedded fonts (see
    # tag_pipeline._plain's docstring); openpyxl raises IllegalCharacterError
    # on them, so run_retrieve's xlsx copy must be sanitized while the
    # parquet keeps the raw chunk_text.
    from openpyxl import load_workbook

    rdir, idir = _embed_and_get_index(
        tmp_path, {"Sec": ["q"]}, monkeypatch, doc_text="alpha\x00beta. " * 200
    )
    results = str(tmp_path / "retrieval_results.parquet")
    rp.run_retrieve(rdir, idir, results, claims_dir=str(tmp_path / "x"), client=_client())

    df = pd.read_parquet(results)
    assert df["chunk_text"].str.contains("\x00").any()  # parquet: raw text preserved

    xlsx = results[:-len(".parquet")] + ".xlsx"
    ws = load_workbook(xlsx).active
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    chunk_col = header.index("chunk_text")
    assert all("\x00" not in (row[chunk_col].value or "") for row in ws.iter_rows(min_row=2))


# --- additional coverage: branches this task implements but the brief's own
# test block (Step 1) doesn't reach ---


def test_retrieve_empty_index_raises(tmp_path, monkeypatch):
    rdir, idir = _embed_and_get_index(tmp_path, {"Sec": ["q"]}, monkeypatch)
    idx_path = os.path.join(idir, "MEMO-1.parquet")
    pd.DataFrame(columns=rp._INDEX_COLUMNS).to_parquet(idx_path, index=False)
    with pytest.raises(ValueError, match="empty"):
        rp.run_retrieve(rdir, idir, str(tmp_path / "r.parquet"),
                        claims_dir=str(tmp_path / "x"), client=_client())


def test_retrieve_ragged_index_embeddings_raise(tmp_path, monkeypatch):
    # A cached index whose chunk_id/chunk_text still match the current source
    # (so it clears _validate_retrieve_inputs' sync check) but whose stored
    # vectors are ragged must still fail — in _load_search_matrix, not the
    # sync check.
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "alpha beta. " * 200)
    src = _src_folder(tmp_path, {"a.pdf": None})
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-1", src, {"Sec": ["q"]})
    idir = tmp_path / "retrieval_index"
    idir.mkdir()
    recs = rp._build_chunk_records(src, "MEMO-1")
    assert len(recs) >= 2
    rows = [{**r, "embedding": [0.1, 0.2, 0.3], "model": "m"} for r in recs]
    rows[-1]["embedding"] = [0.1, 0.2]  # ragged: one shorter vector
    pd.DataFrame(rows)[rp._INDEX_COLUMNS].to_parquet(idir / "MEMO-1.parquet")

    with pytest.raises(ValueError, match="inconsistent dimensions"):
        rp.run_retrieve(str(rdir), str(idir), str(tmp_path / "r.parquet"),
                        claims_dir=str(tmp_path / "x"), client=_client())


def test_retrieve_phrase_index_dim_mismatch_raises_named_error(tmp_path, monkeypatch):
    # Same root cause as the embed-side model-provenance gap, other side: the
    # index is one dimension, the phrase embeddings come back a different
    # dimension (index built with model A, retrieve run against model B).
    # `M @ q` alone raises a bare numpy shape error naming nothing; this must
    # instead raise a ValueError naming the memo and both dimensions.
    rdir, idir = _embed_and_get_index(tmp_path, {"Sec": ["q"]}, monkeypatch)  # 8-dim index
    monkeypatch.setattr(
        rp, "call_embeddings",
        lambda texts, client, *, input_type=None: [[0.1, 0.2, 0.3] for _ in texts],  # 3-dim
    )
    with pytest.raises(ValueError) as ei:
        rp.run_retrieve(rdir, idir, str(tmp_path / "r.parquet"),
                        claims_dir=str(tmp_path / "x"), client=_client())
    msg = str(ei.value)
    assert "MEMO-1" in msg
    assert "3-dim" in msg or "3)" in msg
    assert "8-dim" in msg or "8)" in msg
    assert "EMBED_MODEL" in msg


def test_retrieve_zero_norm_chunk_vector_gets_zero_score(tmp_path, monkeypatch):
    # _load_search_matrix guards norm==0 (would otherwise divide by zero);
    # the resulting normalized-to-itself zero vector must score exactly 0
    # against any query, not NaN/inf.
    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "alpha beta. " * 200)
    monkeypatch.setattr(
        rp, "call_embeddings",
        lambda texts, client, *, input_type=None: [_fake_vec(t) for t in texts],
    )
    src = _src_folder(tmp_path, {"a.pdf": None})
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-1", src, {"Sec": ["hello world"]})
    idir = str(tmp_path / "retrieval_index")
    rp.run_embed(rdir, idir, client=_client())

    idx_path = os.path.join(idir, "MEMO-1.parquet")
    idx = pd.read_parquet(idx_path)
    dim = len(idx.iloc[0]["embedding"])
    zero_chunk_id = idx.iloc[0]["chunk_id"]
    idx.at[idx.index[0], "embedding"] = [0.0] * dim
    idx.to_parquet(idx_path, index=False)

    df = rp.run_retrieve(rdir, idir, str(tmp_path / "r.parquet"),
                         claims_dir=str(tmp_path / "x"), client=_client())
    zero_rows = df[(df.method == "dense") & (df["chunk_id"] == zero_chunk_id)]
    assert not zero_rows.empty
    assert (zero_rows["score"] == 0.0).all()
    assert not zero_rows["score"].isna().any()


def test_retrieve_all_scores_le_zero_warns(tmp_path, monkeypatch, caplog):
    # A degenerate all-zero phrase vector (the guarded qn==0 branch) makes
    # every chunk score exactly 0 <= 0, which must WARNING-log, naming
    # memo/section/phrase, rather than silently returning a useless top-k.
    import logging

    monkeypatch.setattr(gsp, "load_pdf_text", lambda p: "alpha beta. " * 200)

    def fake_call(texts, client, *, input_type=None):
        # chunk texts are built from "alpha beta. " content; the phrase text
        # ("zero phrase") is not, so this distinguishes the two calls without
        # relying on call order.
        return [
            _fake_vec(t) if "alpha beta" in t else [0.0] * 8
            for t in texts
        ]

    monkeypatch.setattr(rp, "call_embeddings", fake_call)
    src = _src_folder(tmp_path, {"a.pdf": None})
    rdir = _write_retrieval_yaml(tmp_path, "MEMO-1", src, {"Sec": ["zero phrase"]})
    idir = str(tmp_path / "retrieval_index")
    rp.run_embed(rdir, idir, client=_client())

    caplog.set_level(logging.WARNING, logger="retrieval")
    df = rp.run_retrieve(rdir, idir, str(tmp_path / "r.parquet"),
                         claims_dir=str(tmp_path / "x"), client=_client())
    assert (df[df.method == "dense"]["score"] == 0.0).all()
    assert any(
        "cosine score" in rec.message and "MEMO-1" in rec.message and "zero phrase" in rec.message
        for rec in caplog.records
    )


def test_retrieve_results_path_without_parquet_suffix(tmp_path, monkeypatch):
    rdir, idir = _embed_and_get_index(tmp_path, {"Sec": ["q"]}, monkeypatch)
    results = str(tmp_path / "results_data")  # deliberately no ".parquet" suffix
    df = rp.run_retrieve(rdir, idir, results, claims_dir=str(tmp_path / "x"), client=_client())
    assert not df.empty
    assert os.path.exists(results)
    assert os.path.exists(results + ".xlsx")


def test_dedupe_by_section_filters_rank_and_keeps_best_rank():
    df = pd.DataFrame(
        [
            {"memo_id": "M", "section": "S", "phrase": "p1", "phrase_index": 0, "method": "dense",
             "doc_id": "d", "chunk_id": "c1", "chunk_text": "t1", "rank": 1, "score": 0.9},
            {"memo_id": "M", "section": "S", "phrase": "p2", "phrase_index": 1, "method": "dense",
             "doc_id": "d", "chunk_id": "c1", "chunk_text": "t1", "rank": 2, "score": 0.99},
            {"memo_id": "M", "section": "S", "phrase": "p1", "phrase_index": 0, "method": "dense",
             "doc_id": "d", "chunk_id": "c2", "chunk_text": "t2", "rank": 6, "score": 0.5},
            {"memo_id": "M", "section": "S", "phrase": "p2", "phrase_index": 1, "method": "dense",
             "doc_id": "d", "chunk_id": "c3", "chunk_text": "t3", "rank": 3, "score": 0.3},
        ],
        columns=rp._RESULTS_COLUMNS,
    )
    out = rp.dedupe_by_section(df, top_k=5)
    # c2 only ever appears at rank 6 (> top_k=5) — must be dropped entirely.
    assert "c2" not in set(out["chunk_id"])
    # c1 appears at rank 1 (score 0.9) and rank 2 (score 0.99) — keep-best-RANK
    # must win over higher score: the rank-1 row survives.
    c1_row = out[out["chunk_id"] == "c1"].iloc[0]
    assert c1_row["rank"] == 1
    assert c1_row["score"] == 0.9
    assert len(out) == 2  # c1, c3


def test_dedupe_by_section_tie_broken_by_score():
    df = pd.DataFrame(
        [
            {"memo_id": "M", "section": "S", "phrase": "p1", "phrase_index": 0, "method": "dense",
             "doc_id": "d", "chunk_id": "c1", "chunk_text": "t1", "rank": 1, "score": 0.5},
            {"memo_id": "M", "section": "S", "phrase": "p2", "phrase_index": 1, "method": "dense",
             "doc_id": "d", "chunk_id": "c1", "chunk_text": "t1", "rank": 1, "score": 0.9},
        ],
        columns=rp._RESULTS_COLUMNS,
    )
    out = rp.dedupe_by_section(df)
    assert len(out) == 1
    assert out.iloc[0]["score"] == 0.9  # same rank -> higher score wins the tie


def _write_claims_md(tmp_path, memo_id, section_names):
    cdir = tmp_path / "claims"
    cdir.mkdir(exist_ok=True)
    body = "---\nmemo_id: {}\nsource_folder: s\n---\n\n".format(memo_id)
    for s in section_names:
        body += f"## {s}\n\n1. a claim.\n\n"
    (cdir / f"{memo_id}.md").write_text(body, encoding="utf-8")
    return str(cdir)


def test_drift_warns_both_directions(tmp_path, caplog):
    cdir = _write_claims_md(tmp_path, "MEMO-1", ["Business Profile", "Ownership"])
    with caplog.at_level("WARNING", logger=rp.logger.name):
        rp._claims_section_drift_warnings(
            "MEMO-1", ["Business Profile", "Industry Overview"], claims_dir=cdir
        )
    # Exactly 2 WARNINGs: one for phrase-only section, one for claims-only section
    assert len(caplog.records) == 2
    messages = [r.getMessage() for r in caplog.records]

    # Verify phrase-only section (in phrases but not in claims)
    assert any("Industry Overview" in msg for msg in messages)

    # Verify claims-only section (in claims but not in phrases)
    assert any("Ownership" in msg for msg in messages)

    # Clean case: matching section produces NO warning (no false positive)
    assert not any("Business Profile" in msg for msg in messages)


def test_drift_no_claims_file_is_silent(tmp_path, caplog):
    with caplog.at_level("WARNING", logger=rp.logger.name):
        rp._claims_section_drift_warnings("MEMO-1", ["Any"], claims_dir=str(tmp_path / "none"))
    assert caplog.text == ""


def test_drift_unreadable_claims_file_is_silent(tmp_path, caplog, monkeypatch):
    cdir = _write_claims_md(tmp_path, "MEMO-1", ["Business Profile"])

    def boom(*a, **k):
        raise UnicodeDecodeError("utf-8", b"", 0, 1, "bad")

    # Task 6 Step 3 imports _scan_claims_file_sections BY NAME into
    # retrieval_pipeline's namespace, so patch it there (rp.*), not on gsp.
    monkeypatch.setattr(rp, "_scan_claims_file_sections", boom)
    with caplog.at_level("WARNING", logger=rp.logger.name):
        rp._claims_section_drift_warnings("MEMO-1", ["Business Profile"], claims_dir=cdir)
    # no crash, no warning
    assert "Business Profile" not in caplog.text


# ---------------------------------------------------------------------------
# Section 8: CLI dispatch + demo
# ---------------------------------------------------------------------------


def test_main_no_arg_exits_with_usage(monkeypatch):
    called = {"from_env": False}
    monkeypatch.setattr(rp.EmbeddingClient, "from_env",
                        classmethod(lambda cls: called.__setitem__("from_env", True) or None))
    with pytest.raises(SystemExit) as ei:
        rp._main(["retrieval_pipeline.py"])
    assert "usage" in str(ei.value).lower()
    assert called["from_env"] is False


def test_main_unknown_command_exits_before_client(monkeypatch):
    called = {"from_env": False}
    monkeypatch.setattr(rp.EmbeddingClient, "from_env",
                        classmethod(lambda cls: called.__setitem__("from_env", True) or None))
    with pytest.raises(SystemExit, match="frobnicate"):
        rp._main(["retrieval_pipeline.py", "frobnicate"])
    assert called["from_env"] is False


def test_main_retrieve_rejects_memo_args(monkeypatch):
    called = {"from_env": False}
    monkeypatch.setattr(rp.EmbeddingClient, "from_env",
                        classmethod(lambda cls: called.__setitem__("from_env", True) or None))
    with pytest.raises(SystemExit, match="retrieve"):
        rp._main(["retrieval_pipeline.py", "retrieve", "MEMO-001"])
    assert called["from_env"] is False


def test_main_embed_with_memo_ids(monkeypatch):
    seen = {}
    monkeypatch.setattr(rp.EmbeddingClient, "from_env", classmethod(lambda cls: "CLIENT"))
    monkeypatch.setattr(rp, "run_embed",
                        lambda rd, idr, memo_ids, *, client: seen.update(memo_ids=memo_ids, client=client))
    rp._main(["retrieval_pipeline.py", "embed", "MEMO-001", "MEMO-002"])
    assert seen["memo_ids"] == ["MEMO-001", "MEMO-002"]
    assert seen["client"] == "CLIENT"


def test_main_embed_no_memo_ids(monkeypatch):
    seen = {}
    monkeypatch.setattr(rp.EmbeddingClient, "from_env", classmethod(lambda cls: "CLIENT"))
    monkeypatch.setattr(rp, "run_embed",
                        lambda rd, idr, memo_ids, *, client: seen.update(memo_ids=memo_ids, client=client))
    rp._main(["retrieval_pipeline.py", "embed"])
    assert seen["memo_ids"] is None
    assert seen["client"] == "CLIENT"


def test_main_retrieve_runs(monkeypatch):
    seen = {}
    monkeypatch.setattr(rp.EmbeddingClient, "from_env", classmethod(lambda cls: "CLIENT"))
    monkeypatch.setattr(
        rp, "run_retrieve",
        # dict.__setitem__ returns None, so `None or DF` yields the DataFrame;
        # dict.setdefault would return True and short-circuit — don't use it here.
        lambda *a, **k: seen.__setitem__("ran", True) or pd.DataFrame(columns=rp._RESULTS_COLUMNS),
    )
    monkeypatch.setattr(rp, "dedupe_by_section", lambda df, **k: df)
    rp._main(["retrieval_pipeline.py", "retrieve"])
    assert seen["ran"] is True


# ---------------------------------------------------------------------------
# Keyword (BM25) and combined (RRF) search
# ---------------------------------------------------------------------------
from rank_bm25 import BM25Okapi  # noqa: E402


RARE_WORD_DOC = "alpha beta. " * 300 + "gamma delta. " * 20  # "gamma" only in the last chunk


def test_retrieve_writes_all_three_methods(tmp_path, monkeypatch):
    # In this tiny corpus every word but "gamma"/"delta" is in every chunk, so
    # BM25's idf floor (0.25 x the average idf) is itself negative and a common
    # word scores nothing. A rare word guarantees keyword rows. (In a real
    # corpus the floor is small and positive, so common words still score.)
    rdir, idir = _embed_and_get_index(tmp_path, {"Sec": ["gamma"]}, monkeypatch, doc_text=RARE_WORD_DOC)
    df = rp.run_retrieve(rdir, idir, str(tmp_path / "r.parquet"),
                         claims_dir=str(tmp_path / "x"), client=_client())
    assert set(df["method"]) == set(rp._RETRIEVAL_METHODS)
    for method, grp in df.groupby("method"):
        assert list(grp["rank"]) == list(range(1, len(grp) + 1)), method


def test_keyword_scores_are_bm25_over_the_index(tmp_path, monkeypatch):
    # "gamma" occurs only near the end, so it is rare (positive BM25 idf) and
    # most chunks score 0 for it — which also exercises the zero-score drop.
    rdir, idir = _embed_and_get_index(tmp_path, {"Sec": ["gamma"]}, monkeypatch, doc_text=RARE_WORD_DOC)
    df = rp.run_retrieve(rdir, idir, str(tmp_path / "r.parquet"),
                         claims_dir=str(tmp_path / "x"), client=_client())
    kw = df[df.method == "keyword"]
    idx = pd.read_parquet(os.path.join(idir, "MEMO-1.parquet"))
    bm25 = BM25Okapi([gsp._tokenize(t) for t in idx["chunk_text"]])
    scores = dict(zip(idx["chunk_id"], bm25.get_scores(["gamma"])))
    expected = sorted((c for c in scores if scores[c] > 0), key=lambda c: (-scores[c], c))
    assert expected, "fixture must give 'gamma' a positive score somewhere"
    assert kw["chunk_id"].tolist() == expected[: rp._RETRIEVE_DEPTH]
    assert kw["score"].tolist() == pytest.approx([scores[c] for c in expected[: rp._RETRIEVE_DEPTH]])


def test_keyword_emits_nothing_for_a_phrase_sharing_no_word(tmp_path, monkeypatch):
    rdir, idir = _embed_and_get_index(tmp_path, {"Sec": ["zzz"]}, monkeypatch)
    df = rp.run_retrieve(rdir, idir, str(tmp_path / "r.parquet"),
                         claims_dir=str(tmp_path / "x"), client=_client())
    assert not (df.method == "keyword").any()
    dense = df[df.method == "dense"].sort_values("rank")["chunk_id"].tolist()
    both = df[df.method == "both"].sort_values("rank")["chunk_id"].tolist()
    assert both == dense  # fusing a single list keeps its order


def test_rank_by_score_drops_nonpositive_only_when_asked():
    scores = [0.0, 2.0, -1.0, 2.0]
    ids = ["a", "b", "c", "d"]
    assert rp._rank_by_score(scores, ids, 10, drop_nonpositive=True) == [(1, 2.0), (3, 2.0)]
    assert [i for i, _ in rp._rank_by_score(scores, ids, 10, drop_nonpositive=False)] == [1, 3, 0, 2]


def test_fuse_rrf_sums_reciprocal_ranks():
    dense = [(0, 0.9), (1, 0.8)]
    keyword = [(1, 5.0), (2, 4.0)]
    out = rp._fuse_rrf(dense, keyword, ["c0", "c1", "c2"], k=20)
    assert [i for i, _ in out] == [1, 0, 2]
    assert [s for _, s in out] == pytest.approx([1 / 62 + 1 / 61, 1 / 61, 1 / 62])


def test_fuse_rrf_breaks_an_exact_tie_in_favour_of_dense():
    # dense #1 (c1) and keyword #1 (c0) both score exactly 1/61; the dense
    # rank decides, not the chunk_id's alphabetical order.
    out = rp._fuse_rrf([(1, 0.9)], [(0, 3.0)], ["c0", "c1"], k=20)
    assert [i for i, _ in out] == [1, 0]


def test_fuse_rrf_truncates_to_k():
    out = rp._fuse_rrf([(0, 1.0), (1, 0.9), (2, 0.8)], [], ["a", "b", "c"], k=2)
    assert [i for i, _ in out] == [0, 1]


def test_dedupe_by_section_keeps_methods_apart():
    base = {"memo_id": "M", "section": "S", "phrase": "p", "phrase_index": 0,
            "doc_id": "d", "chunk_id": "c1", "chunk_text": "t", "rank": 1, "score": 1.0}
    df = pd.DataFrame([{**base, "method": "dense"}, {**base, "method": "keyword"}],
                      columns=rp._RESULTS_COLUMNS)
    out = rp.dedupe_by_section(df)
    assert sorted(out["method"]) == ["dense", "keyword"]


# ---------------------------------------------------------------------------
# recheck — each claim's own text as a dense query
# ---------------------------------------------------------------------------
CLAIMS_T = ["Acme sells widgets.", "Acme runs five plants."]


def _claims_dir_with(tmp_path, memo_id="MEMO-1", section="Business Profile"):
    cdir = tmp_path / "claims"
    cdir.mkdir(exist_ok=True)
    gsp.write_claims_file(str(cdir / f"{memo_id}.md"), memo_id, "src", {}, [(section, CLAIMS_T)])
    return str(cdir)


def test_recheck_writes_top_k_per_claim(tmp_path, monkeypatch):
    rdir, idir = _embed_and_get_index(tmp_path, {"Business Profile": ["q"]}, monkeypatch)
    cdir = _claims_dir_with(tmp_path)
    out = str(tmp_path / "claim_queries.parquet")
    df = rp.run_recheck(rdir, idir, cdir, out, client=_client())

    assert list(df.columns) == rp._CLAIM_QUERY_COLUMNS
    n_chunks = len(pd.read_parquet(os.path.join(idir, "MEMO-1.parquet")))
    per_claim = min(rp._TOP_K, n_chunks)
    assert len(df) == len(CLAIMS_T) * per_claim
    assert set(df["claim_id"]) == {gsp._derive_claim_id("MEMO-1", "Business Profile", c, 0) for c in CLAIMS_T}
    for _, grp in df.groupby("claim_id"):
        assert list(grp["rank"]) == list(range(1, per_claim + 1))
    pd.testing.assert_frame_equal(pd.read_parquet(out), df)


def test_recheck_sends_query_input_type_when_asymmetric(tmp_path, monkeypatch):
    rdir, idir = _embed_and_get_index(tmp_path, {"Business Profile": ["q"]}, monkeypatch)
    cdir = _claims_dir_with(tmp_path)
    seen = []
    monkeypatch.setattr(
        rp, "call_embeddings",
        lambda texts, client, *, input_type=None: seen.append(input_type) or [_fake_vec(t) for t in texts],
    )
    client = rp.EmbeddingClient("https://x/v1", "k", "m", asymmetric=True)
    rp.run_recheck(rdir, idir, cdir, str(tmp_path / "q.parquet"), client=client)
    assert seen and all(t == "query" for t in seen)


def test_recheck_missing_claims_file_raises_before_embedding(tmp_path, monkeypatch):
    rdir, idir = _embed_and_get_index(tmp_path, {"Sec": ["q"]}, monkeypatch)
    calls = []
    monkeypatch.setattr(
        rp, "call_embeddings",
        lambda texts, client, *, input_type=None: calls.append(1) or [_fake_vec(t) for t in texts],
    )
    with pytest.raises(ValueError, match="recheck embeds each claim"):
        rp.run_recheck(rdir, idir, str(tmp_path / "no_claims"), str(tmp_path / "q.parquet"), client=_client())
    assert calls == []


def test_main_recheck_runs(monkeypatch):
    seen = {}
    monkeypatch.setattr(rp.EmbeddingClient, "from_env", classmethod(lambda cls: "CLIENT"))
    monkeypatch.setattr(
        rp, "run_recheck",
        lambda *a, **k: seen.__setitem__("client", k["client"]) or pd.DataFrame(columns=rp._CLAIM_QUERY_COLUMNS),
    )
    rp._main(["retrieval_pipeline.py", "recheck"])
    assert seen["client"] == "CLIENT"


def test_main_recheck_rejects_args(monkeypatch):
    called = {"from_env": False}
    monkeypatch.setattr(rp.EmbeddingClient, "from_env",
                        classmethod(lambda cls: called.__setitem__("from_env", True) or None))
    with pytest.raises(SystemExit, match="recheck"):
        rp._main(["retrieval_pipeline.py", "recheck", "MEMO-001"])
    assert called["from_env"] is False


# ---------------------------------------------------------------------------
# retrieve / recheck refuse an index embedded with a different model
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("command", ["retrieve", "recheck"])
def test_query_model_must_match_the_index_model(tmp_path, monkeypatch, command):
    # Same dimension, different model: the dimension check cannot see it, and
    # the results would silently mix two vector spaces.
    rdir, idir = _embed_and_get_index(tmp_path, {"Business Profile": ["q"]}, monkeypatch)
    cdir = _claims_dir_with(tmp_path)
    calls = []
    monkeypatch.setattr(
        rp, "call_embeddings",
        lambda texts, client, *, input_type=None: calls.append(1) or [_fake_vec(t) for t in texts],
    )
    other = rp.EmbeddingClient("https://x/v1", "k", "other-model", asymmetric=False)
    with pytest.raises(ValueError, match="EMBED_MODEL"):
        if command == "retrieve":
            rp.run_retrieve(rdir, idir, str(tmp_path / "r.parquet"), claims_dir=cdir, client=other)
        else:
            rp.run_recheck(rdir, idir, cdir, str(tmp_path / "q.parquet"), client=other)
    assert calls == []   # refused in the pre-pass, before any embedding call


# ---------------------------------------------------------------------------
# provenance: retrieve / recheck record what produced their file
# ---------------------------------------------------------------------------
def test_retrieve_and_recheck_record_their_provenance(tmp_path, monkeypatch):
    rdir, idir = _embed_and_get_index(tmp_path, {"Business Profile": ["q"]}, monkeypatch)
    cdir = _claims_dir_with(tmp_path)
    fingerprint = rp._index_fingerprint(pd.read_parquet(os.path.join(idir, "MEMO-1.parquet")))
    results, queries = str(tmp_path / "r.parquet"), str(tmp_path / "q.parquet")
    rp.run_retrieve(rdir, idir, results, claims_dir=cdir, client=_client())
    rp.run_recheck(rdir, idir, cdir, queries, client=_client())
    assert rp.read_provenance(results) == {"command": "retrieve", "model": "m", "depth": rp._RETRIEVE_DEPTH,
                                           "indexes": {"MEMO-1": fingerprint}}
    assert rp.read_provenance(queries) == {"command": "recheck", "model": "m", "depth": rp._TOP_K,
                                           "indexes": {"MEMO-1": fingerprint}}
    assert not pd.read_parquet(results).empty   # still an ordinary parquet for every other reader


def test_an_interrupted_provenance_write_keeps_the_last_good_file(tmp_path, monkeypatch):
    path = str(tmp_path / "r.parquet")
    rp._write_with_provenance(pd.DataFrame({"a": [1]}), path, {"command": "retrieve"})

    def crash(table, where, *args, **kwargs):
        open(where, "wb").write(b"PAR1 trunc")          # a half-written file, then the crash
        raise KeyboardInterrupt
    monkeypatch.setattr(rp.pq, "write_table", crash)
    with pytest.raises(KeyboardInterrupt):
        rp._write_with_provenance(pd.DataFrame({"a": [2]}), path, {"command": "retrieve"})
    assert pd.read_parquet(path)["a"].tolist() == [1]
    assert os.listdir(tmp_path) == ["r.parquet"]


def test_index_fingerprint_follows_the_chunk_text_not_the_row_order():
    index = pd.DataFrame({"chunk_id": ["a.pdf_0", "a.pdf_1"], "chunk_text": ["one", "two"]})
    assert rp._index_fingerprint(index) == rp._index_fingerprint(index.iloc[::-1])
    assert rp._index_fingerprint(index) != rp._index_fingerprint(index.assign(chunk_text=["one", "TWO"]))


def test_a_plain_parquet_has_no_provenance(tmp_path):
    path = str(tmp_path / "plain.parquet")
    pd.DataFrame({"a": [1]}).to_parquet(path, index=False)
    assert rp.read_provenance(path) is None
