# %% [markdown]
# # tag_pipeline.py — evidence tagging for the Phase 3 eval
#
# Standalone; not wired into extract/build. Two commands:
#
#   python tag_pipeline.py draft [memo_id ...]   # AI drafts one verdict per claim -> review/<memo_id>.xlsx
#   python tag_pipeline.py finalize <memo_id>     # the checked sheet -> reviewed/<memo_id>.xlsx (the eval's input)
#
# `draft` never writes `tag`; `finalize` writes it only from the user's
# CHECKED answers, so `tag` stays human-approved. The model is pinned in code
# (_BUNDLE_MODEL), NOT from .env — .env's LLM_MODEL drives extract/build only.
#
# Design: docs/superpowers/specs/2026-09-13-tag-pipeline-v2-bundle-design.md
# (§12-§13 addenda). The earlier row-mode tagger it replaced is recorded in
# CLAUDE.md design decision 18 and docs/design-decisions/18-evidence-tagging.md.

# %%
import hashlib
import logging
import os
import re
import sys
import tempfile
from datetime import date

import pandas as pd
from dotenv import load_dotenv
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from golden_set_pipeline import (
    LLMClient,
    _call_llm_with_json_retry,
    export_for_review,
    _is_deterministic_claim_id,
    build_chunk_index,
    parse_claims_file,
    read_filing_entity,
    _load_source_documents,
    _derive_claim_id,
    _claims_with_occurrence,
    _chunk_index_of,
    _valid_memo_id,
    _HUMAN_ADDED_PREFIX,
    _QUOTE_SEPARATOR,
)

load_dotenv()
logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
logger = logging.getLogger("tag")

# %% [markdown]
# ## 1. Filing entities

# %%
def _filing_entity(memo_id: str, claims_dir: str) -> str:
    """
    The company the memo is about, from its claims file's `filing_entity`
    frontmatter key; the draft prompt tells the model that "we", "our", "the
    Group" mean this entity. Kept in the (gitignored) claims file rather than
    in this code so no real company name has to be committed. Raises
    DraftRefused naming the memo if the claims file is missing, unreadable or
    malformed (read_claims_file's own messages) or does not set the key.
    Reads no PDF.
    """
    path = os.path.join(claims_dir, f"{memo_id}.md")
    try:
        read_claims_file(memo_id, claims_dir)     # the whole file must parse, not only its frontmatter
        entity = read_filing_entity(path)
    except ValueError as exc:
        raise DraftRefused(f"{memo_id}: {exc}") from exc
    if entity is None:
        raise DraftRefused(f"{memo_id}: {path} has no filing_entity — add `filing_entity: <company name>` "
                           f"to its frontmatter")
    return entity


# %% [markdown]
# ## 2. draft — one verdict per claim over all its chunks
#
# Spec: docs/superpowers/specs/2026-09-13-tag-pipeline-v2-bundle-design.md.
# §9 MEMO-004 acceptance: measured FAIL on chunk recall only, accepted by the
# user 2026-09-14 (docs/prompt-verification-log.md).

# %%
# Pre-registered primary model for the §9 acceptance test (spec §8/§9.1).
# NOT from .env: a model change is a reviewed code edit, and it voids the
# acceptance result. The corporate on-prem fork must re-point this AND re-run
# §9 against its model (eval spec, port ledger).
_BUNDLE_MODEL = "google/gemini-3.1-pro-preview"

# Hidden reasoning tokens count against max_tokens, and call_llm sends no
# thinking budget. At 4096, gemini-2.5-pro truncated 31% of JSON answers
# (decision 18's probe; mechanism: decision 4). 16000 passed the §9.0
# pre-flight on a 22-chunk claim (gemini-3.1 used 12% of it for reasoning) and
# the §9 acceptance run (192 answers, 0 retries) — docs/prompt-verification-log.md.
_BUNDLE_MAX_TOKENS = 16000

# Source text shown before/after each chunk. Carries the table column
# headings pypdf strips from the chunk itself.
_CONTEXT_CHARS = 500

# About double the largest bundle seen (22 chunks, measured before memos 1-3
# are rebuilt at a wider min_candidates). An anomaly tripwire that fails loud,
# not an API limit.
_MAX_CHUNKS_PER_CLAIM = 40

_VERDICTS = ("stated_directly", "needs_combining", "not_supported")
_DRAFT_FAILED = "draft_failed"

# The eval spec keys its dense re-review cross-check on the prefix
# "auto: found=False" — keep it exactly.
_AUTO_NOT_FOUND_REASON = "auto: found=False, nothing was found by search"

# FROZEN for the spec's §9 acceptance test (Appendix A). Any wording change
# voids that result; its SHA-256 is recorded in docs/prompt-verification-log.md.
# Filled with str.replace, not str.format: the JSON block has literal braces.
_BUNDLE_PROMPT = """You are checking whether source passages support one claim from a financial
memo about __ENTITY__. In the passages, "we", "our", "the Group" and "the
Company" mean __ENTITY__.

CLAIM: __CLAIM__

PASSAGES (each with the source text just before and after it, for context):
__PASSAGES__

Choose exactly one verdict for the claim, looking at all passages together:

- "stated_directly": at least one passage, on its own, states the claim's
  fact — same subject, same measure, same period. Different wording, units or
  rounding still count (EUR 4,203m supports "EUR 4.2bn"; 29.04% supports
  "29%").
- "needs_combining": the claim is true, but only by putting passages together,
  by a calculation (for example a part divided by a total), or by describing
  numbers with a word the passages do not use (for example calling a +0.4%
  change "flat").
- "not_supported": the passages, even together, do not establish the claim.
  This includes passages about a different measure, period or subject that
  merely look similar, and a calculation whose result does not round to the
  claim's number. Do not pick the closest passage from a set that does not
  support the claim.

Then list the passages needed:
- stated_directly: each passage that states the claim on its own;
- needs_combining: every passage the combination or calculation uses;
- not_supported: none.
Only list a passage if its own PASSAGE text carries what the claim needs. Use
the before/after text to understand what a passage's numbers mean (for
example table column headings), but never list a passage for its context
alone.

Worked examples (Acme Ltd, made up):
1. Claim "Acme's FY24 revenue was EUR 4.2bn." [A] "Revenue for FY24 was
   EUR 4,203m." -> stated_directly, needed [A]. Same number, different units.
2. Claim "Acme's revenue was flat in FY24." [A] "Revenue EUR 1,203m
   (FY23: EUR 1,198m)." -> needs_combining, needed [A]. "Flat" is a judgement
   about the numbers, not something the passage says.
3. Claim "Chilled products are 25% of Acme's revenue." [A] "Chilled revenue
   EUR 300m" [B] "Group revenue EUR 1,200m" -> needs_combining, needed [A, B].
   300 / 1,200 = 25%.
4. Claim "Chilled products are 25% of Acme's revenue." [A] "Chilled
   accounts for 25% of Acme's headcount." -> not_supported, needed [].
   Same number, different measure.

Answer with JSON only, no other text:
{"verdict": "stated_directly" | "needs_combining" | "not_supported",
 "needed": ["A", ...],
 "reason": "one or two sentences naming the passages and, for a calculation,
            the numbers used"}"""


def bundle_prompt_sha256() -> str:
    """SHA-256 of the frozen template — recorded with every acceptance run so a
    silently edited prompt can't be scored as if it were the pre-registered one."""
    return hashlib.sha256(_BUNDLE_PROMPT.encode("utf-8")).hexdigest()


# %%
def read_claims_file(memo_id: str, claims_dir: str = "claims") -> tuple[str, list[tuple[str, str, str]]]:
    """
    One memo's claims file, read once for everything tag_pipeline needs from
    it: (source_folder, [(section, claim_text, claim_id), ...]) in file order,
    with ids derived exactly as `build` derives them. Reads no PDF.

    Raises ValueError if the file is missing, is not a regular file, cannot be
    read (an OSError such as permission denied; a non-UTF-8 file's
    UnicodeDecodeError already is a ValueError), does not parse, or declares
    a different memo_id. Messages do not start with the
    memo id: the caller adds it once (DraftRefused, FinalizeRefused).
    """
    path = os.path.join(claims_dir, f"{memo_id}.md")
    if not os.path.exists(path):
        raise ValueError(f"{path} not found — it gives the memo's source_folder and claim order")
    if not os.path.isfile(path):
        raise ValueError(f"{path} is not a file — it should be the memo's claims file")
    try:
        file_memo_id, source_folder, _, sections = parse_claims_file(path)
    except OSError as exc:
        raise ValueError(f"{path} could not be read: {exc}") from exc
    if file_memo_id != memo_id:
        raise ValueError(f"{path} declares memo_id {file_memo_id!r}")
    claims = [(section, text, _derive_claim_id(memo_id, section, text, occurrence))
              for section, texts in sections
              for text, occurrence in _claims_with_occurrence(texts)]
    return source_folder, claims


def load_memo_chunks(memo_id: str, source_folder: str) -> tuple[dict[str, dict], dict[str, str]]:
    """
    Rebuild one memo's chunk index the way `build` did: load the PDFs in
    `source_folder` (from read_claims_file — the checkpoint does not store
    it) with pypdf and chunk at the 1000/200 defaults. Returns
    (chunks_by_id, documents) — chunk dicts keyed by chunk_id, and doc_id ->
    full text for slicing context. `memo_id` only labels log and error
    messages. Slow on large filings (MEMO-004's three PDFs: ~26 s).

    Raises ValueError when the folder has no PDFs, or a PDF cannot be read
    (an OSError such as permission denied, or a pypdf error for a corrupt
    file), so draft and finalize refuse instead of showing a traceback.
    """
    from pypdf.errors import PyPdfError

    try:
        documents = _load_source_documents(source_folder, memo_id)
    except (OSError, PyPdfError) as exc:
        raise ValueError(f"source_folder {source_folder!r}: a PDF could not be read "
                         f"({type(exc).__name__}: {exc})") from exc
    chunks_by_id = {c["chunk_id"]: c for c in build_chunk_index(documents)}
    return chunks_by_id, dict(documents)


def _source_docs(documents: dict[str, str]) -> str:
    """The memo's PDFs as one comparable string: every loaded document (also
    one whose extracted text is empty and so yields no chunks), sorted by
    doc_id, each as '<doc_id> (<first 12 hex of the SHA-256 of its extracted
    text>)', joined by ' | '. So a PDF added, removed, or edited in place
    under its old name all change it. draft writes it on every claim row;
    finalize compares it with what source_folder gives now."""
    return " | ".join(f"{doc_id} ({hashlib.sha256(text.encode('utf-8', 'surrogatepass')).hexdigest()[:12]})"
                      for doc_id, text in sorted(documents.items()))


def _label(i: int) -> str:
    """0 -> 'A', 25 -> 'Z', 26 -> 'AA' (spreadsheet-column style)."""
    label = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        label = chr(65 + r) + label
    return label


def _collapse_ws(text: str) -> str:
    """Collapse every run of whitespace (PDF layout line breaks, tabs, repeated
    spaces) to one space and strip both ends."""
    return " ".join(text.split())


def build_bundle(claim_id: str,
                 chunk_ids: list[str],
                 chunks_by_id: dict[str, dict],
                 documents: dict[str, str]) -> list[dict]:
    """
    Turn one claim's found chunk_ids (in checkpoint order, duplicates dropped)
    into labelled passages for the prompt: label A, B, … plus doc_id, the
    chunk text, and _CONTEXT_CHARS of source text before and after it
    (sliced at start_offset; whitespace collapsed in all three).

    Raises ValueError naming the claim when it has more than
    _MAX_CHUNKS_PER_CLAIM chunks (never truncates) or when a chunk_id is not
    in the rebuilt index (the golden set was built with non-default chunking
    or different PDFs, so context would be sliced from the wrong place).
    """
    unique = list(dict.fromkeys(chunk_ids))
    if len(unique) > _MAX_CHUNKS_PER_CLAIM:
        raise ValueError(f"claim {claim_id}: {len(unique)} found chunks exceeds "
                         f"_MAX_CHUNKS_PER_CLAIM ({_MAX_CHUNKS_PER_CLAIM}) — refusing to truncate")
    missing = [c for c in unique if c not in chunks_by_id]
    if missing:
        raise ValueError(f"claim {claim_id}: chunk_id {missing[0]!r} is not in the chunk index "
                         f"rebuilt from source_folder (non-default chunking or different PDFs?)")
    bundle = []
    for i, chunk_id in enumerate(unique):
        chunk = chunks_by_id[chunk_id]
        doc = documents[chunk["doc_id"]]
        start, end = chunk["start_offset"], chunk["start_offset"] + len(chunk["chunk_text"])
        bundle.append({
            "label": _label(i),
            "chunk_id": chunk_id,
            "doc_id": chunk["doc_id"],
            "text": _collapse_ws(chunk["chunk_text"]),
            "before": _collapse_ws(doc[max(0, start - _CONTEXT_CHARS):start]),
            "after": _collapse_ws(doc[end:end + _CONTEXT_CHARS]),
        })
    return bundle


def _build_bundle_prompt(claim_text: str, entity: str, bundle: list[dict]) -> str:
    """Fill the frozen _BUNDLE_PROMPT for one claim. Passages are shown by
    label only (never chunk_id) — shorter, and nothing long for the model to
    echo back wrongly. The three slots are filled in one pass, so a value that
    contains a slot token (a claim quoting "__PASSAGES__") is inserted as-is,
    never filled again."""
    passages = "\n\n".join(
        f"[{b['label']}] document: {b['doc_id']}\n"
        f"  before: {b['before']}\n"
        f"  PASSAGE: {b['text']}\n"
        f"  after: {b['after']}"
        for b in bundle
    )
    values = {"__ENTITY__": entity, "__CLAIM__": claim_text, "__PASSAGES__": passages}
    return re.sub("__(?:ENTITY|CLAIM|PASSAGES)__", lambda m: values[m.group(0)], _BUNDLE_PROMPT)


# %%
def _validate_bundle_answer(parsed, labels: list[str]) -> tuple[str, list[str], str]:
    """
    Check one parsed LLM answer against the bundle it was asked about.
    Returns (verdict, needed_labels, reason) with the verdict lower-cased,
    labels upper-cased and de-duplicated in answer order. Raises ValueError
    naming the first problem: not a JSON object; verdict outside _VERDICTS;
    `needed` not a list; a label not in this bundle; not_supported with
    passages listed, or another verdict with none; empty reason.
    """
    if not isinstance(parsed, dict):
        raise ValueError(f"answer is {type(parsed).__name__}, not a JSON object")
    verdict = str(parsed.get("verdict", "")).strip().lower()
    if verdict not in _VERDICTS:
        raise ValueError(f"verdict {verdict!r} is not one of {_VERDICTS}")
    raw_needed = parsed.get("needed")
    if not isinstance(raw_needed, list):
        raise ValueError(f"needed must be a list, got {type(raw_needed).__name__}")
    needed = list(dict.fromkeys(str(x).strip().upper() for x in raw_needed))
    unknown = [x for x in needed if x not in labels]
    if unknown:
        raise ValueError(f"unknown passage label(s) {unknown} (bundle has {labels})")
    if verdict == "not_supported" and needed:
        raise ValueError(f"not_supported must list no passages, got {needed}")
    if verdict != "not_supported" and not needed:
        raise ValueError(f"{verdict} must list at least one passage")
    reason = parsed.get("reason")
    if not isinstance(reason, str) or not reason.strip():
        raise ValueError("reason is missing or empty")
    return verdict, needed, reason.strip()


def draft_claim(claim_text: str, entity: str, bundle: list[dict], llm_client: LLMClient | None) -> dict:
    """
    Draft one claim's verdict over its whole bundle.

    An empty bundle (no found chunks) makes no LLM call and returns
    not_supported with _AUTO_NOT_FOUND_REASON — for that case only,
    `llm_client` may be None (run_draft's caller builds no client at all when
    no claim has a bundle). Otherwise one call with the frozen prompt; an
    answer that fails to arrive, parse or validate is retried exactly once;
    if that fails too the verdict is _DRAFT_FAILED and `reason` carries both
    failures, so a human sees it in the review sheet. Never raises for an LLM
    or answer problem.

    Model, endpoint and temperature come from `llm_client` and `call_llm`
    (temperature is call_llm's 0.0 default); only max_tokens is set here, to
    _BUNDLE_MAX_TOKENS. Log lines name the claim by its first 60 characters.

    Returns {"verdict", "needed_chunk_ids", "reason", "attempts"}.
    """
    if not bundle:
        return {"verdict": "not_supported", "needed_chunk_ids": [],
                "reason": _AUTO_NOT_FOUND_REASON, "attempts": 0}
    if llm_client is None:
        raise RuntimeError(f"draft_claim {claim_text[:60]!r}: bundle has {len(bundle)} chunk(s) but no LLM "
                            f"client is configured")
    prompt = _build_bundle_prompt(claim_text, entity, bundle)
    labels = [b["label"] for b in bundle]
    chunk_id_by_label = {b["label"]: b["chunk_id"] for b in bundle}
    snippet = claim_text[:60]
    errors = []
    for attempt in (1, 2):
        try:
            parsed = _call_llm_with_json_retry(prompt, llm_client, context=f"draft_claim {snippet!r}",
                                               max_attempts=1, max_tokens=_BUNDLE_MAX_TOKENS)
            verdict, needed, reason = _validate_bundle_answer(parsed, labels)
        except Exception as exc:  # noqa: BLE001 — one claim's failure never stops a memo
            errors.append(f"attempt {attempt}: {exc}")
            logger.warning("[draft_claim %r] attempt %d/2 failed: %s", snippet, attempt, exc)
            continue
        return {"verdict": verdict, "needed_chunk_ids": [chunk_id_by_label[x] for x in needed],
                "reason": reason, "attempts": attempt}
    return {"verdict": _DRAFT_FAILED, "needed_chunk_ids": [],
            "reason": "DRAFT FAILED: " + " | ".join(errors), "attempts": 2}


# %%
class DraftRefused(ValueError):
    """`draft` refused before any LLM call (spec §5.2, §12). Nothing has been
    spent; `_main` turns it into a SystemExit carrying the message."""


# The checkpoint columns a review sheet carries per chunk row (hidden), so the
# sheet is self-contained — finalize never re-reads the checkpoint (spec §6).
_CHECKPOINT_FIELDS = ("doc_id", "chunk_id", "chunk_text", "bm25_score", "evidence_span",
                      "found", "confidence", "ambiguous_match", "verbatim_match")


def _review_path(review_dir: str, memo_id: str) -> str:
    return os.path.join(review_dir, f"{memo_id}.xlsx")


def _prepare_memo_claims(df: pd.DataFrame, memo_id: str, claims_dir: str) -> list[dict]:
    """prepare_draft's step 3 for one memo that will be drafted: refuse build-error
    rows and a claims file that disagrees with the checkpoint, then read the PDFs
    and build every claim's bundle. Returns the memo's "claims" list."""
    memo = df[df["memo_id"].eq(memo_id)]
    err = memo["found"].isna() | memo["confidence"].eq("error")
    if err.any():
        raise DraftRefused(f"{memo_id}: {int(err.sum())} build-error row(s) (found empty / confidence "
                           f"'error') — re-run `python golden_set_pipeline.py build` for this memo")

    try:
        source_folder, file_claims = read_claims_file(memo_id, claims_dir)
    except ValueError as exc:
        raise DraftRefused(f"{memo_id}: {exc}") from exc
    file_ids = [claim_id for _, _, claim_id in file_claims]
    ckpt_ids = list(dict.fromkeys(memo["claim_id"]))
    only_file = [c for c in file_ids if c not in set(ckpt_ids)]
    only_ckpt = [c for c in ckpt_ids if c not in set(file_ids)]
    if only_file or only_ckpt:
        raise DraftRefused(f"{memo_id}: claims file and checkpoint disagree — "
                           f"{len(only_file)} claim(s) only in the claims file (first: {only_file[:1]}), "
                           f"{len(only_ckpt)} only in the checkpoint (first: {only_ckpt[:1]}); "
                           f"an edited claim needs `build` re-run before drafting")
    # The same ids are not enough: the bundle is judged against the claims
    # file's text, so the checkpoint rows under an id must be about that text.
    path = os.path.join(claims_dir, f"{memo_id}.md")
    file_by_id = {claim_id: (section, text) for section, text, claim_id in file_claims}
    for row in memo[["claim_id", "section", "claim_text"]].drop_duplicates().itertuples(index=False):
        section, text = file_by_id[row.claim_id]
        for field, value, expected in (("section", row.section, section), ("claim_text", row.claim_text, text)):
            if value != expected:
                raise DraftRefused(f"{memo_id}: claim {row.claim_id}: the checkpoint's {field} {value!r} is not "
                                   f"{path}'s {expected!r} (an edited checkpoint?) — re-run "
                                   f"`python golden_set_pipeline.py build`")

    try:
        chunks_by_id, documents = load_memo_chunks(memo_id, source_folder)
    except ValueError as exc:
        raise DraftRefused(f"{memo_id}: {exc}") from exc
    sources = _source_docs(documents)

    found_rows = memo[memo["found"].eq(True)]
    claims = []
    for number, (section, claim_text, claim_id) in enumerate(file_claims, start=1):
        rows = found_rows[found_rows["claim_id"].eq(claim_id)]       # one per chunk: prepare_draft refused repeats
        row_dicts = rows[list(_CHECKPOINT_FIELDS)].to_dict("records")
        for row in row_dicts:
            # An id that exists in both indexes can still name different text
            # (non-default chunk_size/overlap, or an edited PDF). Refuse, so the
            # draft never judges one passage while the sheet carries another.
            # A missing id is left to build_bundle, which names it.
            chunk = chunks_by_id.get(row["chunk_id"])
            if chunk is not None and (chunk["doc_id"], chunk["chunk_text"]) != (row["doc_id"], row["chunk_text"]):
                what = "document" if chunk["doc_id"] != row["doc_id"] else "text"
                raise DraftRefused(f"{memo_id}: claim {claim_id}: chunk {row['chunk_id']!r} rebuilt from "
                                   f"source_folder has a different {what} than the checkpoint "
                                   f"(non-default chunking or different PDFs?) — rebuild the golden set")
        try:
            bundle = build_bundle(claim_id, [r["chunk_id"] for r in row_dicts], chunks_by_id, documents)
        except ValueError as exc:
            raise DraftRefused(f"{memo_id}: {exc}") from exc
        claims.append({"number": number, "claim_id": claim_id, "section": section,
                       "claim_text": claim_text, "rows": row_dicts, "bundle": bundle, "source_docs": sources})
    return claims


def prepare_draft(memo_ids: list[str],
                  checkpoint_path: str = "golden_set_checkpoint.parquet",
                  claims_dir: str = "claims",
                  review_dir: str = "review") -> list[dict]:
    """
    Everything `draft` does before spending: every check and every bundle,
    for every requested memo. Takes no LLM client on purpose — it cannot make
    a call, so "no credentials and no spend until every memo passes" holds by
    construction (spec §5.2, §12). Raises DraftRefused naming the memo on the
    first problem.

    memo_ids: memos to draft, in order; empty means every memo in the
    checkpoint (checkpoint order).

    Order, and why:
    1. Checkpoint-only checks run for EVERY requested memo, even one whose
       review sheet already exists, so a checkpoint fault always surfaces:
       the memo is in the checkpoint; every claim_id is
       a deterministic uuid5; every non-empty found is TRUE/FALSE (text
       like 'FALSE' would silently drop evidence); no claim has more than _MAX_CHUNKS_PER_CLAIM
       found chunks (never truncated); no (claim_id, chunk_id) has two found
       rows (a checkpoint from before _rows_for_claim's dedupe — keeping one
       would guess its confidence and span); review/<memo_id>.xlsx, if it
       exists, is a file (a folder of that name would otherwise be skipped
       forever). Before all of these, `review_dir` must be a folder if it
       exists: run_draft creates it only after the LLM calls are spent.
    2. An existing review/<memo_id>.xlsx marks the memo "skip". Its claims
       file and PDFs are never read — nothing from them would be used.
    3. Only for memos to draft:
       - first, for every such memo before any memo's PDFs are read: its
         claims file parses and sets `filing_entity` (_filing_entity);
       - a build-error row (found empty / confidence "error") refuses the
         memo: re-run build. It is never drafted as "not supported", which
         would give a build failure the "auto: found=False" prefix the eval
         reads as "search found nothing";
       - the claims file must hold exactly the checkpoint's claim ids (an
         edited claim changes its id, and the sheet's order comes from the
         file), and every checkpoint row's section and claim_text must equal
         the file's for its id (an edited checkpoint);
       - every found chunk_id must be in the chunk index rebuilt from the
         claims file's source_folder, with the checkpoint's doc_id and
         chunk_text (the PDFs are read here, and only here).

    Returns one dict per memo:
    {"memo_id", "action": "draft" | "skip", "review_path", "entity" (None for a skip),
     "claims": [{"number", "claim_id", "section", "claim_text", "rows", "bundle", "source_docs"}, ...]}
    "claims" is [] for a skip. "rows" holds one checkpoint row dict
    (_CHECKPOINT_FIELDS) per bundle passage, aligned with "bundle" by
    position; both are [] for a claim with no found rows.
    """
    if not os.path.exists(checkpoint_path):
        raise DraftRefused(f"{checkpoint_path} not found — run `python golden_set_pipeline.py build` first")
    # lexists, not exists: a symlink to nothing is not an absent path (exists
    # is False for it), and makedirs / the publish would fail on it later.
    if os.path.lexists(review_dir) and not os.path.isdir(review_dir):
        raise DraftRefused(f"{review_dir} exists but is not a folder — move it away; draft writes review sheets there")
    try:
        df = pd.read_parquet(checkpoint_path).reset_index(drop=True)
    except (OSError, ValueError) as exc:          # pyarrow's ArrowInvalid is a ValueError
        raise DraftRefused(f"{checkpoint_path} could not be read ({type(exc).__name__}: {exc}) — re-run "
                           f"`python golden_set_pipeline.py build`") from exc
    missing = [c for c in ("memo_id", "claim_id", "section", "claim_text", *_CHECKPOINT_FIELDS) if c not in df.columns]
    if missing:
        raise DraftRefused(f"{checkpoint_path} is missing column(s) {missing} — re-run "
                           f"`python golden_set_pipeline.py build`")
    in_checkpoint = list(dict.fromkeys(df["memo_id"]))
    requested = list(dict.fromkeys(memo_ids)) or in_checkpoint

    for memo_id in requested:
        if memo_id not in in_checkpoint:
            raise DraftRefused(f"{memo_id}: not in {checkpoint_path} (memos there: {in_checkpoint})")
        sheet = _review_path(review_dir, memo_id)
        if os.path.lexists(sheet) and not os.path.isfile(sheet):
            raise DraftRefused(f"{memo_id}: {sheet} exists but is not a file — remove it, then re-run")

    found = df["found"].eq(True)
    for memo_id in requested:
        memo = df["memo_id"].eq(memo_id)
        bad = [c for c in df.loc[memo, "claim_id"].unique() if not _is_deterministic_claim_id(c)]
        if bad:
            raise DraftRefused(f"{memo_id}: {len(bad)} claim_id(s) are not deterministic uuid5 "
                               f"(pre-decision-17 golden set) — rebuild with `build`: {bad[:3]}")
        # found must be a real boolean (empty is a build-error row, refused later):
        # text such as 'FALSE' would fail .eq(True) and silently drop the evidence.
        found_values = df.loc[memo, "found"]
        not_bool = found_values.notna() & ~found_values.isin([True, False])
        if not_bool.any():
            raise DraftRefused(f"{memo_id}: {int(not_bool.sum())} checkpoint row(s) have a found value that is not "
                               f"TRUE/FALSE (first: {found_values[not_bool].iloc[0]!r}) — re-run "
                               f"`python golden_set_pipeline.py build`")
        repeated = df[memo & found & df.duplicated(subset=["claim_id", "chunk_id"], keep=False)]
        if len(repeated):
            first = repeated.iloc[0]
            n = int((repeated["claim_id"].eq(first["claim_id"]) & repeated["chunk_id"].eq(first["chunk_id"])).sum())
            raise DraftRefused(f"{memo_id}: claim {first['claim_id']}: chunk {first['chunk_id']!r} appears in {n} "
                               f"found rows (a checkpoint built before same-chunk matches were collapsed) — "
                               f"re-run `python golden_set_pipeline.py build`")
        per_claim = df[memo & found].groupby("claim_id", sort=False)["chunk_id"].nunique()
        over = per_claim[per_claim > _MAX_CHUNKS_PER_CLAIM]
        if len(over):
            raise DraftRefused(f"{memo_id}: claim {over.index[0]} has {int(over.iloc[0])} found chunks, "
                               f"over _MAX_CHUNKS_PER_CLAIM ({_MAX_CHUNKS_PER_CLAIM}) — refusing to truncate")

    to_draft = [m for m in requested if not os.path.exists(_review_path(review_dir, m))]
    entities = {memo_id: _filing_entity(memo_id, claims_dir) for memo_id in to_draft}

    plan = []
    for memo_id in requested:
        entry = {"memo_id": memo_id, "review_path": _review_path(review_dir, memo_id),
                 "entity": entities.get(memo_id)}
        if memo_id not in entities:
            plan.append({**entry, "action": "skip", "claims": []})
            continue
        plan.append({**entry, "action": "draft", "claims": _prepare_memo_claims(df, memo_id, claims_dir)})
    return plan


# %%
# The review sheet (spec §6). Plain words for the verdicts a human reads and
# picks; the hidden ai_verdict column keeps the code form.
_VERDICT_WORDS = {"stated_directly": "stated directly", "needs_combining": "needs combining",
                  "not_supported": "not supported", _DRAFT_FAILED: "draft failed"}

# (key, header, width) of the visible columns, left to right.
_SHEET_VISIBLE = (
    ("block", "block / letter", 8),
    ("text", "text", 70),
    ("document", "document", 22),
    ("before", "text just before", 40),
    ("after", "text just after", 40),
    ("quoted_span", "quoted span (from build)", 40),
    ("verdict_needed", "VERDICT / NEEDED?", 20),
    ("ai_reason", "AI reason", 50),
    ("checked", "CHECKED", 10),
    ("note", "note", 30),
)
# Hidden columns make the sheet self-contained. row_kind is claim / chunk / add.
# chunk_count (claim rows only) is how many chunk rows draft wrote under the
# claim, so finalize can tell a deleted LAST chunk row from a shorter bundle.
# source_docs (claim rows only) names the memo's PDFs draft read, each with a
# fingerprint of its extracted text (_source_docs), so finalize can tell that
# source_folder's PDFs changed since drafting — added, removed or edited.
_SHEET_HIDDEN = (("row_kind", "claim_id", "memo_id", "section", "claim_text")
                 + _CHECKPOINT_FIELDS + ("ai_verdict", "ai_needed", "chunk_count", "source_docs"))
# Hidden columns a sheet drafted before they were added lacks (MEMO-004's was
# drafted without either); read_review_sheet accepts a sheet without them.
_SHEET_OPTIONAL = ("chunk_count", "source_docs")
# THE column order, shared with finalize's reader (step-5 Plan B) — one list,
# not two copies kept identical by hand. Row 1 holds _SHEET_HEADERS[key].
_SHEET_COLUMNS = tuple(key for key, _, _ in _SHEET_VISIBLE) + _SHEET_HIDDEN
_SHEET_HEADERS = {**{key: header for key, header, _ in _SHEET_VISIBLE},
                  **{key: key for key in _SHEET_HIDDEN}}

_CLAIM_FILL = "DDEBF7"   # light blue claim rows

# Spec Appendix B, verbatim; {memo_id} is filled with str.replace.
_HOW_TO_TEXT = """TAG REVIEW — {memo_id}

Each blue row is a claim. The rows under it are the pieces of source text
search found for it. The AI has already filled in an answer. Check it and
change anything wrong.

1. VERDICT (blue claim row):
   stated directly  - at least one piece says the claim on its own
   needs combining  - true only by putting pieces together, by calculating,
                      or by describing numbers in a word the text doesn't use
   not supported    - these pieces, even together, don't establish the claim

2. NEEDED? (each piece row) - 'yes' on:
   stated directly  -> the piece(s) that say it on their own
   needs combining  -> every piece the combination uses
   not supported    -> nothing
   Mark a piece only when ITS OWN text carries what the claim needs. The grey
   columns show text just before/after it (often table headings).

3. MISSING A PIECE? Paste a quote (25+ characters) from the PDF into an
   empty '+' row under the claim and set NEEDED? to 'yes'.

4. CHECKED (blue claim row): set 'yes' once you have looked at the claim -
   also when you agree with the AI. Claims without it are refused at finalize.

5. DRAFT FAILED? The AI could not answer that claim. Pick the verdict
   yourself and mark the pieces that support it (mark nothing only for
   'not supported').

Save with the SAME file name."""


def _plain(value):
    """A checkpoint value as openpyxl writes it cleanly: None / NaN / "" become
    an empty cell, numpy scalars (parquet bools read back as numpy.bool)
    become Python scalars, and control characters Excel cannot store are
    removed from text — pypdf emits NUL bytes (242 in one real memo's PDFs, 42 in
    another's) and openpyxl raises on them. The sheet's hidden chunk_text can
    therefore differ from the checkpoint's by those bytes; chunk_id is the key."""
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return None
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    return None if value == "" else value


def _claim_block(memo_id: str, claim: dict) -> list[dict]:
    """One claim's sheet rows as {column key: value}: the claim row, one row per
    bundle passage (AI marks pre-filled), then two empty add-a-piece rows.
    Every row carries claim_id, so a reader never depends on row position."""
    draft = claim["draft"]
    needed = set(draft["needed_chunk_ids"])
    base = {"claim_id": claim["claim_id"], "memo_id": memo_id,
            "section": claim["section"], "claim_text": claim["claim_text"]}
    rows = [{**base, "row_kind": "claim", "block": claim["number"], "text": claim["claim_text"],
             "verdict_needed": _VERDICT_WORDS[draft["verdict"]], "ai_reason": draft["reason"],
             "ai_verdict": draft["verdict"], "chunk_count": len(claim["bundle"]),
             "source_docs": claim["source_docs"]}]
    for passage, ckpt in zip(claim["bundle"], claim["rows"], strict=True):
        mark = "yes" if passage["chunk_id"] in needed else None
        rows.append({**base, **ckpt, "row_kind": "chunk", "block": passage["label"],
                     "text": passage["text"], "document": passage["doc_id"],
                     "before": passage["before"], "after": passage["after"],
                     "quoted_span": ckpt["evidence_span"], "verdict_needed": mark, "ai_needed": mark})
    rows += [{**base, "row_kind": "add", "block": "+"} for _ in range(2)]
    return rows


def write_review_sheet(path: str, memo_id: str, claims: list[dict]) -> None:
    """
    Write one memo's review sheet (spec §6): a "how to" tab (Appendix B) and a
    "bundles" tab — per claim a blue claim row, one row per passage, two
    add-a-piece rows — with the AI draft pre-filled in VERDICT / NEEDED?.
    `claims` is prepare_draft's "claims" list, each item with a "draft" key
    holding draft_claim's return.

    Columns follow _SHEET_COLUMNS: visible first, then hidden. Verdicts show
    as plain words; `draft failed` is written as-is, outside the dropdown
    list, so finalize refuses it. Text starting with "=" is stored as text,
    never as a formula.

    Never overwrites: FileExistsError if `path` exists, checked first and
    again at publish. Atomic: saved to a unique temp file next to `path`, then
    published with os.link, which fails if `path` exists — unlike
    write_claims_file's os.replace, a sheet another run published in the
    meantime is never replaced. A failed save leaves neither file.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    if os.path.exists(path):
        raise FileExistsError(f"{path} already exists — a review sheet is never overwritten")

    wb = Workbook()
    how_to = wb.active
    if how_to is None:
        # openpyxl returns None only for a workbook with no sheets at all,
        # which a freshly constructed Workbook() can't produce (it always
        # starts with one active sheet).
        raise ValueError(f"{path}: new workbook has no active worksheet")
    how_to.title = "how to"
    for i, line in enumerate(_HOW_TO_TEXT.replace("{memo_id}", memo_id).split("\n"), start=1):
        how_to.cell(row=i, column=1, value=line or None)
    how_to.column_dimensions["A"].width = 90

    ws = wb.create_sheet("bundles")
    col = {key: i for i, key in enumerate(_SHEET_COLUMNS, start=1)}
    letter = {key: get_column_letter(i) for key, i in col.items()}
    for key, i in col.items():
        ws.cell(row=1, column=i, value=_SHEET_HEADERS[key]).font = Font(bold=True)
    for key, _, width in _SHEET_VISIBLE:
        ws.column_dimensions[letter[key]].width = width
    for key in _SHEET_HIDDEN:
        ws.column_dimensions[letter[key]].hidden = True
    ws.freeze_panes = "C2"

    verdict_list = DataValidation(type="list", allow_blank=True,
                                  formula1='"' + ",".join(_VERDICT_WORDS[v] for v in _VERDICTS) + '"')
    yes_list = DataValidation(type="list", allow_blank=True, formula1='"yes"')
    ws.add_data_validation(verdict_list)
    ws.add_data_validation(yes_list)
    claim_fill = PatternFill("solid", fgColor=_CLAIM_FILL)
    grey = Font(color="808080")
    wrap = Alignment(wrap_text=True, vertical="top")

    r = 1
    for claim in claims:
        for row in _claim_block(memo_id, claim):
            r += 1
            for key, value in row.items():
                cell = ws.cell(row=r, column=col[key], value=_plain(value))
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"          # PDF text, never a formula
            for key, _, _ in _SHEET_VISIBLE:
                cell = ws.cell(row=r, column=col[key])
                cell.alignment = wrap
                if row["row_kind"] == "claim":
                    cell.fill = claim_fill
                if key in ("before", "after"):
                    cell.font = grey
            if row["row_kind"] == "claim":
                verdict_list.add(f"{letter['verdict_needed']}{r}")
                yes_list.add(f"{letter['checked']}{r}")
            else:
                yes_list.add(f"{letter['verdict_needed']}{r}")

    fd, tmp = tempfile.mkstemp(dir=os.path.dirname(path) or ".",
                               prefix=os.path.basename(path) + ".", suffix=".tmp")
    os.close(fd)
    try:
        wb.save(tmp)
        os.link(tmp, path)        # FileExistsError if another run published first
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)        # after a link, `path` keeps the saved workbook


# %%
def run_draft(plan: list[dict], llm_client: LLMClient | None) -> dict[str, int]:
    """
    Spend: draft every claim of every "draft" memo in prepare_draft's plan, in
    claims-file order, and write its review sheet (spec §5.4). Per memo:
    - written — at least one LLM call gave a valid draft, or the memo needed
      no call at all; the sheet is saved, `draft failed` claims visible inline;
    - skipped — the plan says so (its review sheet already exists; no calls),
      or the sheet appeared while this memo was drafting (another run): that
      sheet is left untouched and this run's drafts are discarded;
    - failed — the memo made LLM calls and every one came back `draft failed`
      (an outage, a bad key), or its sheet could not be saved (disk full,
      permissions). Claims that needed no call do not count as successes
      here: otherwise almost every memo would be "written" during an outage,
      and a re-run would then skip it. No file, so a re-run simply retries.
    A failed or skipped memo never stops the memos after it.
    Logs one line per claim (number and verdict, no claim text) and one
    summary line per memo. `llm_client` may be None only when no "draft" memo
    has a claim with a non-empty bundle. Returns {"written", "skipped", "failed"}.
    """
    counts = {"written": 0, "skipped": 0, "failed": 0}
    for memo in plan:
        memo_id = memo["memo_id"]
        if memo["action"] == "skip":
            counts["skipped"] += 1
            logger.info("%s: skipped (%s already exists)", memo_id, memo["review_path"])
            continue
        drafted = []
        total = len(memo["claims"])
        for claim in memo["claims"]:
            draft = draft_claim(claim["claim_text"], memo["entity"], claim["bundle"], llm_client)
            drafted.append({**claim, "draft": draft})
            logger.info("%s: claim %d/%d -> %s", memo_id, claim["number"], total, draft["verdict"])
        n_failed = sum(c["draft"]["verdict"] == _DRAFT_FAILED for c in drafted)
        n_no_pieces = sum(not c["bundle"] for c in drafted)
        n_called = len(drafted) - n_no_pieces
        n_drafted = n_called - n_failed
        if n_called and n_failed == n_called:
            counts["failed"] += 1
            logger.error("%s: failed (all %d LLM calls came back draft failed) — no file written",
                         memo_id, n_failed)
            continue
        try:
            # Separate from the save: makedirs raises FileExistsError when
            # `review` is a regular file, which is a failure, not a sheet
            # that appeared.
            os.makedirs(os.path.dirname(memo["review_path"]) or ".", exist_ok=True)
        except OSError as exc:
            counts["failed"] += 1
            logger.error("%s: failed — could not create the folder for %s: %s", memo_id, memo["review_path"], exc)
            continue
        try:
            write_review_sheet(memo["review_path"], memo_id, drafted)
        except FileExistsError:
            counts["skipped"] += 1
            logger.warning("%s: skipped — %s appeared while drafting (another run?); it was left "
                           "untouched and this run's drafts were discarded", memo_id, memo["review_path"])
            continue
        except OSError as exc:
            counts["failed"] += 1
            logger.error("%s: failed — could not save %s: %s", memo_id, memo["review_path"], exc)
            continue
        counts["written"] += 1
        logger.info("%s: written (%d drafted, %d with no pieces, %d draft failed) -> %s",
                    memo_id, n_drafted, n_no_pieces, n_failed, memo["review_path"])
    return counts


# %% [markdown]
# ## 3. finalize — a checked review sheet becomes the eval's input
#
# Spec §7 as refined by §13 (2026-09-14). Reads review/<memo_id>.xlsx, the
# claims file and the source PDFs; never the checkpoint, never an LLM.

# %%
class FinalizeRefused(ValueError):
    """`finalize` could not read what it checks: no sheet, a numbered copy next
    to it, no bundles tab or a missing header, no claims file or PDFs. Nothing
    is written; `_main` turns it into a SystemExit carrying the message."""


def read_review_sheet(path: str) -> list[dict]:
    """
    Read the "bundles" tab of a review sheet: one {column key: value} dict
    per non-empty row, plus "sheet_row" (the row number Excel shows). Reads
    only; judges nothing (that is check_review).

    Built for sheets that came back through Excel or Apple Numbers (spec
    §13.1): the tab is found by name (Numbers inserts an "Export Summary"
    sheet first), columns by their row-1 header (position and extra columns
    do not matter), and hidden state is ignored (Numbers un-hides columns).

    Raises FinalizeRefused if the file is not a readable workbook, has no
    "bundles" tab, lacks one of _SHEET_HEADERS other than _SHEET_OPTIONAL's,
    or repeats one. A missing optional column is simply absent from every
    record's keys.
    """
    from xml.etree.ElementTree import ParseError
    from zipfile import BadZipFile

    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        wb = load_workbook(path)
    except (InvalidFileException, BadZipFile, KeyError, OSError, ParseError) as exc:   # ParseError: damaged XML inside the zip
        raise FinalizeRefused(f"{path} is not a readable .xlsx workbook: {exc}") from exc
    if "bundles" not in wb.sheetnames:
        raise FinalizeRefused(f"{path} has no 'bundles' tab (tabs: {wb.sheetnames})")
    ws = wb["bundles"]
    key_by_header = {header: key for key, header in _SHEET_HEADERS.items()}
    col = {}
    for i, cell in enumerate(ws[1]):
        key = key_by_header.get(cell.value) if isinstance(cell.value, str) else None
        if key is None:
            continue
        if key in col:
            raise FinalizeRefused(f"{path}: column header {cell.value!r} appears twice")
        col[key] = i
    missing = [_SHEET_HEADERS[key] for key in _SHEET_COLUMNS if key not in col and key not in _SHEET_OPTIONAL]
    if missing:
        raise FinalizeRefused(f"{path}: the 'bundles' tab is missing column header(s) {missing}")
    records = []
    for number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        record = {key: (row[i] if i < len(row) else None) for key, i in col.items()}
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in record.values()):
            continue
        records.append({**record, "sheet_row": number})
    return records


# %%
# Shortest pasted quote accepted: anything shorter can be a bare number or a
# generic phrase that matches in many places (spec §8).
_MIN_QUOTE_CHARS = 25


def _squeeze(text: str) -> str:
    """`text` with every whitespace character removed and the control
    characters Excel cannot store stripped — the key a pasted quote is
    matched on (spec §13.4)."""
    return "".join(ILLEGAL_CHARACTERS_RE.sub("", text).split())


def _resolve_quote(quote: str, chunks_by_id: dict[str, dict]) -> list[str] | str:
    """
    Find the chunk(s) of the memo a pasted quote comes from (spec §13.4).
    Returns a list of chunk ids, or a problem sentence for the user.

    - Containing _QUOTE_SEPARATOR once stored (whitespace collapsed), or
      starting or ending with its bar -> problem: finalize joins one chunk's
      quotes with it, and the eval splits them there again, so such a quote
      would be scored as fragments.
    - Shorter than _MIN_QUOTE_CHARS after whitespace collapse -> problem.
    - Matching removes ALL whitespace from both sides, not just collapses
      it: pypdf sometimes drops the space between table cells ("EUR4,203m"),
      which a collapsed quote could never equal. No ligature or hyphen
      folding — a quote copied from a PDF viewer that differs that way is
      reported as not found.
    - Exactly one chunk contains it -> that chunk.
    - Exactly two chunks of the same document, next to each other
      (_chunk_index_of differs by 1) -> both, in document order: the quote
      sits in their shared overlap, one passage (decision 1). A quote in
      one place sits in at most two chunks, since the 200-character overlap
      is shorter than the 800-character step, so three or more hits always
      mean several places. Residual: a quote genuinely repeated in two
      neighbouring chunks also lands here and over-marks one of them; the
      finalize summary labels this case so the user can look.
    - None, or any other set -> problem.
    """
    # Tested on the text as finalize stores it (whitespace collapsed), padded
    # so a leading "| " or trailing " |" — which joining would turn into a
    # separator — is caught too.
    if _QUOTE_SEPARATOR in f" {_collapse_ws(quote)} ":
        return f"quote contains {_QUOTE_SEPARATOR.strip()!r} between spaces or at an end — paste the parts on " \
               f"either side as separate quotes"
    stripped = ILLEGAL_CHARACTERS_RE.sub("", quote)
    collapsed = _collapse_ws(stripped)
    if len(collapsed) < _MIN_QUOTE_CHARS:
        return f"quote is {len(collapsed)} characters — paste at least {_MIN_QUOTE_CHARS}"
    key = "".join(stripped.split())                     # _squeeze, reusing the strip above
    hits = [chunk_id for chunk_id, chunk in chunks_by_id.items() if key in _squeeze(chunk["chunk_text"])]
    if len(hits) == 1:
        return hits
    if len(hits) == 2:
        first, second = (_chunk_index_of(h) for h in hits)
        same_doc = chunks_by_id[hits[0]]["doc_id"] == chunks_by_id[hits[1]]["doc_id"]
        if same_doc and first is not None and second is not None and abs(first - second) == 1:
            return [hits[0], hits[1]] if first < second else [hits[1], hits[0]]
    if not hits:
        return ("quote not found in the PDFs' extracted text — it may run across two pieces, or the PDF "
                "viewer's copy differs (ligatures, hyphenation); paste a shorter part")
    documents = sorted({chunks_by_id[h]["doc_id"] for h in hits})
    return f"quote found in {len(hits)} places ({', '.join(documents)}) — paste a longer quote"


# %%
# A verdict word as the user picks it in the sheet -> its code form.
_VERDICT_BY_WORD = {_VERDICT_WORDS[v]: v for v in _VERDICTS}


def _answer(value) -> str:
    """A typed cell as finalize compares it: stripped and case-folded, so
    'Yes ' equals 'yes' (Excel autocorrect capitalises a cell's first letter).
    An empty cell is ''."""
    return "" if value is None else str(value).strip().casefold()


def check_review(memo_id: str, records: list[dict], file_claims: list[tuple[str, str, str]],
                 chunks_by_id: dict[str, dict], documents: dict[str, str]) -> tuple[list[str], list[dict]]:
    """
    Every check finalize makes on a read sheet, in one pass (spec §13.3):
    returns (problems, claims). `problems` lists every problem found, each
    naming the sheet row and claim block; any problem means nothing may be
    written. `claims` is only meaningful when `problems` is empty.

    Inputs: read_review_sheet's records; read_claims_file's claims for the
    memo; load_memo_chunks' rebuilt index and its documents (doc_id -> text). Writes nothing, raises nothing.

    Why each hidden column is checked rather than trusted: a sheet saved by
    Numbers comes back with its hidden columns visible and editable, so
    claim_id, chunk_id, memo_id and chunk_text are verified against the
    claims file and the PDFs.

    Checks: CHECKED is 'yes'; VERDICT is one of the three words (never
    'draft failed'); marks agree with the verdict (a '+' row with a quote and
    'yes' counts as a mark even if its quote fails to resolve); NEEDED? is
    'yes' or blank; the claim rows' source_docs equal the PDFs source_folder
    gives now (_source_docs; skipped on a sheet without the column); the
    sheet's claims and the claims file's claims are the same set, both
    directions; a claim row's visible text and hidden
    claim_text and section equal the claims file's for its claim_id (the
    census alone passes ids swapped between two whole blocks); a '+' row has both a quote and 'yes' or
    neither; row_kind is known; every row after a claim row carries that
    claim's claim_id and this memo_id; no chunk appears twice in one claim;
    every chunk row's chunk_id is in the rebuilt index with the same doc_id
    and chunk_text, and its visible text and document are that chunk's (a
    mark applies to the passage the reviewer read; residual: a row whose
    visible AND hidden cells were all re-pointed to another chunk passes —
    only the checkpoint knows draft's bundle, and finalize does not read it); found/ambiguous_match/verbatim_match are booleans, found
    is TRUE (draft writes chunk rows only for found chunks, and the reviewed
    file writes them found=True), and bm25_score a number or blank; every quote resolves (_resolve_quote); no
    chunk row was deleted — a claim with no chunk rows must carry
    _AUTO_NOT_FOUND_REASON as its AI reason (draft writes no chunk rows only
    then), and a claim's chunk rows must be lettered exactly _label(0),
    _label(1), … in order, and as many as the claim row's chunk_count (a
    whole number). Residual: on a sheet drafted before chunk_count existed
    (no such column), deleting a claim's LAST chunk row (B of A, B) passes;
    run_finalize warns about it.
    'yes' and the verdict words are compared with _answer.

    Each returned claim: {"number", "claim_id", "verdict", "ai_verdict",
    "ai_reason", "note", "plus_notes", "chunks": [{"record", "marked",
    "ai_marked", "by_quote", "quote_notes"}], "added": [{"chunk_id", "quotes",
    "notes", "how"}]} where "by_quote" / "how" is None or "1 chunk" / "2
    adjacent chunks — check for repeated text". No note cell is dropped (spec
    §7.4): "plus_notes" holds notes of '+' rows with no quote (a comment on
    the claim); "quote_notes" / "notes" hold the note of each '+' row whose
    quote marked that bundle chunk / added that chunk.
    """
    problems = []
    blocks = []
    for record in records:
        kind = _answer(record["row_kind"])
        if kind == "claim":
            blocks.append((record, []))
        elif kind not in ("chunk", "add"):
            problems.append(f"row {record['sheet_row']}: unknown row_kind {record['row_kind']!r}")
        elif not blocks:
            problems.append(f"row {record['sheet_row']}: a {kind} row before the first claim row")
        else:
            blocks[-1][1].append(record)

    # A PDF added to (or removed from) source_folder since drafting: chunk rows
    # can all still verify, but a claim drafted with nothing found was never
    # searched in the new PDF. One problem for the memo, not one per claim row.
    sources = _source_docs(documents)
    drifted = [row for row, _ in blocks if "source_docs" in row and (row["source_docs"] or "") != sources]
    if drifted:
        problems.append(f"row {drifted[0]['sheet_row']}: this sheet was drafted from the PDFs "
                        f"{drifted[0]['source_docs']!r}, but claims/{memo_id}.md's source_folder now gives "
                        f"{sources!r} (names and extracted text are compared) — run "
                        f"`python golden_set_pipeline.py build`, then "
                        f"`python tag_pipeline.py draft {memo_id}`, and review again (answers are not carried over)")

    file_ids = {claim_id for _, _, claim_id in file_claims}
    file_by_id = {claim_id: (section, text) for section, text, claim_id in file_claims}
    census = False
    claims, seen_claims = [], set()
    for claim_row, children in blocks:
        number = claim_row["block"]

        def problem(record, text, number=number):
            problems.append(f"row {record['sheet_row']} (claim {number}): {text}")

        claim_id = claim_row["claim_id"]
        for record in (claim_row, *children):
            if record["memo_id"] != memo_id:
                problem(record, f"memo_id {record['memo_id']!r} is not {memo_id} "
                                f"(a row pasted from another memo's sheet?)")
        if not claim_id:
            problem(claim_row, "claim row has no claim_id")
            continue
        if claim_id in seen_claims:
            problem(claim_row, f"claim_id {claim_id} has a second claim row")
            continue
        seen_claims.add(claim_id)
        if claim_id not in file_ids:
            census = True
            problem(claim_row, f"claim_id {claim_id} is not in claims/{memo_id}.md "
                               f"(an edited cell, or a claim reworded after drafting)")
        else:
            # The census alone passes two claims whose ids were swapped across
            # their whole blocks; the claims file, not the sheet, says which text
            # an id belongs to.
            section, text = file_by_id[claim_id]
            for field, expected in (("text", text), ("claim_text", text), ("section", section)):
                if (claim_row[field] or "") != (_plain(expected) or ""):
                    problem(claim_row, f"{field} {claim_row[field]!r} is not claims/{memo_id}.md's {expected!r} "
                                       f"for this claim_id (an edited cell, or claim ids swapped between claims)")
                    break
        if _answer(claim_row["checked"]) != "yes":
            problem(claim_row, "CHECKED is not 'yes'")
        word = _answer(claim_row["verdict_needed"])
        verdict = _VERDICT_BY_WORD.get(word)
        if word == _VERDICT_WORDS[_DRAFT_FAILED]:
            problem(claim_row, "VERDICT is still 'draft failed' — pick a verdict yourself and mark "
                               "the pieces that support it")
        elif verdict is None:
            problem(claim_row, f"VERDICT {claim_row['verdict_needed']!r} is not one of "
                               f"{[_VERDICT_WORDS[v] for v in _VERDICTS]}")

        chunks, quotes, plus_notes, seen_chunks = [], [], [], set()
        for record in children:
            kind = _answer(record["row_kind"])
            if record["claim_id"] != claim_id:
                problem(record, f"claim_id {record['claim_id']!r} differs from its claim row's "
                                f"(an edited cell, or a row moved between claims)")
                continue
            mark = _answer(record["verdict_needed"])
            if mark not in ("", "yes"):
                problem(record, f"NEEDED? {record['verdict_needed']!r} is not 'yes' or blank")
            if kind == "add":
                quote = "" if record["text"] is None else str(record["text"])
                if quote.strip() and mark != "yes":
                    problem(record, "a quote without NEEDED? 'yes' — set 'yes' to add it, or clear the quote")
                elif mark == "yes" and not quote.strip():
                    problem(record, "NEEDED? 'yes' on a '+' row without a quote")
                elif quote.strip():
                    quotes.append((record, quote))
                elif record["note"] is not None and str(record["note"]).strip():
                    plus_notes.append(record["note"])            # a comment on the claim
                continue
            chunk_id = record["chunk_id"]
            if not chunk_id:
                problem(record, "chunk row has no chunk_id")
                continue
            if chunk_id in seen_chunks:
                problem(record, f"chunk {chunk_id!r} appears twice under this claim")
                continue
            seen_chunks.add(chunk_id)
            real = chunks_by_id.get(chunk_id)
            if real is None:
                problem(record, f"chunk_id {chunk_id!r} is not in the chunk index rebuilt from the "
                                f"source PDFs (an edited cell?)")
                continue
            if record["doc_id"] != real["doc_id"]:
                problem(record, f"chunk {chunk_id!r}: document {record['doc_id']!r} differs from the PDFs' "
                                f"{real['doc_id']!r} (an edited cell?)")
            elif (record["chunk_text"] or "") != (_plain(real["chunk_text"]) or ""):
                problem(record, f"chunk {chunk_id!r}: its text differs from the PDFs' (an edited cell?)")
            # The reviewer marks what they read: the visible passage and document
            # must be this chunk's, as draft wrote them (whitespace collapsed).
            elif (record["text"] or "") != (_plain(_collapse_ws(real["chunk_text"])) or ""):
                problem(record, f"chunk {chunk_id!r}: the passage shown in 'text' is not this chunk's text "
                                f"(hidden cells replaced with another chunk's?)")
            elif record["document"] != real["doc_id"]:
                problem(record, f"chunk {chunk_id!r}: the 'document' shown is {record['document']!r}, "
                                f"not {real['doc_id']!r} (an edited cell?)")
            for field in ("found", "ambiguous_match", "verbatim_match"):
                if not isinstance(record[field], bool):
                    problem(record, f"{field} {record[field]!r} is not TRUE/FALSE (an edited cell?)")
            if record["found"] is False:                  # an openpyxl cell value: a plain Python bool
                problem(record, "found is FALSE, but draft writes chunk rows only for found chunks (an edited cell?)")
            score = record["bm25_score"]
            if score is not None and (isinstance(score, bool) or not isinstance(score, (int, float))):
                problem(record, f"bm25_score {score!r} is not a number (an edited cell?)")
            chunks.append({"record": record, "marked": mark == "yes",
                           "ai_marked": _answer(record["ai_needed"]) == "yes", "by_quote": None,
                           "quote_notes": []})

        # finalize never reads the checkpoint, so a deleted chunk row is caught
        # from what draft wrote: a claim has no chunk rows only when search found
        # nothing (that exact AI reason), its chunk rows are lettered A, B, C…,
        # and there are chunk_count of them (the one check that sees the last row)
        chunk_rows = [record for record in children if _answer(record["row_kind"]) == "chunk"]
        letters = [record["block"] for record in chunk_rows]
        expected = [_label(i) for i in range(len(chunk_rows))]
        recovery = (f"restore it from the drafted sheet, or delete review/{memo_id}.xlsx and re-run "
                    f"`python tag_pipeline.py draft {memo_id}`")
        if not chunk_rows and claim_row["ai_reason"] != _AUTO_NOT_FOUND_REASON:
            problem(claim_row, f"no chunk rows, but its AI reason is not '{_AUTO_NOT_FOUND_REASON}', so search "
                               f"found chunks for it — a chunk row was deleted: {recovery}")
        elif letters != expected:
            first = next(record for record, want in zip(chunk_rows, expected) if record["block"] != want)
            problem(first, f"chunk rows are lettered {letters}, not {expected} as draft wrote them — a chunk "
                           f"row was deleted or moved: {recovery}")
        elif "chunk_count" in claim_row:                         # absent on a sheet drafted before it
            count = claim_row["chunk_count"]
            if isinstance(count, bool) or not isinstance(count, (int, float)) or count != int(count):
                problem(claim_row, f"chunk_count {count!r} is not a whole number (an edited cell?)")
            elif len(chunk_rows) != count:
                problem(claim_row, f"{len(chunk_rows)} chunk row(s) under this claim, but draft wrote {int(count)} "
                                   f"— a chunk row was deleted: {recovery}")

        if verdict == "not_supported" and (any(c["marked"] for c in chunks) or quotes):
            problem(claim_row, "'not supported' with pieces marked — clear the marks or change the verdict")
        elif verdict in ("stated_directly", "needs_combining") and not (any(c["marked"] for c in chunks) or quotes):
            problem(claim_row, f"'{_VERDICT_WORDS[verdict]}' with no piece marked")

        in_bundle = {c["record"]["chunk_id"]: c for c in chunks}
        added = {}
        for record, quote in quotes:
            resolved = _resolve_quote(quote, chunks_by_id)
            if isinstance(resolved, str):
                problem(record, resolved)
                continue
            how = "1 chunk" if len(resolved) == 1 else "2 adjacent chunks — check for repeated text"
            for chunk_id in resolved:
                if chunk_id in in_bundle:
                    in_bundle[chunk_id]["by_quote"] = in_bundle[chunk_id]["by_quote"] or how
                    in_bundle[chunk_id]["quote_notes"].append(record["note"])
                else:
                    entry = added.setdefault(chunk_id, {"chunk_id": chunk_id, "quotes": [], "notes": [], "how": how})
                    entry["quotes"].append(quote)
                    entry["notes"].append(record["note"])

        claims.append({"number": number, "claim_id": claim_id, "verdict": verdict,
                       "ai_verdict": claim_row["ai_verdict"], "ai_reason": claim_row["ai_reason"],
                       "note": claim_row["note"], "plus_notes": plus_notes, "chunks": chunks,
                       "added": list(added.values())})

    for _, claim_text, claim_id in file_claims:
        if claim_id not in seen_claims:
            census = True
            problems.append(f"claims/{memo_id}.md claim {claim_text[:60]!r} ({claim_id}) is not in the "
                            f"sheet (added to the claims file after drafting?)")
    if census:
        problems.append(f"to recover from a changed claims file: delete review/{memo_id}.xlsx, re-run "
                        f"`python golden_set_pipeline.py build`, then `python tag_pipeline.py draft "
                        f"{memo_id}` and review again (answers are not carried over)")
    return problems, claims


# %%
# The golden-set schema reviewed/<memo_id>.xlsx is written in (CLAUDE.md
# "DataFrame schema"). The eval reads these columns by name.
_REVIEWED_COLUMNS = ("claim_id", "memo_id", "section", "claim_text", "doc_id", "chunk_id", "chunk_text",
                     "bm25_score", "evidence_span", "found", "confidence", "ambiguous_match",
                     "verbatim_match", "human_reviewed", "tag", "tag_draft", "tag_rationale")


def _row_tag(verdict: str, needed: bool) -> str:
    """Spec §3: a needed chunk of a stated_directly claim is extractive, of a
    needs_combining claim synthesized; every other chunk is unverifiable."""
    if needed and verdict == "stated_directly":
        return "extractive"
    if needed and verdict == "needs_combining":
        return "synthesized"
    return "unverifiable"


def _notes(*cells) -> str:
    """'; note: <text>' for each non-empty note cell, in the order given."""
    return "".join(f"; note: {_collapse_ws(str(c))}" for c in cells if c is not None and str(c).strip())


def build_reviewed_rows(memo_id: str, claims: list[dict], file_claims: list[tuple[str, str, str]],
                        chunks_by_id: dict[str, dict], run_date: str) -> pd.DataFrame:
    """
    The eval's input for one memo, from check_review's claims (call only when
    it returned no problems): one row per (claim_id, chunk_id) in
    _REVIEWED_COLUMNS, ordered by the claims file, then bundle order, then
    human-added chunks (spec §13.5).

    - claim_id/section/claim_text come from the claims file and memo_id from
      the command; doc_id/chunk_text from the rebuilt index — never from the
      sheet's editable copies. bm25_score, evidence_span, confidence and the
      two match flags have no other source and come from the sheet (type
      checked by check_review).
    - tag: _row_tag on the user's verdict and marks (a quote counts as a
      mark). tag_draft: _row_tag on the AI's ai_verdict/ai_needed; blank for
      a draft_failed claim and for a human-added chunk (the AI never judged
      either).
    - human_reviewed is True on every row: each reaches the file only through
      a CHECKED claim.
    - A claim with no chunks and no resolved quote keeps one found=False row,
      tag unverifiable, tag_rationale starting with the exact
      _AUTO_NOT_FOUND_REASON the eval keys on.
    - tag_rationale then appends '; note: …' for the claim's note and any
      quote-less '+' row note, then the row's own note(s) — including the
      note of a '+' row whose quote marked or added that chunk.
    """
    by_id = {claim["claim_id"]: claim for claim in claims}
    rows = []
    for section, claim_text, claim_id in file_claims:
        claim = by_id[claim_id]
        word = _VERDICT_WORDS[claim["verdict"]]
        ai = claim["ai_verdict"] if claim["ai_verdict"] in _VERDICTS else None
        base = {"claim_id": claim_id, "memo_id": memo_id, "section": section, "claim_text": claim_text,
                "human_reviewed": True}
        for chunk in claim["chunks"]:
            record = chunk["record"]
            real = chunks_by_id[record["chunk_id"]]
            how = "marked" if chunk["marked"] else "marked by quote" if chunk["by_quote"] else "not marked"
            rows.append({**base, "doc_id": real["doc_id"], "chunk_id": record["chunk_id"],
                         "chunk_text": real["chunk_text"], "bm25_score": record["bm25_score"],
                         "evidence_span": record["evidence_span"], "found": True,
                         "confidence": record["confidence"], "ambiguous_match": record["ambiguous_match"],
                         "verbatim_match": record["verbatim_match"],
                         "tag": _row_tag(claim["verdict"], chunk["marked"] or chunk["by_quote"] is not None),
                         "tag_draft": _row_tag(ai, chunk["ai_marked"]) if ai else None,
                         "tag_rationale": (f"bundle review {run_date}: verdict '{word}'; chunk {how}; "
                                           f"AI: {_collapse_ws(str(claim['ai_reason'] or ''))}"
                                           + _notes(claim["note"], *claim["plus_notes"], record["note"],
                                                    *chunk["quote_notes"]))})
        for added in claim["added"]:
            real = chunks_by_id[added["chunk_id"]]
            quoted = "; ".join(f"quote '{_collapse_ws(q)[:60]}'" for q in added["quotes"])
            rows.append({**base, "doc_id": real["doc_id"], "chunk_id": added["chunk_id"],
                         "chunk_text": real["chunk_text"], "bm25_score": None,
                         "evidence_span": _QUOTE_SEPARATOR.join(_collapse_ws(q) for q in added["quotes"]), "found": True,
                         "confidence": None, "ambiguous_match": None, "verbatim_match": None,
                         "tag": _row_tag(claim["verdict"], True), "tag_draft": None,
                         "tag_rationale": (f"{_HUMAN_ADDED_PREFIX} {run_date}: {quoted}; verdict '{word}'"
                                           + _notes(claim["note"], *claim["plus_notes"], *added["notes"]))})
        if not claim["chunks"] and not claim["added"]:
            rows.append({**base, "doc_id": None, "chunk_id": None, "chunk_text": None, "bm25_score": None,
                         "evidence_span": None, "found": False, "confidence": None, "ambiguous_match": None,
                         "verbatim_match": None, "tag": "unverifiable",
                         "tag_draft": _row_tag(ai, False) if ai else None,
                         "tag_rationale": _AUTO_NOT_FOUND_REASON + _notes(claim["note"], *claim["plus_notes"])})
    return pd.DataFrame(rows, columns=list(_REVIEWED_COLUMNS))


# %%
def _numbered_copies(path: str) -> list[str]:
    """Names of '<name> <n>.xlsx' files next to `path` — the copy Excel or
    Numbers saves when the original is open or locked. No folder, no copies."""
    folder = os.path.dirname(path) or "."
    if not os.path.isdir(folder):
        return []
    stem, ext = os.path.splitext(os.path.basename(path))
    pattern = re.compile(re.escape(stem) + r" \d+" + re.escape(ext) + r"$")
    return sorted(name for name in os.listdir(folder) if pattern.match(name))


def _write_reviewed(df: pd.DataFrame, path: str) -> None:
    """
    Write reviewed/<memo_id>.xlsx through golden_set_pipeline's
    export_for_review, overwriting any earlier one (it is a pure function of
    the review sheet). Every value goes through _plain first (export_for_review
    strips nothing, and pypdf text can hold NUL bytes openpyxl refuses).

    export_for_review stores text starting with "=" as a formula, which reads
    back as blank; such cells are re-marked as text before publishing.
    Atomic: written to a unique '<name>.<random>.tmp.xlsx' next to `path`,
    then os.replace. Unique, like write_review_sheet's temp file, so two
    finalize runs of one memo never publish or delete each other's
    half-written file; the extension stays last because export_for_review
    picks the format from it and raises on anything else. A failed write
    leaves no temp file.
    """
    from openpyxl import load_workbook

    folder = os.path.dirname(path) or "."
    os.makedirs(folder, exist_ok=True)
    root, ext = os.path.splitext(os.path.basename(path))
    fd, tmp = tempfile.mkstemp(dir=folder, prefix=f"{root}.", suffix=f".tmp{ext}")
    os.close(fd)
    try:
        export_for_review(df.map(_plain), tmp)
        wb = load_workbook(tmp)
        ws = wb.active
        if ws is None:
            # openpyxl returns None only for a workbook with no sheets at
            # all, which export_for_review's DataFrame.to_excel call above
            # can't produce.
            raise ValueError(f"{tmp}: exported workbook has no active worksheet")
        formulas = [cell for row in ws.iter_rows() for cell in row if cell.data_type == "f"]
        for cell in formulas:
            cell.data_type = "s"
        if formulas:
            wb.save(tmp)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)


def run_finalize(memo_id: str, review_dir: str = "review", reviewed_dir: str = "reviewed",
                 claims_dir: str = "claims", run_date: str | None = None) -> dict:
    """
    `finalize` for one memo (spec §7, §13): read review/<memo_id>.xlsx, check
    it against the claims file and the source PDFs, and write
    reviewed/<memo_id>.xlsx in the golden-set schema. Makes no LLM call.

    Order: `memo_id` must be a plain filename (_valid_memo_id) — it names the
    review, claims and reviewed files, so a path in it could read or overwrite
    files outside those folders; `reviewed_dir` must be a folder and
    reviewed/<memo_id>.xlsx a file, if they exist (otherwise the write would
    fail only after the PDFs were read); then no
    '<name> <n>.xlsx' copy may sit beside the sheet (Excel/Numbers
    save one when the file is open — the user's answers may be in either).
    This is checked first, so a copy left without its original is never
    answered with "run draft", which would pay to re-draft answers that
    exist; then the sheet must exist; then the sheet, the claims file and
    the PDFs are read (always: chunk rows are verified against the rebuilt
    index); then check_review; only with zero problems are rows built and
    written.

    Raises FinalizeRefused for anything that stops the checks from running.
    Returns {"problems": [...], "written": path or None, "claims",
    "verdicts_changed", "marks_changed", "marked_by_quote", "human_added"};
    problems are also logged, one line each, then a refusal line. On success
    logs a summary: counts, each chunk marked by a quote and each human-added
    chunk with how it resolved and the first 80 characters of its text.
    """
    if not _valid_memo_id(memo_id):
        raise FinalizeRefused(f"{memo_id!r} is not a valid memo id (letters, digits, '.', '_' and '-' only)")
    # lexists, not exists: a symlink to nothing must be refused here too.
    if os.path.lexists(reviewed_dir) and not os.path.isdir(reviewed_dir):
        raise FinalizeRefused(f"{reviewed_dir} exists but is not a folder — move it away; finalize writes "
                              f"reviewed files there")
    out = _review_path(reviewed_dir, memo_id)        # same <dir>/<memo_id>.xlsx rule
    if os.path.lexists(out) and not os.path.isfile(out):
        raise FinalizeRefused(f"{out} exists but is not a file — remove it; finalize writes the reviewed file there")
    path = _review_path(review_dir, memo_id)
    copies = _numbered_copies(path)
    if copies:
        raise FinalizeRefused(f"{copies} next to {path}: Excel or Numbers saved a copy. Keep the file with "
                              f"your answers as {os.path.basename(path)}, delete the other, and re-run")
    if not os.path.exists(path):
        raise FinalizeRefused(f"{path} not found — run `python tag_pipeline.py draft {memo_id}` first")
    records = read_review_sheet(path)
    if records and "chunk_count" not in records[0]:
        logger.warning("%s: %s has no chunk_count column (drafted before it was added), so a claim's deleted "
                       "LAST chunk row cannot be detected — check by eye that no claim lost one", memo_id, path)
    if records and "source_docs" not in records[0]:
        logger.warning("%s: %s has no source_docs column (drafted before it was added), so a PDF added to or "
                       "removed from source_folder since drafting cannot be detected — if the PDFs changed, run "
                       "build and draft again", memo_id, path)
    try:
        source_folder, file_claims = read_claims_file(memo_id, claims_dir)
        chunks_by_id, documents = load_memo_chunks(memo_id, source_folder)
    except ValueError as exc:
        raise FinalizeRefused(f"{memo_id}: {exc}") from exc

    problems, claims = check_review(memo_id, records, file_claims, chunks_by_id, documents)
    result = {"problems": problems, "written": None, "claims": len(claims), "verdicts_changed": 0,
              "marks_changed": 0, "marked_by_quote": 0, "human_added": 0}
    if problems:
        for text in problems:
            logger.error("%s: %s", memo_id, text)
        logger.error("%s: refused — %d problem(s); nothing written", memo_id, len(problems))
        return result

    df = build_reviewed_rows(memo_id, claims, file_claims, chunks_by_id, run_date or date.today().isoformat())
    _write_reviewed(df, out)

    chunks = [(claim, chunk) for claim in claims for chunk in claim["chunks"]]
    by_quote = [(claim, chunk) for claim, chunk in chunks if chunk["by_quote"] and not chunk["marked"]]
    result.update(
        written=out,
        verdicts_changed=sum(claim["verdict"] != claim["ai_verdict"] for claim in claims),
        marks_changed=sum((chunk["marked"] or chunk["by_quote"] is not None) != chunk["ai_marked"]
                          for _, chunk in chunks),
        marked_by_quote=len(by_quote),
        human_added=sum(len(claim["added"]) for claim in claims),
    )
    logger.info("%s: written -> %s (%d rows)", memo_id, out, len(df))
    logger.info("%s: %d claims checked; AI verdict kept on %d, changed on %d; marks kept on %d chunk(s), "
                "changed on %d", memo_id, len(claims), len(claims) - result["verdicts_changed"],
                result["verdicts_changed"], len(chunks) - result["marks_changed"], result["marks_changed"])
    for claim, chunk in by_quote:
        logger.info("%s: claim %s: chunk %s marked by quote (%s)", memo_id, claim["number"],
                    chunk["record"]["chunk_id"], chunk["by_quote"])
    for claim in claims:
        for added in claim["added"]:
            logger.info("%s: claim %s: human-added chunk %s (%s): %r", memo_id, claim["number"], added["chunk_id"],
                        added["how"], _collapse_ws(chunks_by_id[added["chunk_id"]]["chunk_text"])[:80])
    return result


# %%
def _bundle_client() -> LLMClient:
    """
    The LLM client `draft` uses: LLM_BASE_URL and LLM_API_KEY from the
    environment, the model pinned to _BUNDLE_MODEL. Not LLMClient.from_env(),
    which also requires LLM_MODEL — a value draft ignores (.env.example), so an
    environment set up only for tagging would be refused for a variable it
    never uses. Raises EnvironmentError naming each missing variable.
    """
    base_url, api_key = os.environ.get("LLM_BASE_URL"), os.environ.get("LLM_API_KEY")
    if not base_url or not api_key:
        missing = [name for name, value in (("LLM_BASE_URL", base_url), ("LLM_API_KEY", api_key)) if not value]
        raise EnvironmentError(f"Missing required environment variable(s): {', '.join(missing)} — draft needs "
                               f"only LLM_BASE_URL and LLM_API_KEY (its model is pinned: {_BUNDLE_MODEL})")
    return LLMClient(base_url=base_url, api_key=api_key, model=_BUNDLE_MODEL)


def _main(argv: list[str]) -> None:
    """
    `python tag_pipeline.py draft [memo_id ...]` — no memo ids means every
    memo in the checkpoint. Every check runs in prepare_draft BEFORE
    _bundle_client(), so a typo or a bad checkpoint needs no credentials
    and costs nothing; a run with no claim to send to the LLM (every memo
    skipped, or every claim with no found chunks) builds no client at
    all. Exits with status 1 if any memo failed.

    `python tag_pipeline.py finalize <memo_id>` — exactly one memo. Makes no
    LLM call and never builds a client. Exits with status 1 if the sheet has
    problems (each is logged); a FinalizeRefused becomes the exit message.
    """
    usage = "usage: python tag_pipeline.py draft [memo_id ...] | finalize <memo_id>"
    if len(argv) < 2:
        raise SystemExit(usage)
    command = argv[1]
    if command == "finalize":
        if len(argv) != 3:
            raise SystemExit(f"{usage}  (finalize takes exactly one memo_id)")
        try:
            result = run_finalize(argv[2])
        except FinalizeRefused as exc:
            raise SystemExit(f"finalize refused: {exc}") from None
        if result["problems"]:
            raise SystemExit(1)
        return
    if command != "draft":
        raise SystemExit(f"{usage}  (unknown command {command!r})")
    try:
        plan = prepare_draft(argv[2:])
    except DraftRefused as exc:
        raise SystemExit(f"draft refused: {exc}") from None
    client = None
    if any(claim["bundle"] for memo in plan if memo["action"] == "draft" for claim in memo["claims"]):
        client = _bundle_client()                 # an empty bundle is answered without a call
    counts = run_draft(plan, client)
    if counts["failed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    _main(sys.argv)
